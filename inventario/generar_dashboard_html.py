#!/usr/bin/env python3
"""Genera dashboard HTML estático del inventario (solo lectura, visual)."""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

SRC = Path("/home/ubuntu/.cursor/projects/workspace/uploads/INVENTARIO_SERV_JOSE_eb9f.xlsx")
OUT = Path("/workspace/inventario/dashboard_inventario.html")


def load_data():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    E, S = defaultdict(float), defaultdict(float)
    meta = {}

    ws = wb["ENTRADA"]
    for r in range(17, ws.max_row + 1):
        cod, cant = ws.cell(r, 2).value, ws.cell(r, 6).value
        if not cod:
            continue
        try:
            E[str(cod).strip()] += float(cant)
        except (TypeError, ValueError):
            pass

    ws = wb["SALIDA"]
    for r in range(18, ws.max_row + 1):
        cod, cant = ws.cell(r, 2).value, ws.cell(r, 6).value
        if not cod:
            continue
        try:
            S[str(cod).strip()] += float(cant)
        except (TypeError, ValueError):
            pass

    products = []
    ws = wb["STOCK"]
    for r in range(16, ws.max_row):
        cod = ws.cell(r, 2).value
        if not cod:
            continue
        cod = str(cod).strip()
        cat = ws.cell(r, 3).value or ""
        name = ws.cell(r, 4).value or ""
        stock = ws.cell(r, 8).value
        try:
            stock_f = float(stock) if stock is not None else None
        except (TypeError, ValueError):
            stock_f = None
        calc = E.get(cod, 0) - S.get(cod, 0)
        products.append(
            {
                "codigo": cod,
                "categoria": str(cat),
                "producto": str(name),
                "stock_excel": stock_f,
                "stock_calc": calc,
                "entradas": E.get(cod, 0),
                "salidas": S.get(cod, 0),
            }
        )

    by_cat = defaultdict(lambda: {"items": 0, "stock": 0.0, "neg": 0})
    for p in products:
        c = p["categoria"] or "SIN CATEGORIA"
        by_cat[c]["items"] += 1
        s = p["stock_calc"]
        by_cat[c]["stock"] += max(s, 0)
        if s < 0:
            by_cat[c]["neg"] += 1

    salidas_tienda = defaultdict(float)
    ws = wb["SALIDA"]
    for r in range(18, ws.max_row + 1):
        tienda = ws.cell(r, 7).value
        cant = ws.cell(r, 6).value
        if not tienda:
            continue
        try:
            salidas_tienda[str(tienda).strip()] += float(cant)
        except (TypeError, ValueError):
            pass

    meta["generado"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    meta["total_skus"] = len(products)
    meta["total_entradas"] = sum(E.values())
    meta["total_salidas"] = sum(S.values())
    meta["sin_stock"] = sum(1 for p in products if p["stock_calc"] == 0)
    meta["negativos"] = sum(1 for p in products if p["stock_calc"] < 0)
    meta["mov_entrada"] = sum(1 for _ in E for __ in [1])  # placeholder
    meta["mov_entrada"] = len(
        [1 for r in range(17, wb["ENTRADA"].max_row + 1) if wb["ENTRADA"].cell(r, 2).value]
    )
    meta["mov_salida"] = len(
        [1 for r in range(18, wb["SALIDA"].max_row + 1) if wb["SALIDA"].cell(r, 2).value]
    )

    top_salidas = sorted(
        [(p["producto"], p["salidas"], p["codigo"]) for p in products if p["salidas"] > 0],
        key=lambda x: -x[1],
    )[:15]

    alertas = sorted(
        [p for p in products if p["stock_calc"] <= 0],
        key=lambda x: x["stock_calc"],
    )[:50]

    return {
        "meta": meta,
        "by_cat": dict(by_cat),
        "top_salidas": top_salidas,
        "salidas_tienda": dict(sorted(salidas_tienda.items(), key=lambda x: -x[1])[:20]),
        "alertas": alertas,
    }


def render(data: dict) -> str:
    payload = json.dumps(data, ensure_ascii=False)
    m = data["meta"]
    return f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Inventario S.S.G.G. — Panel</title>
  <style>
    :root {{
      --bg: #0f172a; --card: #1e293b; --text: #e2e8f0; --muted: #94a3b8;
      --accent: #38bdf8; --ok: #4ade80; --warn: #fbbf24; --bad: #f87171;
    }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; font-family: 'Segoe UI', system-ui, sans-serif; background: var(--bg); color: var(--text); }}
    header {{ padding: 1.5rem 2rem; background: linear-gradient(135deg,#1e3a5f,#0f172a); border-bottom: 1px solid #334155; }}
    h1 {{ margin: 0 0 .25rem; font-size: 1.5rem; }}
    .sub {{ color: var(--muted); font-size: .9rem; }}
    main {{ padding: 1.5rem 2rem 3rem; max-width: 1400px; margin: 0 auto; }}
    .kpis {{ display: grid; grid-template-columns: repeat(auto-fit,minmax(160px,1fr)); gap: 1rem; margin-bottom: 1.5rem; }}
    .kpi {{ background: var(--card); border-radius: 12px; padding: 1rem; border: 1px solid #334155; }}
    .kpi .val {{ font-size: 1.75rem; font-weight: 700; color: var(--accent); }}
    .kpi .lbl {{ font-size: .75rem; text-transform: uppercase; letter-spacing: .05em; color: var(--muted); margin-top: .35rem; }}
    .kpi.bad .val {{ color: var(--bad); }}
    .kpi.warn .val {{ color: var(--warn); }}
    .grid2 {{ display: grid; grid-template-columns: 1fr 1fr; gap: 1rem; }}
    @media (max-width: 900px) {{ .grid2 {{ grid-template-columns: 1fr; }} }}
    section {{ background: var(--card); border-radius: 12px; padding: 1rem 1.25rem; border: 1px solid #334155; margin-bottom: 1rem; }}
    h2 {{ margin: 0 0 1rem; font-size: 1rem; color: var(--accent); }}
    table {{ width: 100%; border-collapse: collapse; font-size: .85rem; }}
    th, td {{ padding: .45rem .5rem; text-align: left; border-bottom: 1px solid #334155; }}
    th {{ color: var(--muted); font-weight: 600; }}
    .bar-wrap {{ display: flex; align-items: center; gap: .5rem; }}
    .bar {{ height: 8px; background: #334155; border-radius: 4px; flex: 1; overflow: hidden; }}
    .bar > i {{ display: block; height: 100%; background: var(--accent); border-radius: 4px; }}
    input {{ width: 100%; padding: .6rem .75rem; border-radius: 8px; border: 1px solid #475569; background: #0f172a; color: var(--text); margin-bottom: .75rem; }}
    .tag {{ display: inline-block; padding: .15rem .45rem; border-radius: 4px; font-size: .7rem; }}
    .tag.neg {{ background: #7f1d1d; color: #fecaca; }}
    .tag.zero {{ background: #78350f; color: #fde68a; }}
  </style>
</head>
<body>
  <header>
    <h1>Inventario S.S.G.G. — Panel de control</h1>
    <p class="sub">Generado: {m["generado"]} · Solo lectura (datos del Excel original)</p>
  </header>
  <main>
    <div class="kpis">
      <div class="kpi"><div class="val">{m["total_skus"]:,.0f}</div><div class="lbl">SKUs</div></div>
      <div class="kpi"><div class="val">{m["total_entradas"]:,.0f}</div><div class="lbl">Unidades ingresadas</div></div>
      <div class="kpi"><div class="val">{m["total_salidas"]:,.0f}</div><div class="lbl">Unidades egresadas</div></div>
      <div class="kpi warn"><div class="val">{m["sin_stock"]:,.0f}</div><div class="lbl">Sin stock</div></div>
      <div class="kpi bad"><div class="val">{m["negativos"]:,.0f}</div><div class="lbl">Stock negativo</div></div>
      <div class="kpi"><div class="val">{m["mov_salida"]:,.0f}</div><div class="lbl">Mov. salida</div></div>
    </div>
    <div class="grid2">
      <section><h2>Salidas por tienda (top)</h2><div id="tiendas"></div></section>
      <section><h2>Top consumo (salidas acumuladas)</h2><div id="top"></div></section>
    </div>
    <section>
      <h2>Alertas — sin stock o negativo</h2>
      <input type="search" id="q" placeholder="Buscar código o producto…" />
      <table><thead><tr><th>Código</th><th>Producto</th><th>Categoría</th><th>Entradas</th><th>Salidas</th><th>Stock calc.</th></tr></thead><tbody id="alertas"></tbody></table>
    </section>
  </main>
  <script>
    const DATA = {payload};
    const maxT = Math.max(...Object.values(DATA.salidas_tienda), 1);
    document.getElementById('tiendas').innerHTML = Object.entries(DATA.salidas_tienda).map(([t,v]) => {{
      const pct = (100*v/maxT).toFixed(0);
      return `<div class="bar-wrap"><span style="width:140px;font-size:.8rem">${{t}}</span><div class="bar"><i style="width:${{pct}}%"></i></div><strong>${{v.toLocaleString('es')}}</strong></div>`;
    }}).join('');
    document.getElementById('top').innerHTML = '<table><thead><tr><th>Producto</th><th>Salidas</th></tr></thead><tbody>' +
      DATA.top_salidas.map(([n,s,c]) => `<tr><td>${{n}}<br><small style="color:#94a3b8">${{c}}</small></td><td>${{s.toLocaleString('es')}}</td></tr>`).join('') + '</tbody></table>';
    function renderAlertas(filter='') {{
      const f = filter.toLowerCase();
      const rows = DATA.alertas.filter(p => !f || p.codigo.toLowerCase().includes(f) || p.producto.toLowerCase().includes(f));
      document.getElementById('alertas').innerHTML = rows.map(p => {{
        const tag = p.stock_calc < 0 ? '<span class="tag neg">NEG</span>' : '<span class="tag zero">0</span>';
        return `<tr><td>${{p.codigo}}</td><td>${{p.producto}}</td><td>${{p.categoria}}</td><td>${{p.entradas}}</td><td>${{p.salidas}}</td><td>${{p.stock_calc}} ${{tag}}</td></tr>`;
      }}).join('');
    }}
    renderAlertas();
    document.getElementById('q').addEventListener('input', e => renderAlertas(e.target.value));
  </script>
</body>
</html>"""


def main():
    data = load_data()
    OUT.write_text(render(data), encoding="utf-8")
    print(f"Dashboard: {OUT}")


if __name__ == "__main__":
    main()
