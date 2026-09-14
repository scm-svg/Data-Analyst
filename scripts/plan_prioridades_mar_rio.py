#!/usr/bin/env python3
"""Plan de prioridades Mar / Rio según tela Jabón disponible.

Reglas:
1. Usar cantidades MÁXIMAS sugeridas (alerta de planificación).
2. Prioridad KIDS en ambos modelos. Objetivo = 50% del pendiente
   (Mar Kids descuenta el lote ya cortado: 1.300 und).
3. El lote cortado son UND de producción (WIP). La foto de tela es
   inventario actual y se usa completa.
4. Tela restante → lotes CAB / DAMA de Mar y Rio, con flexibilidad
   por color (reparto proporcional a demanda de kg, sin pasar el máx).
5. Rojo (0,18 kg) se excluye. Morado de inventario cubre Púrpura.
"""
from __future__ import annotations

import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUT = ROOT / "output"
MAR_XLSX = DATA / "MAR_PROYECCION_CANTIDADES_SUGERIDAS.xlsx"
RIO_XLSX = DATA / "RIO_PROYECCION_CANTIDADES_SUGERIDAS.xlsx"

# Máximos del Excel de alerta "CANTIDADES SUGERIDAS (2)" — MAR.
# RIO se lee del archivo (coincide con la alerta).
MAR_MAX_COLORES = {
    "CAB": {
        "Negro": 555, "Azul Marino": 473, "Blanco": 390, "Verde Militar": 355,
        "Azul Lavanda": 268, "Azul Rey": 205, "Gris Claro": 190, "Aguamarina": 151,
        "Vinotinto": 124, "Rojo": 122, "Amarillo Neón": 67,
    },
    "DAMA": {
        "Negro": 417, "Blanco": 393, "Azul Marino": 357, "Verde Militar": 272,
        "Azul Lavanda": 232, "Púrpura": 196, "Lila": 193, "Rojo": 168,
        "Vinotinto": 162, "Aguamarina": 135, "Rosado Pastel": 108, "Amarillo Neón": 67,
    },
    "KIDS": {
        "Blanco": 460, "Azul Marino": 409, "Negro": 379, "Lila": 254, "Aguamarina": 243,
        "Rojo": 239, "Azul Rey": 232, "Verde Militar": 224, "Rosado Pastel": 214,
        "Púrpura": 180, "Azul Lavanda": 153, "Amarillo Neón": 114,
    },
}

TELA_KG = {
    "Aguamarina": 348.34,
    "Amarillo Neón": 57.94,
    "Azul Lavanda": 221.00,
    "Azul Marino": 689.90,
    "Azul Rey": 271.46,
    "Blanco": 162.04,
    "Gris Claro": 47.68,
    "Lila": 188.44,
    "Morado": 94.20,
    "Negro": 657.16,
    "Púrpura": 0.00,
    "Rojo": 0.18,
    "Rosado Pastel": 31.44,
    "Verde Militar": 46.02,
    "Vinotinto": 175.38,
}

LOTE_COLORES = [
    "Verde Militar", "Vinotinto", "Azul Marino", "Azul Rey", "Rojo",
    "Gris Claro", "Azul Lavanda", "Aguamarina", "Rosado Pastel", "Negro", "Lila",
]
LOTE_TALLAS = {"8": 260, "10": 260, "12": 260, "14": 520}
LOTE_TOTAL = sum(LOTE_TALLAS.values())
PCT_KIDS = 0.50
MIN_KG = 5.0
MIN_LOT_ADULT = 10
SKIP_TALLAS = {"1"}

COLOR_HEX = {
    "Aguamarina": "#2ec4b6", "Amarillo Neón": "#c6ff00", "Azul Lavanda": "#9b8ec4",
    "Azul Marino": "#1b365d", "Azul Rey": "#3d5af1", "Blanco": "#e8e8ee",
    "Gris Claro": "#9aa0a6", "Lila": "#c77dff", "Morado": "#7b2cbf",
    "Negro": "#2a2a32", "Púrpura": "#7b2cbf", "Rojo": "#e63946",
    "Rosado Pastel": "#f4a6c1", "Verde Militar": "#4a7c23", "Vinotinto": "#6b1d2a",
}

PRIORIDAD = ["MAR KIDS", "RIO KIDS", "MAR CAB", "MAR DAMA", "RIO CAB", "RIO DAMA"]


def stock_key(color: str) -> str:
    return "Morado" if color == "Púrpura" else color


def tela_de(color: str) -> float:
    if color == "Rojo":
        return 0.0
    return float(TELA_KG.get(stock_key(color), 0.0))


def largest_remainder(total: int, weights: dict[str, float]) -> dict[str, int]:
    keys = list(weights)
    pos = {k: max(0.0, float(weights[k])) for k in keys}
    s = sum(pos.values())
    if total <= 0 or s <= 0:
        return {k: 0 for k in keys}
    raw = {k: total * pos[k] / s for k in keys}
    floors = {k: int(math.floor(raw[k])) for k in keys}
    rem = total - sum(floors.values())
    order = sorted(keys, key=lambda k: (raw[k] - floors[k]), reverse=True)
    for k in order:
        if rem <= 0:
            break
        floors[k] += 1
        rem -= 1
    return floors


def parse_consumo(path: Path) -> dict[str, float]:
    ws = openpyxl.load_workbook(path, data_only=True)["Compra de Tela"]
    out: dict[str, float] = {}
    for row in ws.iter_rows(values_only=True):
        v0 = str(row[0] or "").strip()
        if v0 in ("CABALLERO", "DAMA", "KIDS") and row[1] is not None:
            key = "CAB" if v0 == "CABALLERO" else v0
            out[key] = float(row[1])
    return out


def parse_colores(path: Path) -> dict[str, dict[str, int]]:
    ws = openpyxl.load_workbook(path, data_only=True)["Cantidades por Colores"]
    result = {"CAB": {}, "DAMA": {}, "KIDS": {}}
    current = None
    for row in ws.iter_rows(values_only=True):
        v0 = str(row[0] or "").strip()
        u = v0.upper()
        if "CABALLERO" in u:
            current = "CAB"
            continue
        if "DAMA" in u and "CAB" not in u:
            current = "DAMA"
            continue
        if "KIDS" in u:
            current = "KIDS"
            continue
        if not current or v0 in ("Color", "TOTAL", "") or row[3] is None:
            continue
        try:
            result[current][v0] = int(float(row[3]))
        except (TypeError, ValueError):
            pass
    return result


def parse_tallas(path: Path) -> dict[str, dict[str, int]]:
    ws = openpyxl.load_workbook(path, data_only=True)["Producción por Talla"]
    result = {"CAB": {}, "DAMA": {}, "KIDS": {}}
    current = None
    for row in ws.iter_rows(values_only=True):
        v0 = str(row[0] or "").strip()
        u = v0.upper()
        if "CABALLERO" in u:
            current = "CAB"
            continue
        if "DAMA" in u and "CAB" not in u:
            current = "DAMA"
            continue
        if "KIDS" in u:
            current = "KIDS"
            continue
        if not current or v0 in ("Talla", "TOTAL", "") or v0 in SKIP_TALLAS or row[3] is None:
            continue
        try:
            result[current][v0] = int(float(row[3]))
        except (TypeError, ValueError):
            pass
    return result


def parse_color_talla(path: Path) -> dict[str, dict[str, dict[str, int]]]:
    ws = openpyxl.load_workbook(path, data_only=True)["Producción Color × Talla"]
    result = {"CAB": {}, "DAMA": {}, "KIDS": {}}
    current = None
    sizes: list[str] = []
    for row in ws.iter_rows(values_only=True):
        v0 = str(row[0] or "").strip()
        u = v0.upper()
        if "CABALLERO" in u:
            current, sizes = "CAB", []
            continue
        if v0.startswith("DAMA"):
            current, sizes = "DAMA", []
            continue
        if v0.startswith("KIDS"):
            current, sizes = "KIDS", []
            continue
        if v0.startswith("Color"):
            sizes = [str(x) for x in row[1:] if x not in (None, "")]
            continue
        if not current or ("Máx" not in v0 and "Max" not in v0):
            continue
        color = v0.replace("— Máx", "").replace("— Max", "").replace("- Máx", "").strip()
        vals = {}
        for i, s in enumerate(sizes):
            if s.lower() == "total" or s in SKIP_TALLAS:
                continue
            val = row[i + 1] if i + 1 < len(row) else None
            if val is None:
                continue
            try:
                n = int(float(val))
            except (TypeError, ValueError):
                continue
            if n > 0:
                vals[str(s)] = n
        if vals:
            result[current][color] = vals
    return result


def scale_matrix(matrix: dict[str, int], new_total: int) -> dict[str, int]:
    return largest_remainder(new_total, {k: float(v) for k, v in matrix.items()})


def load_modelos() -> dict:
    mar_ct = parse_color_talla(MAR_XLSX)
    rio_ct = parse_color_talla(RIO_XLSX)
    mar_col_file = parse_colores(MAR_XLSX)
    rio_col = parse_colores(RIO_XLSX)
    modelos = {
        "MAR": {
            "consumo": parse_consumo(MAR_XLSX),
            "tallas": parse_tallas(MAR_XLSX),
            "colores": MAR_MAX_COLORES,
            "color_talla": {},
        },
        "RIO": {
            "consumo": parse_consumo(RIO_XLSX),
            "tallas": parse_tallas(RIO_XLSX),
            "colores": rio_col,
            "color_talla": rio_ct,
        },
    }
    # Escalar curva color×talla MAR del Excel de proyección a los máx de la alerta.
    for gen, colores in MAR_MAX_COLORES.items():
        modelos["MAR"]["color_talla"][gen] = {}
        for color, mx in colores.items():
            src = mar_ct.get(gen, {}).get(color) or mar_col_file.get(gen, {})
            if isinstance(src, dict) and src and color in mar_ct.get(gen, {}):
                modelos["MAR"]["color_talla"][gen][color] = scale_matrix(mar_ct[gen][color], mx)
            else:
                curve = {k: v for k, v in modelos["MAR"]["tallas"][gen].items() if k not in SKIP_TALLAS}
                modelos["MAR"]["color_talla"][gen][color] = scale_matrix(curve, mx)
    return modelos


def lote_por_color() -> dict[str, float]:
    n = len(LOTE_COLORES)
    return {c: LOTE_TOTAL / n for c in LOTE_COLORES}


def lote_por_color_talla() -> dict[str, dict[str, float]]:
    por_color = lote_por_color()
    out = {}
    for color, und in por_color.items():
        out[color] = {t: und * qty / LOTE_TOTAL for t, qty in LOTE_TALLAS.items()}
    return out


def remaining_curve(color_talla: dict[str, int], cut: dict[str, float] | None) -> dict[str, float]:
    rem = {t: float(v) for t, v in color_talla.items() if t not in SKIP_TALLAS}
    if cut:
        for t, u in cut.items():
            rem[t] = max(0.0, rem.get(t, 0.0) - u)
    return rem


def plan(modelos: dict) -> dict:
    saldo = {c: float(kg) for c, kg in TELA_KG.items()}
    saldo["Púrpura"] = saldo["Morado"]
    saldo["Rojo"] = 0.0

    cortado = lote_por_color()
    cortado_talla = lote_por_color_talla()
    filas = []
    lots = []

    def take(color: str, und: int, consumo: float) -> tuple[int, float]:
        key = stock_key(color)
        kg_disp = saldo.get(key, 0.0)
        max_und = int(kg_disp / consumo) if consumo > 0 else 0
        und = max(0, min(und, max_und))
        kg = round(und * consumo, 4)
        saldo[key] = max(0.0, kg_disp - kg)
        if color == "Púrpura":
            saldo["Morado"] = saldo[key]
        return und, kg

    # ── KIDS (prioridad estricta Mar → Rio) ──
    for prio, (modelo, genero) in enumerate((("MAR", "KIDS"), ("RIO", "KIDS")), 1):
        consumo = modelos[modelo]["consumo"][genero]
        for color, meta in sorted(modelos[modelo]["colores"][genero].items(), key=lambda x: -x[1]):
            ya = cortado.get(color, 0.0) if modelo == "MAR" else 0.0
            # Colores del lote sin meta kids (Vinotinto, Gris Claro) no se descuentan aquí.
            if modelo == "MAR" and color not in modelos["MAR"]["colores"]["KIDS"]:
                ya = 0.0
            if modelo == "MAR" and color not in LOTE_COLORES:
                ya = 0.0
            pendiente = max(0, int(round(meta - ya)))
            objetivo = int(pendiente * PCT_KIDS)
            kg_before = saldo.get(stock_key(color), 0.0)
            estado_tela = "SIN TELA" if kg_before < MIN_KG else "OK"
            und, kg = (0, 0.0) if estado_tela == "SIN TELA" else take(color, objetivo, consumo)
            cut_ct = cortado_talla.get(color) if modelo == "MAR" and color in LOTE_COLORES else None
            curve = remaining_curve(modelos[modelo]["color_talla"][genero].get(color, {}), cut_ct)
            if modelo == "MAR":
                # Talla 14 ya cubierta a nivel lote (520 ≈ máximo sugerido).
                curve["14"] = 0.0
            dist = largest_remainder(und, curve)
            estado = (
                "SIN TELA" if und == 0 and objetivo > 0
                else "PARCIAL" if und < objetivo
                else "OK"
            )
            fila = {
                "prioridad": prio, "orden": f"{modelo} {genero}", "modelo": modelo,
                "genero": genero, "color": color, "meta_max": meta,
                "ya_cortado": round(ya, 1), "pendiente": pendiente,
                "objetivo": objetivo, "und": und, "kg": round(kg, 2),
                "consumo": consumo, "tela_antes": round(kg_before, 2),
                "tela_despues": round(saldo.get(stock_key(color), 0.0), 2),
                "estado": estado, "tallas": dist,
            }
            filas.append(fila)
            if und:
                lots.append(fila)

    # ── Adultos: flexibilidad por color ──
    adultos = []
    for modelo in ("MAR", "RIO"):
        for genero in ("CAB", "DAMA"):
            prio = PRIORIDAD.index(f"{modelo} {genero}") + 1
            consumo = modelos[modelo]["consumo"][genero]
            for color, meta in modelos[modelo]["colores"][genero].items():
                adultos.append({
                    "prioridad": prio, "orden": f"{modelo} {genero}",
                    "modelo": modelo, "genero": genero, "color": color,
                    "meta_max": meta, "ya_cortado": 0, "pendiente": meta,
                    "objetivo": meta, "consumo": consumo,
                })

    by_color: dict[str, list] = defaultdict(list)
    for a in adultos:
        by_color[a["color"]].append(a)

    for color, grupo in by_color.items():
        key = stock_key(color)
        kg_disp = saldo.get(key, 0.0)
        kg_before = kg_disp
        demandas = []
        for a in grupo:
            kg_need = a["objetivo"] * a["consumo"]
            demandas.append((a, kg_need))
        total_need = sum(k for _, k in demandas) or 1.0

        if kg_disp < MIN_KG:
            for a, _ in demandas:
                a.update(und=0, kg=0.0, tela_antes=round(kg_before, 2),
                         tela_despues=round(kg_disp, 2), estado="SIN TELA", tallas={})
                filas.append(a)
            continue

        asignados = []
        used = 0.0
        for a, kg_need in demandas:
            share = kg_disp * (kg_need / total_need)
            und = min(a["objetivo"], int(share / a["consumo"]))
            kg = und * a["consumo"]
            asignados.append([a, und, kg])
            used += kg

        leftover = kg_disp - used
        # Completar con tela sobrante (flexibilidad: llena el hueco de mayor demanda).
        progressed = True
        while progressed:
            progressed = False
            order = sorted(asignados, key=lambda x: -((x[0]["objetivo"] - x[1]) * x[0]["consumo"]))
            for item in order:
                a, und, kg = item
                if und >= a["objetivo"]:
                    continue
                if leftover + 1e-9 < a["consumo"]:
                    continue
                item[1] += 1
                item[2] += a["consumo"]
                leftover -= a["consumo"]
                progressed = True
                break

        # Lotes adultos < 10 und no son cortables: consolidar en un solo modelo/género.
        small = [item for item in asignados if 0 < item[1] < MIN_LOT_ADULT]
        if small:
            released = 0.0
            for item in small:
                released += item[2]
                leftover += item[2]
                item[1], item[2] = 0, 0.0
            best = sorted(
                asignados,
                key=lambda x: min(x[0]["objetivo"] - x[1], int(leftover / x[0]["consumo"])) * x[0]["consumo"],
                reverse=True,
            )[0]
            extra = min(best[0]["objetivo"] - best[1], int(leftover / best[0]["consumo"]))
            if extra >= MIN_LOT_ADULT or (best[1] + extra >= MIN_LOT_ADULT and extra > 0):
                best[1] += extra
                best[2] += extra * best[0]["consumo"]
                leftover -= extra * best[0]["consumo"]
            # si extra < 10 y el receptor también era 0, se deja como merma de tela

        for a, und, kg in asignados:
            saldo[key] = max(0.0, saldo.get(key, 0.0) - kg)
            if color == "Púrpura":
                saldo["Morado"] = saldo[key]
            curve = remaining_curve(modelos[a["modelo"]]["color_talla"][a["genero"]].get(color, {}), None)
            dist = largest_remainder(und, curve)
            estado = "SIN TELA" if und == 0 else ("PARCIAL" if und < a["objetivo"] else "OK")
            a.update(
                und=und, kg=round(kg, 2), tela_antes=round(kg_before, 2),
                tela_despues=round(saldo.get(key, 0.0), 2), estado=estado, tallas=dist,
            )
            filas.append(a)
            if und:
                lots.append(a)

    # Resumen por bloque
    por_prio = {}
    for f in filas:
        key = f["orden"]
        por_prio.setdefault(key, {"und": 0, "kg": 0.0, "objetivo": 0})
        por_prio[key]["und"] += f["und"]
        por_prio[key]["kg"] += f["kg"]
        por_prio[key]["objetivo"] += f["objetivo"]

    flex = []
    for color in sorted(set(TELA_KG) | set(MAR_MAX_COLORES["KIDS"]) | set(MAR_MAX_COLORES["DAMA"])):
        if color == "Púrpura":
            continue
        display = color
        kg0 = TELA_KG.get(color, 0.0)
        used = kg0 - saldo.get(color, 0.0)
        row = {"color": display, "stock_kg": kg0, "usado_kg": round(used, 2),
               "resta_kg": round(saldo.get(color, 0.0), 2)}
        for ord_ in PRIORIDAD:
            row[ord_] = sum(f["und"] for f in filas if f["color"] in (color, "Púrpura") and f["orden"] == ord_ and (color != "Morado" or f["color"] == "Púrpura" or f["color"] == "Morado"))
        # Fix Morado/Púrpura mapping in matrix
        if color == "Morado":
            for ord_ in PRIORIDAD:
                row[ord_] = sum(f["und"] for f in filas if f["color"] == "Púrpura" and f["orden"] == ord_)
        else:
            for ord_ in PRIORIDAD:
                row[ord_] = sum(f["und"] for f in filas if f["color"] == color and f["orden"] == ord_)
        und_total = sum(row[o] for o in PRIORIDAD)
        if kg0 < MIN_KG:
            rec = "Excluir — sin tela (comprar o esperar)"
        elif row["resta_kg"] < 8 and und_total:
            rec = "Tela agotada en este plan — sin margen de swap"
        elif kg0 >= 400:
            rec = "Alta flexibilidad: se puede mover CAB/DAMA entre Mar y Rio"
        elif kg0 >= 150:
            rec = "Flex media: kids 50% + adultos según demanda"
        elif kg0 >= 40:
            rec = "Escasa: priorizar KIDS; adultos solo si sobra"
        else:
            rec = "Muy escasa: solo KIDS parcial"
        row["recomendacion"] = rec
        row["und_total"] = und_total
        flex.append(row)

    return {
        "filas": filas,
        "lots": lots,
        "por_prioridad": por_prio,
        "flex": flex,
        "saldo": {k: round(v, 2) for k, v in saldo.items()},
        "cortado": cortado,
        "cortado_talla": cortado_talla,
        "consumo": {m: modelos[m]["consumo"] for m in modelos},
        "metas": {m: {g: sum(modelos[m]["colores"][g].values()) for g in modelos[m]["colores"]} for m in modelos},
        "tela_total": round(sum(TELA_KG.values()), 2),
        "tela_usable": round(sum(v for k, v in TELA_KG.items() if k not in ("Rojo", "Púrpura")), 2),
        "total_und": sum(f["und"] for f in filas),
        "total_kg": round(sum(f["kg"] for f in filas), 2),
        "lote_total": LOTE_TOTAL,
        "lote_tallas": LOTE_TALLAS,
        "lote_colores": LOTE_COLORES,
    }


def write_excel(plan_data: dict, path: Path) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    kids_fill = PatternFill("solid", fgColor="FFF2CC")
    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    alert_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Border(
        left=Side(style="thin", color="AAAAAA"),
        right=Side(style="thin", color="AAAAAA"),
        top=Side(style="thin", color="AAAAAA"),
        bottom=Side(style="thin", color="AAAAAA"),
    )

    def hdr(ws, row, n):
        for c in range(1, n + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin

    def autosize(ws, cap=28):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = min(max(len(str(c.value or "")) for c in col) + 2, cap)
            ws.column_dimensions[letter].width = width

    ws = wb.active
    ws.title = "RESUMEN"
    ws["A1"] = "PLAN DE PRIORIDADES — MAR / RIO · TELA JABÓN"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Tela foto (kg)", plan_data["tela_total"]),
        ("Tela usable (kg, sin Rojo/Púrpura 0)", plan_data["tela_usable"]),
        ("Kg asignados en este plan", plan_data["total_kg"]),
        ("Unidades a producir (nuevas)", plan_data["total_und"]),
        ("Lote Mar Kids YA CORTADO (und)", plan_data["lote_total"]),
        ("Objetivo Kids", "50% del pendiente (máx sugerido − cortado)"),
        ("Consumo MAR kg/und", "CAB 0.50 · DAMA 0.40 · KIDS 0.26"),
        ("Consumo RIO kg/und", "CAB 0.66 · DAMA 0.52 · KIDS 0.32"),
        ("", ""),
        ("Bloque", "Und plan / Objetivo / Kg"),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
        r += 1
    for i, nombre in enumerate(PRIORIDAD, 1):
        d = plan_data["por_prioridad"].get(nombre, {"und": 0, "kg": 0, "objetivo": 0})
        ws.cell(row=r, column=1, value=f"P{i} {nombre}")
        ws.cell(row=r, column=2, value=f"{d['und']} und  /  obj {d['objetivo']}  /  {d['kg']:.1f} kg")
        if "KIDS" in nombre:
            ws.cell(row=r, column=1).fill = kids_fill
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="REGLAS").font = Font(bold=True)
    r += 1
    for line in [
        "1. Prioridad Kids: Mar primero (descuenta 1.300 und ya cortadas), luego Rio. Meta = 50% del pendiente.",
        "2. Talla 14 Mar Kids ya está cubierta (~520 cortadas ≈ máximo). El lote nuevo rellena 2-6-8-10-12.",
        "3. Vinotinto y Gris Claro se cortaron en kids pero no están en la meta kids del Excel: no se descuentan de otros colores.",
        "4. Adultos (CAB/DAMA) se reparte la tela restante por color, proporcional a kg de demanda máxima, con swap Mar↔Rio.",
        "5. Rojo 0,18 kg: no producir. Morado de inventario (94,2 kg) cubre Púrpura del Excel.",
        "6. Blanco, Verde Militar, Rosado Pastel y Gris Claro son cuellos de botella: primero kids.",
    ]:
        ws.cell(row=r, column=1, value=line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
    autosize(ws, 90)

    ws2 = wb.create_sheet("INVENTARIO TELA")
    h = ["Color", "Stock kg", "Usado plan kg", "Resta kg", "Mar Kids", "Rio Kids", "Mar CAB", "Mar DAMA", "Rio CAB", "Rio DAMA", "Recomendación"]
    ws2.append(h)
    hdr(ws2, 1, len(h))
    for row in plan_data["flex"]:
        ws2.append([
            row["color"], row["stock_kg"], row["usado_kg"], row["resta_kg"],
            row["MAR KIDS"], row["RIO KIDS"], row["MAR CAB"], row["MAR DAMA"],
            row["RIO CAB"], row["RIO DAMA"], row["recomendacion"],
        ])
    autosize(ws2, 55)

    ws3 = wb.create_sheet("MAR KIDS CORTADO")
    ws3.append(["Talla", "Und cortadas"])
    hdr(ws3, 1, 2)
    for t, u in LOTE_TALLAS.items():
        ws3.append([t, u])
    ws3.append(["TOTAL", LOTE_TOTAL])
    ws3.append([])
    ws3.append(["Color", "Und cortadas (prop.)", "En meta KIDS Excel", "Nota"])
    hdr(ws3, ws3.max_row, 4)
    meta_kids = set(MAR_MAX_COLORES["KIDS"])
    for c, u in lote_por_color().items():
        flag = "Sí" if c in meta_kids else "No — extra vs Excel kids"
        ws3.append([c, round(u, 1), flag, "Reparto igual entre 11 colores"])
    autosize(ws3)

    ws4 = wb.create_sheet("PRIORIDADES")
    cols = ["P", "Orden", "Modelo", "Género", "Color", "Máx sugerido", "Ya cortado",
            "Pendiente", "Objetivo", "Und asignar", "Kg", "Tela resta kg", "Estado"]
    ws4.append(cols)
    hdr(ws4, 1, len(cols))
    for f in plan_data["filas"]:
        ws4.append([
            f["prioridad"], f["orden"], f["modelo"], f["genero"], f["color"],
            f["meta_max"], f["ya_cortado"], f["pendiente"], f["objetivo"],
            f["und"], f["kg"], f["tela_despues"], f["estado"],
        ])
        r = ws4.max_row
        if f["genero"] == "KIDS":
            for c in range(1, len(cols) + 1):
                ws4.cell(row=r, column=c).fill = kids_fill
        fill = ok_fill if f["estado"] == "OK" else (alert_fill if f["estado"] != "PARCIAL" else PatternFill("solid", fgColor="FCE4D6"))
        ws4.cell(row=r, column=13).fill = fill
    autosize(ws4)

    ws5 = wb.create_sheet("LOTES COLOR x TALLA")
    ws5.append(["Prioridad", "Modelo", "Género", "Color", "Talla", "Und", "Kg color", "Estado"])
    hdr(ws5, 1, 8)
    size_order = ["2", "4", "6", "8", "10", "12", "14", "XS", "S", "M", "L", "XL", "2XL"]
    for f in plan_data["filas"]:
        if not f["und"]:
            continue
        tallas = f.get("tallas") or {}
        ordered = [t for t in size_order if t in tallas] + [t for t in tallas if t not in size_order]
        for t in ordered:
            if tallas[t] <= 0:
                continue
            ws5.append([f["prioridad"], f["modelo"], f["genero"], f["color"], t, tallas[t], f["kg"], f["estado"]])
    autosize(ws5)

    wb.save(path)


def esc(s) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_html(plan_data: dict, path: Path) -> None:
    payload = json.dumps({
        "filas": plan_data["filas"],
        "flex": plan_data["flex"],
        "por": plan_data["por_prioridad"],
        "kpis": {
            "und": plan_data["total_und"],
            "kg": plan_data["total_kg"],
            "tela": plan_data["tela_usable"],
            "cortado": plan_data["lote_total"],
            "metas": plan_data["metas"],
            "consumo": plan_data["consumo"],
        },
        "lote_tallas": plan_data["lote_tallas"],
        "cortado": plan_data["cortado"],
        "colors": COLOR_HEX,
    }, ensure_ascii=False)
    template = (Path(__file__).with_name("plan_prioridades_template.html")).read_text(encoding="utf-8")
    path.write_text(template.replace("__PAYLOAD__", payload), encoding="utf-8")



def assert_plan(plan_data: dict, modelos: dict) -> None:
    used = defaultdict(float)
    for f in plan_data["filas"]:
        used[stock_key(f["color"])] += f["kg"]
        assert f["und"] >= 0
        assert f["und"] <= f["objetivo"] + 1  # rounding
        assert f["und"] <= f["meta_max"] + 1
        if f["tallas"]:
            assert sum(f["tallas"].values()) == f["und"]
    for color, kg_used in used.items():
        stock = TELA_KG.get(color, 0.0)
        assert kg_used <= stock + 0.05, (color, kg_used, stock)
    # Mar kids talla 14 should be ~0 new units overall
    t14 = sum((f.get("tallas") or {}).get("14", 0) for f in plan_data["filas"] if f["orden"] == "MAR KIDS")
    assert t14 == 0, t14
    print("OK asserts · talla 14 Mar Kids nueva =", t14)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    modelos = load_modelos()
    plan_data = plan(modelos)
    assert_plan(plan_data, modelos)

    xlsx = OUT / "PLAN_PRIORIDADES_MAR_RIO.xlsx"
    html = ROOT / "PLAN_PRIORIDADES_MAR_RIO.html"
    js = OUT / "plan_prioridades.json"
    write_excel(plan_data, xlsx)
    write_html(plan_data, html)
    dump = {k: v for k, v in plan_data.items()}
    js.write_text(json.dumps(dump, ensure_ascii=False, indent=2), encoding="utf-8")

    print("Excel:", xlsx)
    print("HTML:", html)
    print("JSON:", js)
    print("Total und:", plan_data["total_und"], "kg:", plan_data["total_kg"])
    for k, v in plan_data["por_prioridad"].items():
        cubre = (100 * v["und"] / v["objetivo"]) if v["objetivo"] else 0
        print(f"  {k:10s}  {v['und']:5d} und  obj {v['objetivo']:5d}  {cubre:5.1f}%  {v['kg']:7.1f} kg")


if __name__ == "__main__":
    main()
