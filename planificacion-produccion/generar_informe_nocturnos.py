#!/usr/bin/env python3
"""Genera el informe Excel + HTML de escenarios de turno nocturno A–D."""
from __future__ import annotations

import json
import os
import shutil
import html as htmlmod
from datetime import date
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from escenarios_nocturnos import (
    BONO_USD,
    CORTE_ALMACEN,
    CORTE_COSTURA_ALMACEN,
    CORTE_COSTURA_DICIEMBRE,
    ESCENARIOS,
    PERSONAS,
    celdas_a_dict,
    comparar_lotes,
    cargar_plan,
    etiquetas_semana,
    fmt_fecha,
    kpis_escenario,
    modelos_a_dict,
    personas_total,
    resumen_semanas,
    simular_nocturnos,
)

NUEVO_XLSX = os.environ.get(
    "PLAN_NUEVO_XLSX",
    "/home/ubuntu/.cursor/projects/workspace/uploads/Planificacion_Produccion_LOTE_NUEVO_f414.xlsx",
)
ACTUAL_XLSX = os.environ.get(
    "PLAN_ACTUAL_XLSX",
    "/home/ubuntu/.cursor/projects/workspace/uploads/Planificacion_Produccion_ACTUAL_15df.xlsx",
)
OUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__)))
OUT_XLSX = os.path.join(OUT_DIR, "Informe_Turnos_Nocturnos.xlsx")
OUT_HTML = os.path.join(OUT_DIR, "Informe_Turnos_Nocturnos.html")
ARTIFACT_DIR = "/opt/cursor/artifacts"


NAVY = "12203C"
AC = "12203C"
GR = "137333"
RD = "C5221F"
YW = "B06000"
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_AC = PatternFill("solid", fgColor=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_ALT = PatternFill("solid", fgColor="F6F7F9")
FILL_LOTE = PatternFill("solid", fgColor="FEF7E0")
FILL_OK = PatternFill("solid", fgColor="E6F4EA")
FILL_BAD = PatternFill("solid", fgColor="FCE8E6")
FILL_SOFT = PatternFill("solid", fgColor="E8F0FE")
FONT_W = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
FONT_H = Font(name="Calibri", color="FFFFFF", bold=True, size=16)
FONT_T = Font(name="Calibri", bold=True, size=13, color=NAVY)
FONT_N = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="DADCE0"),
    right=Side(style="thin", color="DADCE0"),
    top=Side(style="thin", color="DADCE0"),
    bottom=Side(style="thin", color="DADCE0"),
)
WRAP = Alignment(wrap_text=True, vertical="center")


def fmt_n(x, dec=0) -> str:
    if x is None or x == "":
        return "--"
    try:
        n = float(x)
    except (TypeError, ValueError):
        return str(x)
    if dec == 0:
        return f"{n:,.0f}".replace(",", ".")
    s = f"{n:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_usd(x) -> str:
    if x is None:
        return "--"
    return f"US$ {fmt_n(x, 0)}"


def _kpis_plain(k: dict) -> dict:
    out = dict(k)
    for key, val in list(out.items()):
        if isinstance(val, date):
            out[key] = fmt_fecha(val)
    return out


def _pros_contras(sid: str, k: dict, sim: dict) -> Dict[str, List[str]]:
    noches = ", ".join(fmt_fecha(d) for d in sim["noches"]) or "ninguna"
    idle = sum((k.get("idle_dia") or {}).values())
    lin = k.get("pzas_por_linea") or {}
    l1 = lin.get("1") or lin.get(1) or 0
    pros, cons = [], []
    if sid == "0":
        pros.append("No hay costo de bono nocturno ni desgaste extra de plantilla.")
        pros.append("El tablero diurno y los pedidos especiales no se tocan.")
        cons.append(
            f"Ninguno de los 6 lotes nuevos entra a almacén al {fmt_fecha(CORTE_ALMACEN)} "
            "(cierran entre el 25/11 y el 18/12)."
        )
        cons.append(
            f"Solo {fmt_n(k['pzas_lotes_alm16_esc'])} pzas de lote nuevo se costuran a tiempo "
            f"para almacén 16/11 (corte de costura {fmt_fecha(CORTE_COSTURA_ALMACEN)})."
        )
        cons.append(
            "RIO/MAR lote nuevo CAB-DAMA-KIDS no cubren diciembre ni la tienda nueva a tiempo: "
            "el cierre más tarde es RIO LOTE NUEVO DAMA el 18/12 en almacén."
        )
        cons.append("Capacidad real (101–124 pzas/día, L5=40) sigue por debajo de las máquinas nuevas a 130.")
        return {"pros": pros, "contras": cons}

    pros.append(
        f"{k['n_noches']} noches ({noches}) producen {fmt_n(k['pzas_nocturno'])} pzas extra "
        f"a {fmt_usd(k['usd_por_pza'])} por pieza."
    )
    pros.append(
        f"+{fmt_n(k['extra_alm16'])} pzas adelantadas al corte de almacén 16/11 "
        f"({fmt_n(k['pzas_alm16_base'])} → {fmt_n(k['pzas_alm16_esc'])})."
    )
    pros.append(
        f"Lotes nuevos: {fmt_n(k['pzas_lotes_alm16_esc'])} pzas costuradas para 16/11 "
        f"(+{fmt_n(k['extra_alm16_lotes'])} vs base) y {fmt_n(k['pzas_lotes_dic_esc'])} "
        f"para demanda diciembre ({fmt_fecha(CORTE_COSTURA_DICIEMBRE)})."
    )
    if k.get("lotes_adelantados"):
        pros.append(
            f"{k['lotes_adelantados']} de 6 lotes nuevos cierran antes "
            f"(mediana {k['dias_ganados_lotes']} días de costura)."
        )
    rem_u = k.get("remanente_usado") or 0
    rem_t = k.get("remanente_total") or 0
    if rem_u > 0:
        pros.append(
            f"El cupo diurno que deja el nocturno se rellena con lo que sigue en cola: "
            f"{fmt_n(rem_u)} pzas de faltante no tablero entran al horizonte "
            f"(remanente total {fmt_n(rem_t)} pzas)."
        )
    if idle > 0:
        cons.append(
            f"Tras rellenar la cola quedan {fmt_n(idle)} pzas de ociosidad diurna "
            "(sin modelo asignado en esa línea, p.ej. MOTION LOOP o holgura de L5)."
        )
    if l1:
        pros.append(
            f"L1 de noche aporta {fmt_n(l1)} pzas a la cola de L2 (RIO), que es la que arrastra "
            "Kids urgente y los lotes nuevos RIO."
        )

    cons.append(
        f"Costo de bono: {fmt_usd(k['costo_usd'])} "
        f"({k['n_noches']} × {personas_total()} personas × {fmt_usd(BONO_USD)}), "
        "aunque alguna línea quede a media carga."
    )
    if k.get("noches_perdidas_cobro"):
        cons.append(
            f"Se pierde {k['noches_perdidas_cobro']} noche teórica por cobro "
            f"(30/09, último del mes). Quedan {k['n_noches']} de {k['noches_teoricas']}."
        )
    cons.append(
        "Ningún lote nuevo cierra entero en almacén al 16/11: el turno nocturno adelanta piezas, "
        "no mueve el corte de gerencia por sí solo."
    )
    cons.append(
        "L1 de noche sigue la cola de L2, no la cola propia: RIO LOTE NUEVO CAB y DAMA no adelantan "
        "su fecha de cierre (el cuello queda en el turno diurno de L1)."
    )
    cons.append(
        "No aplica a Por Hacer-Especial: Clásica / Mafe / satélite siguen en el día y pueden "
        "bloquear estaciones mientras la noche corre regular."
    )
    cons.append(
        "La cap nocturna es el 50% de la cap actual (máquinas nuevas aún no a régimen). "
        "Si la planta no sostiene ese 50%, el adelanto se encoge."
    )
    if sid in ("C", "D"):
        cons.append(
            "Tres o cuatro semanas de mar-mié-vie nocturno concentran fatiga, ausentismo y "
            "calidad en el tramo que alimenta noviembre-diciembre."
        )
    return {"pros": pros, "contras": cons}


def _pick_modelo(items: List[dict], nombre: str) -> dict:
    for m in items:
        if m.get("modelo") == nombre:
            return m
    return {}


def _lectura(payload: dict) -> str:
    k = payload["kpis"]
    kd, kc, kb, ka = k["D"], k["C"], k["B"], k["A"]
    md = payload["modelos"]["D"]
    m0 = payload["modelos"]["0"]
    cab = _pick_modelo(md, "MAR LOTE NUEVO CAB")
    cab0 = _pick_modelo(m0, "MAR LOTE NUEVO CAB")
    kids = _pick_modelo(md, "RIO LOTE NUEVO KIDS")
    kids0 = _pick_modelo(m0, "RIO LOTE NUEVO KIDS")
    rem = kd.get("remanente_usado") or 0
    return (
        f"Insertar los 6 lotes nuevos (~{fmt_n(payload['pzas_lotes'])} pzas) sin nocturnos deja "
        f"el almacén del {fmt_fecha(CORTE_ALMACEN)} cubierto solo con el plan actual: 0 de 6 lotes "
        f"nuevos llegan completos. Con D se costuran {fmt_n(kd['pzas_lotes_alm16_esc'])} pzas de esos "
        f"lotes a tiempo para el 16/11 (vs {fmt_n(k['0']['pzas_lotes_alm16_esc'])} sin noches) y "
        f"{fmt_n(kd['pzas_lotes_dic_esc'])} pzas para diciembre/tienda nueva. "
        f"MAR LOTE NUEVO CAB pasa de entrar el {cab0.get('entrada_base') or cab0.get('entrada_esc') or '--'} "
        f"al {cab.get('entrada_esc') or '--'} y RIO LOTE NUEVO KIDS del "
        f"{kids0.get('entrada_base') or kids0.get('entrada_esc') or '--'} al {kids.get('entrada_esc') or '--'}. "
        f"El cupo diurno que libera la noche se rellena con la cola (remanente absorbido: {fmt_n(rem)} pzas). "
        f"C logra parte del adelanto de D ({fmt_n(kc['extra_alm16'])} vs "
        f"{fmt_n(kd['extra_alm16'])} pzas) a menor costo ({fmt_usd(kc['costo_usd'])} vs {fmt_usd(kd['costo_usd'])}). "
        f"A ({fmt_usd(ka['costo_usd'])}) y B ({fmt_usd(kb['costo_usd'])}) sirven de piloto. "
        f"Recomendación: D si el criterio es tiendas + diciembre + tienda nueva; C si se quiere "
        f"limitar fatiga y caja. A no alcanza el objetivo de almacén."
    )


def construir_payload(plan_nuevo: dict, plan_actual: dict) -> dict:
    lotes = comparar_lotes(plan_actual, plan_nuevo)
    sims = {}
    kpis = {}
    modelos = {}
    for spec in ESCENARIOS:
        sim = simular_nocturnos(plan_nuevo, spec["semanas"])
        sims[spec["id"]] = sim
        kpis[spec["id"]] = _kpis_plain(kpis_escenario(sim))
        modelos[spec["id"]] = modelos_a_dict(sim["resultados"])
    pc = {spec["id"]: _pros_contras(spec["id"], kpis[spec["id"]], sims[spec["id"]]) for spec in ESCENARIOS}
    calendario = {sid: celdas_a_dict(sims[sid].get("calendario") or []) for sid in sims}
    base_idx = {}
    for c in calendario.get("0") or []:
        key = (c["fecha_iso"], c["linea"], c["turno"], c["modelo"])
        base_idx[key] = base_idx.get(key, 0.0) + c["qty"]
    for rows in calendario.values():
        for c in rows:
            c["qty_base"] = base_idx.get((c["fecha_iso"], c["linea"], c["turno"], c["modelo"]), 0.0)
            c["delta"] = round(c["qty"] - c["qty_base"], 1)
    resumen_cal = {sid: resumen_semanas(rows) for sid, rows in calendario.items()}
    payload = {
        "corte_almacen": fmt_fecha(CORTE_ALMACEN),
        "corte_costura": fmt_fecha(CORTE_COSTURA_ALMACEN),
        "corte_dic": fmt_fecha(CORTE_COSTURA_DICIEMBRE),
        "personas": PERSONAS,
        "personas_total": personas_total(),
        "bono": BONO_USD,
        "pzas_actual": round(sum(m.faltante for m in plan_actual["modelos"].values()), 0),
        "pzas_nuevo": round(sum(m.faltante for m in plan_nuevo["modelos"].values()), 0),
        "modelos_actual": len(plan_actual["modelos"]),
        "modelos_nuevo": len(plan_nuevo["modelos"]),
        "pzas_lotes": round(sum(m.faltante for m in lotes), 0),
        "lotes": [
            {
                "modelo": m.modelo,
                "faltante": m.faltante,
                "cap": m.cap,
                "lineas": m.lineas,
                "prioridad": m.prioridad,
                "termino_plan": fmt_fecha(m.termino_plan),
                "entrada_plan": fmt_fecha(m.entrada_plan),
                "fecha_obj": fmt_fecha(m.fecha_obj),
            }
            for m in lotes
        ],
        "kpis": kpis,
        "modelos": modelos,
        "noches": {spec["id"]: [fmt_fecha(d) for d in sims[spec["id"]]["noches"]] for spec in ESCENARIOS},
        "pzas_linea": {spec["id"]: sims[spec["id"]]["pzas_por_linea"] for spec in ESCENARIOS},
        "idle_dia": {spec["id"]: sims[spec["id"]]["idle_dia"] for spec in ESCENARIOS},
        "remanente": {
            spec["id"]: {
                "total": sims[spec["id"]].get("remanente_total") or 0,
                "usado": sims[spec["id"]].get("remanente_usado") or 0,
            }
            for spec in ESCENARIOS
        },
        "calendario": calendario,
        "resumen_cal": resumen_cal,
        "etiquetas_semana": etiquetas_semana(),
        "pros_contras": pc,
        "escenarios": ESCENARIOS,
    }
    payload["lectura"] = _lectura(payload)
    return payload


def _header(ws, title, subtitle, ncols=12):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(1, 1, title)
    c.font = FONT_H
    c.fill = FILL_NAVY
    c.alignment = Alignment(vertical="center")
    s = ws.cell(2, 1, subtitle)
    s.font = Font(name="Calibri", color="D9DEF0", size=10)
    s.fill = FILL_HEAD
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    for col in range(1, ncols + 1):
        ws.cell(1, col).fill = FILL_NAVY
        ws.cell(2, col).fill = FILL_HEAD


def _write_row(ws, r, values, fill=None, bold=False, wrap=False):
    for i, v in enumerate(values, 1):
        cell = ws.cell(r, i, v)
        cell.font = Font(name="Calibri", size=10, bold=bold, color=NAVY if bold else "222")
        cell.border = THIN
        cell.alignment = WRAP if wrap else Alignment(vertical="center")
        if fill:
            cell.fill = fill
    return r + 1


def _write_table(ws, r, headers, rows, lote_col=None):
    for i, h in enumerate(headers, 1):
        cell = ws.cell(r, i, h)
        cell.font = FONT_W
        cell.fill = FILL_AC
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[r].height = 28
    r += 1
    for idx, row in enumerate(rows):
        fill = FILL_ALT if idx % 2 == 0 else None
        if lote_col is not None and row and str(row[lote_col]).upper() in ("SI", "SÍ", "TRUE", "LOTE NUEVO"):
            fill = FILL_LOTE
        for i, v in enumerate(row, 1):
            cell = ws.cell(r, i, v)
            cell.font = FONT_N
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill
            if v in ("Sí", "A tiempo", "Cubre 16/11"):
                cell.fill = FILL_OK
            if v in ("No", "Retraso", "No cubre"):
                cell.fill = FILL_BAD
        r += 1
    return r


def _sort_impacto(items: List[dict]) -> List[dict]:
    """Mayor días ganados primero; sin dato al final."""
    return sorted(
        items,
        key=lambda m: (
            -(m["dias_ganados"] if m.get("dias_ganados") is not None else -10**9),
            m.get("modelo") or "",
        ),
    )


def _model_rows(items: List[dict], esc=True):
    rows = []
    for m in _sort_impacto(items):
        cubre = m["cubre_almacen_esc"] if esc else m["cubre_almacen_base"]
        rows.append(
            [
                m["modelo"],
                "Lote nuevo" if m["lote_nuevo"] else "",
                m["faltante"],
                m["termino_base"],
                m["termino_esc"] if esc else m["termino_base"],
                m["entrada_base"],
                m["entrada_esc"] if esc else m["entrada_base"],
                m["dias_ganados"] if esc else 0,
                m["pzas_nocturno"] if esc else 0,
                m["pzas_alm16_esc"] if esc else m["pzas_alm16_base"],
                m["pzas_dic_esc"] if esc else m["pzas_dic_base"],
                "Cubre 16/11" if cubre else ("No cubre" if cubre is False else "--"),
            ]
        )
    return rows


MODEL_HEADERS = [
    "Modelo",
    "Lote",
    "Faltante",
    "Salida costura base",
    "Salida costura escenario",
    "Entrada almacén base",
    "Entrada almacén escenario",
    "Días ganados",
    "Pzas nocturno",
    "Pzas p/ almacén 16/11",
    "Pzas p/ dic (costura 25/11)",
    "¿Cubre 16/11?",
]


def escribir_excel(payload: dict, path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    _header(
        ws,
        "Turnos nocturnos · decisión de planificación",
        "Escenarios A–D sobre el plan con lotes nuevos. Corte almacén 16/11/2026. Solo Por Hacer (no Especial).",
        12,
    )
    r = 4
    ws.cell(r, 1, "Lectura para gerencia").font = FONT_T
    r = 5
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=12)
    cell = ws.cell(r, 1, payload["lectura"])
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.fill = FILL_SOFT
    cell.font = Font(name="Calibri", size=11)
    r = 10
    k = payload["kpis"]
    cmp_headers = [
        "Escenario",
        "Semanas",
        "Noches reales",
        "Noches teóricas",
        "Costo bono US$",
        "Pzas nocturno",
        "US$/pza",
        "Extra p/ 16/11",
        "Pzas lote nuevo p/ 16/11",
        "Pzas lote nuevo p/ dic",
        "Lotes nuevos que cierran antes",
        "Mediana días ganados (lotes)",
        "Modelos adelantados",
    ]
    cmp_rows = []
    tit = {e["id"]: e["titulo"] for e in payload["escenarios"]}
    for sid in "0ABCD":
        x = k[sid]
        cmp_rows.append(
            [
                sid if sid != "0" else "Base",
                tit[sid],
                x["n_noches"],
                x["noches_teoricas"],
                x["costo_usd"],
                x["pzas_nocturno"],
                x["usd_por_pza"] if x["usd_por_pza"] is not None else 0,
                x["extra_alm16"],
                x["pzas_lotes_alm16_esc"],
                x["pzas_lotes_dic_esc"],
                x["lotes_adelantados"],
                x["dias_ganados_lotes"] if x["dias_ganados_lotes"] is not None else 0,
                x["modelos_adelantados"],
            ]
        )
    r = _write_table(ws, r, cmp_headers, cmp_rows)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Piezas extra vs costo"
    chart.y_axis.title = "Pzas extra 16/11"
    data = Reference(ws, min_col=8, min_row=10, max_row=15)
    cats = Reference(ws, min_col=1, min_row=11, max_row=15)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.y_axis.axId = 100
    chart2 = BarChart()
    chart2.type = "col"
    chart2.y_axis.axId = 200
    chart2.y_axis.title = "US$"
    data2 = Reference(ws, min_col=5, min_row=10, max_row=15)
    chart2.add_data(data2, titles_from_data=True)
    chart2.y_axis.crosses = "max"
    chart += chart2
    chart.width = 18
    chart.height = 8
    ws.add_chart(chart, "A16")

    r = 32
    ws.cell(r, 1, "Lotes nuevos (no están en el plan ACTUAL)").font = FONT_T
    r = 33
    lote_headers = ["Modelo", "Faltante", "Cap/día", "Líneas", "Prioridad", "Salida costura plan", "Entrada almacén plan", "Fecha obj.", "¿Cubre 16/11 hoy?"]
    lote_rows = []
    for m in payload["lotes"]:
        cubre = "No cubre"
        lote_rows.append(
            [m["modelo"], m["faltante"], m["cap"], m["lineas"], m["prioridad"], m["termino_plan"], m["entrada_plan"], m["fecha_obj"], cubre]
        )
    r = _write_table(ws, r, lote_headers, lote_rows)
    r += 1
    ws.cell(r, 1, "Contraste plan ACTUAL vs plan LOTE NUEVO").font = FONT_T
    r += 1
    r = _write_table(
        ws,
        r,
        ["Plan", "Modelos Por Hacer", "Faltante (pzas)", "Lotes nuevos", "Pzas lotes nuevos"],
        [
            ["ACTUAL (sin lotes nuevos)", payload["modelos_actual"], payload["pzas_actual"], 0, 0],
            ["LOTE NUEVO", payload["modelos_nuevo"], payload["pzas_nuevo"], len(payload["lotes"]), payload["pzas_lotes"]],
            ["Delta", payload["modelos_nuevo"] - payload["modelos_actual"], payload["pzas_nuevo"] - payload["pzas_actual"], len(payload["lotes"]), payload["pzas_lotes"]],
        ],
    )
    r += 1
    ws.cell(r, 1, "Costo de bono (todas las líneas, 22 personas)").font = FONT_T
    r += 1
    r = _write_table(
        ws,
        r,
        ["Línea", "Personas", "US$ / noche / línea"],
        [[f"Línea {lin}", n, n * BONO_USD] for lin, n in PERSONAS.items()]
        + [["Total", payload["personas_total"], payload["personas_total"] * BONO_USD]],
    )
    r += 1
    ws.cell(r, 1, "Riesgo si NO se hacen nocturnos").font = FONT_T
    r += 1
    riesgos = [
        "Los 6 lotes nuevos (~8.912 pzas) no están en el plan ACTUAL; al meterlos, almacén 16/11 queda corto para tiendas + diciembre + tienda nueva.",
        "Sin noches, MAR/RIO lote nuevo CAB-DAMA-KIDS entran a almacén entre el 25/11 y el 18/12.",
        "RIO KIDS urgente y shorts de L5 ya van retrasados en el plan diurno; la noche de L1 refuerza la cola de L2 (RIO).",
        "La cap actual (101–124, L5=40) no es la de máquinas nuevas a 130: el atraso se agranda si no se adelanta.",
        "LITE PANT DAMA y Basic Line Pant quedan a fines de diciembre y no entran al corte de almacén.",
    ]
    for t in riesgos:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        c = ws.cell(r, 1, "• " + t)
        c.alignment = WRAP
        c.fill = FILL_BAD
        ws.row_dimensions[r].height = 32
        r += 1

    for col in range(1, 14):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.column_dimensions["A"].width = 36
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    for spec in payload["escenarios"]:
        if spec["id"] == "0":
            continue
        sid = spec["id"]
        wsx = wb.create_sheet(f"Escenario {sid}")
        noches = payload["noches"][sid]
        sub = (
            f"{spec['titulo']} · {k[sid]['n_noches']} noches ({', '.join(noches)}) · "
            f"costo {fmt_usd(k[sid]['costo_usd'])} · {fmt_n(k[sid]['pzas_nocturno'])} pzas nocturno"
        )
        _header(wsx, f"Escenario {sid} — {spec['titulo']}", sub, 15)
        r = 4
        kpi_h = ["Noches", "Costo US$", "Pzas nocturno", "US$/pza", "Extra 16/11", "Lote nuevo 16/11", "Lote nuevo dic", "Días ganados (lotes)", "Remanente en cola", "Ociosidad diurna"]
        idle = sum((payload["idle_dia"].get(sid) or {}).values())
        rem_u = (payload.get("remanente") or {}).get(sid, {}).get("usado") or k[sid].get("remanente_usado") or 0
        kpi_v = [
            k[sid]["n_noches"],
            k[sid]["costo_usd"],
            k[sid]["pzas_nocturno"],
            k[sid]["usd_por_pza"],
            k[sid]["extra_alm16"],
            k[sid]["pzas_lotes_alm16_esc"],
            k[sid]["pzas_lotes_dic_esc"],
            k[sid]["dias_ganados_lotes"],
            rem_u,
            idle,
        ]
        r = _write_table(wsx, r, kpi_h, [kpi_v])
        r += 1
        wsx.cell(r, 1, "A favor").font = Font(name="Calibri", bold=True, color=GR, size=12)
        r += 1
        for p in payload["pros_contras"][sid]["pros"]:
            wsx.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
            c = wsx.cell(r, 1, "▸ " + p)
            c.fill = FILL_OK
            c.alignment = WRAP
            wsx.row_dimensions[r].height = 34
            r += 1
        r += 1
        wsx.cell(r, 1, "En contra").font = Font(name="Calibri", bold=True, color=RD, size=12)
        r += 1
        for p in payload["pros_contras"][sid]["contras"]:
            wsx.merge_cells(start_row=r, start_column=1, end_row=r, end_column=15)
            c = wsx.cell(r, 1, "▸ " + p)
            c.fill = FILL_BAD
            c.alignment = WRAP
            wsx.row_dimensions[r].height = 34
            r += 1
        r += 1
        wsx.cell(r, 1, "Todos los modelos Por Hacer · orden por días ganados").font = FONT_T
        r += 1
        mods = list(payload["modelos"][sid])
        r = _write_table(wsx, r, MODEL_HEADERS, _model_rows(mods), lote_col=1)
        r += 1
        wsx.cell(r, 1, "Lotes nuevos · orden por días ganados").font = FONT_T
        r += 1
        lotes_m = [m for m in payload["modelos"][sid] if m["lote_nuevo"]]
        r = _write_table(wsx, r, MODEL_HEADERS, _model_rows(lotes_m), lote_col=1)
        r += 1
        wsx.cell(r, 1, "Piezas nocturnas por línea").font = FONT_T
        r += 1
        lin = payload["pzas_linea"].get(sid) or {}
        r = _write_table(
            wsx,
            r,
            ["Línea", "Personas", "Pzas nocturno", "Ociosidad diurna"],
            [
                [
                    f"Línea {x}",
                    PERSONAS[x],
                    lin.get(x, 0),
                    (payload["idle_dia"].get(sid) or {}).get(x, 0),
                ]
                for x in "12345"
            ],
        )
        for col in range(1, 13):
            wsx.column_dimensions[get_column_letter(col)].width = 18
        wsx.column_dimensions["A"].width = 38
        wsx.freeze_panes = "A4"
        wsx.page_setup.orientation = "landscape"
        wsx.page_setup.fitToPage = True
        wsx.page_setup.fitToWidth = 1
        wsx.page_setup.fitToHeight = 0
        wsx.sheet_properties.pageSetUpPr.fitToPage = True

    wsc = wb.create_sheet("Calendario")
    _header(
        wsc,
        "Calendario de planificación · escenarios A–D",
        "Filtrar por escenario y semana. Delta = piezas del escenario menos la base sin nocturnos. Remanente = faltante que no estaba en el tablero y ahora ocupa cupo diurno.",
        12,
    )
    cal_headers = [
        "Escenario",
        "Semana",
        "Fecha",
        "Día",
        "Línea",
        "Turno",
        "Modelo",
        "Pzas escenario",
        "Pzas base",
        "Delta",
        "¿Remanente?",
        "¿Cambio?",
    ]
    dias_nom = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    cal_rows = []
    for spec in payload["escenarios"]:
        sid = spec["id"]
        label = "Base" if sid == "0" else sid
        for c in payload.get("calendario", {}).get(sid) or []:
            delta = c.get("delta")
            cal_rows.append(
                [
                    label,
                    c.get("semana"),
                    c.get("fecha"),
                    dias_nom[c.get("dow") or 0],
                    f"Línea {c.get('linea')}",
                    "Noche" if c.get("turno") == "noche" else "Día",
                    c.get("modelo"),
                    c.get("qty"),
                    c.get("qty_base"),
                    delta,
                    "Sí" if c.get("remanente") else "",
                    "Sí" if abs(delta or 0) > 0.001 or c.get("turno") == "noche" else "",
                ]
            )
    r = _write_table(wsc, 4, cal_headers, cal_rows)
    if cal_rows:
        wsc.auto_filter.ref = f"A4:L{3 + len(cal_rows)}"
    r += 2
    wsc.cell(r, 1, "Resumen por escenario y semana").font = FONT_T
    r += 1
    sum_headers = ["Escenario", "Semana", "Pzas día", "Pzas noche", "Pzas remanente", "Pzas total", "Modelos"]
    sum_rows = []
    for spec in payload["escenarios"]:
        sid = spec["id"]
        label = "Base" if sid == "0" else sid
        for w in payload.get("resumen_cal", {}).get(sid) or []:
            sum_rows.append(
                [
                    label,
                    w["semana"],
                    w["pzas_dia"],
                    w["pzas_noche"],
                    w["pzas_remanente"],
                    w["pzas"],
                    w["n_modelos"],
                ]
            )
    r = _write_table(wsc, r, sum_headers, sum_rows)
    for col in range(1, 13):
        wsc.column_dimensions[get_column_letter(col)].width = 18
    wsc.column_dimensions["G"].width = 36
    wsc.freeze_panes = "A5"
    wsc.page_setup.orientation = "landscape"
    wsc.page_setup.fitToPage = True
    wsc.page_setup.fitToWidth = 1
    wsc.page_setup.fitToHeight = 0
    wsc.sheet_properties.pageSetUpPr.fitToPage = True

    wss = wb.create_sheet("Supuestos")
    _header(wss, "Supuestos del modelo", "Qué entra, qué no, y cómo se calculan fechas y costo.", 8)
    rows = [
        ["Calendario", "Primer nocturno martes 29/09/2026. Semana arranca lunes 28/09."],
        ["Días de noche", "Martes, miércoles y viernes. 3 noches teóricas por semana."],
        ["Cobro", "Sin nocturno el día 15 ni el último de mes. El 30/09 (miércoles) se salta: el escenario A tiene 2 noches, no 3."],
        ["Alcance", "Cinco líneas. Solo modelos de Por Hacer. Por Hacer – Especial no se toca de noche."],
        ["Línea 1", "De noche tiene 100% del cupo nocturno, pero produce la cola de Línea 2 (RIO), no la cola diurna de L1."],
        ["Capacidad noche", "50% de la cap diurna del modelo que está al frente de esa línea. Cap según Por Hacer (máquinas nuevas aún no a 130)."],
        ["L5", "Cap diurna 40 → noche 20. El resto L2–L4 usa 101 o 124 según modelo."],
        ["Cómo se simula", "Se respeta el orden cronológico del tablero 12 semanas. Cada noche consume de la cola regular; los días siguientes fabrican lo que queda, así se adelantan los lotes de atrás."],
        ["Cupos diurnos libres", "El cupo diurno que deja el nocturno no se deja ocioso: se rellena con lo que sigue en cola (siguiente modelo del tablero y el faltante no programado de esa línea)."],
        ["Almacén", "Entrada = salida de costura + 4 días hábiles. El corte 16/11 (lunes) exige terminar costura el 10/11."],
        ["Diciembre / tienda nueva", "Se cuenta costura al 25/11 (entra a almacén ~01/12) como proxy de reposición diciembre y tienda nueva."],
        ["Costo", "US$ 15 por persona por noche. L1–L3 = 4; L4–L5 = 5; total 22. Se paga el equipo completo aunque una línea no llene el cupo."],
        ["Plan ACTUAL vs NUEVO", "ACTUAL = planificación vigente. NUEVO = misma base + 6 lotes MAR/RIO LOTE NUEVO CAB-DAMA-KIDS (~8.912 pzas) que las proyecciones de tienda van a pedir."],
        ["Qué no se mueve", "Pedidos especiales, satélite, días de cobro, sábados y domingos. No se reordena prioridad: se adelanta la cola tal cual está."],
        ["Cuello L1", "Como L1 noche trabaja para L2, RIO LOTE NUEVO CAB y DAMA no adelantan su cierre: el remanente sigue en el día de L1."],
        ["Piezas no programadas", "MOTION LOOP (sin línea) no entra. El resto de LITE PANT, Basic Line Pant y SEMI MOTION se encolan al final de su línea y ocupan el cupo diurno que libera el nocturno."],
        ["Calendario", "La hoja Calendario y la pestaña HTML filtran por escenario y semana. Cada celda compara el plan del escenario contra la base (sin noches). Amarillo = remanente; azul = cambio vs base; navy = turno noche."],
        ["Conservadurismo", "Floor 50% de cap actual, no de 130. Si las máquinas nuevas llegan a régimen, el mismo esquema de noches rinde más."],
        ["Listas", "Las tablas de modelos van por días ganados (mayor impacto primero), no por calendario ni alfabético."],
        ["Fuente", "Planificacion_Produccion_ACTUAL.xlsx y Planificacion_Produccion_LOTE_NUEVO.xlsx (Por Hacer, tablero Semana 1–12, Proyeccion, Entrada de Almacen)."],
    ]
    r = 4
    r = _write_table(wss, r, ["Tema", "Supuesto"], rows)
    wss.column_dimensions["A"].width = 28
    wss.column_dimensions["B"].width = 110
    for i in range(5, 5 + len(rows)):
        wss.row_dimensions[i].height = 36
    wss.page_setup.orientation = "landscape"
    wss.page_setup.fitToPage = True
    wss.page_setup.fitToWidth = 1
    wss.page_setup.fitToHeight = 0
    wss.sheet_properties.pageSetUpPr.fitToPage = True

    for name in wb.sheetnames:
        sh = wb[name]
        sh.sheet_view.showGridLines = False
        sh.sheet_properties.tabColor = "12203C"
    wb.save(path)


def _hx(s) -> str:
    return htmlmod.escape("" if s is None else str(s))


def _cubre_html(v) -> str:
    if v is True:
        return '<span class="ok">Cubre 16/11</span>'
    if v is False:
        return '<span class="no">No cubre</span>'
    return "--"


def _tabla_modelos(items: List[dict]) -> str:
    rows = _sort_impacto(items)
    out = [
        '<div class="wrap"><table><thead><tr>',
        "<th>Modelo</th><th></th><th>Faltante</th>",
        "<th>Salida base</th><th>Salida esc.</th><th>Almacén base</th><th>Almacén esc.</th>",
        "<th>Días ganados</th><th>Pzas noche</th><th>Pzas 16/11</th><th>Pzas dic</th><th>16/11</th>",
        "</tr></thead><tbody>",
    ]
    for m in rows:
        lote = '<span class="badge">Lote nuevo</span>' if m.get("lote_nuevo") else ""
        cls = ' class="lote"' if m.get("lote_nuevo") else ""
        out.append(
            f"<tr{cls}><td>{_hx(m['modelo'])}</td><td>{lote}</td>"
            f"<td>{fmt_n(m.get('faltante'))}</td>"
            f"<td>{_hx(m.get('termino_base'))}</td><td>{_hx(m.get('termino_esc'))}</td>"
            f"<td>{_hx(m.get('entrada_base'))}</td><td>{_hx(m.get('entrada_esc'))}</td>"
            f"<td>{_hx(m.get('dias_ganados') if m.get('dias_ganados') is not None else '--')}</td>"
            f"<td>{fmt_n(m.get('pzas_nocturno'))}</td><td>{fmt_n(m.get('pzas_alm16_esc'))}</td>"
            f"<td>{fmt_n(m.get('pzas_dic_esc'))}</td><td>{_cubre_html(m.get('cubre_almacen_esc'))}</td></tr>"
        )
    out.append("</tbody></table></div>")
    return "\n".join(out)


def _kpis_html(sid: str, payload: dict) -> str:
    x = payload["kpis"][sid]
    idle = sum((payload["idle_dia"].get(sid) or {}).values())
    usd_p = f"US$ {fmt_n(x['usd_por_pza'], 2)}" if x.get("usd_por_pza") is not None else "--"
    return f"""
    <div class="g g5">
      <div class="kpi"><span>Noches</span><b>{x['n_noches']}</b></div>
      <div class="kpi"><span>Costo bono</span><b>{fmt_usd(x['costo_usd'])}</b></div>
      <div class="kpi"><span>Pzas nocturno</span><b>{fmt_n(x['pzas_nocturno'])}</b></div>
      <div class="kpi"><span>Extra almacén 16/11</span><b>{fmt_n(x['extra_alm16'])}</b></div>
      <div class="kpi"><span>Lote nuevo p/ dic</span><b>{fmt_n(x['pzas_lotes_dic_esc'])}</b></div>
    </div>
    <div class="g g4">
      <div class="kpi"><span>Lote nuevo p/ 16/11</span><b>{fmt_n(x['pzas_lotes_alm16_esc'])}</b></div>
      <div class="kpi"><span>US$ / pza</span><b>{usd_p}</b></div>
      <div class="kpi"><span>Lotes que cierran antes</span><b>{x['lotes_adelantados']} / 6</b></div>
      <div class="kpi"><span>Remanente absorbido</span><b>{fmt_n(x.get('remanente_usado') or 0)}</b></div>
    </div>"""


def _lis(items: List[str], cls: str) -> str:
    return "".join(f'<div class="{cls}">{_hx(t)}</div>' for t in items)


def _sec_esc(sid: str, payload: dict) -> str:
    spec = next(e for e in payload["escenarios"] if e["id"] == sid)
    pc = payload["pros_contras"][sid]
    noches = ", ".join(payload["noches"][sid]) or "—"
    lotes = [m for m in payload["modelos"][sid] if m.get("lote_nuevo")]
    lin = payload["pzas_linea"].get(sid) or {}
    idle = payload["idle_dia"].get(sid) or {}
    lin_rows = "".join(
        f"<tr><td>Línea {x}</td><td>{PERSONAS[x]}</td><td>{fmt_n(lin.get(x, 0))}</td>"
        f"<td>{fmt_n(idle.get(x, 0))}</td></tr>"
        for x in "12345"
    )
    return f"""
    <div class="card"><h3>{_hx(spec['titulo'])} · noches { _hx(noches)}</h3>
      {_kpis_html(sid, payload)}</div>
    <div class="g g2">
      <div class="card"><h3>A favor</h3>{_lis(pc['pros'], 'pro')}</div>
      <div class="card"><h3>En contra</h3>{_lis(pc['contras'], 'con')}</div>
    </div>
    <div class="card"><h3>Todos los modelos Por Hacer (orden por días ganados)</h3>
      {_tabla_modelos(payload['modelos'][sid])}</div>
    <div class="card"><h3>Lotes nuevos (orden por días ganados)</h3>{_tabla_modelos(lotes)}</div>
    <div class="card"><h3>Aprovechamiento por línea</h3>
      <table><thead><tr><th>Línea</th><th>Personas</th><th>Pzas nocturno</th><th>Ociosidad diurna</th></tr></thead>
      <tbody>{lin_rows}</tbody></table>
    </div>
    """


def escribir_html(payload: dict, path: str):
    k = payload["kpis"]
    cmp_rows = []
    for sid in "0ABCD":
        spec = next(e for e in payload["escenarios"] if e["id"] == sid)
        x = k[sid]
        label = "Base" if sid == "0" else sid
        cmp_rows.append(
            "<tr>"
            f"<td>{label}</td><td>{_hx(spec['titulo'])}</td>"
            f"<td>{x['n_noches']}</td><td>{fmt_usd(x['costo_usd'])}</td>"
            f"<td>{fmt_n(x['pzas_nocturno'])}</td><td>{fmt_n(x['extra_alm16'])}</td>"
            f"<td>{fmt_n(x['pzas_lotes_alm16_esc'])}</td><td>{fmt_n(x['pzas_lotes_dic_esc'])}</td>"
            f"<td>{x['lotes_adelantados']}</td><td>{x['dias_ganados_lotes'] if x['dias_ganados_lotes'] is not None else '--'}</td>"
            "</tr>"
        )
    lote_rows = []
    for m in payload["lotes"]:
        lote_rows.append(
            '<tr class="lote">'
            f"<td>{_hx(m['modelo'])}</td><td>{fmt_n(m['faltante'])}</td><td>{_hx(m['cap'])}</td>"
            f"<td>{_hx(m['lineas'])}</td><td>{_hx(m['prioridad'])}</td>"
            f"<td>{_hx(m['termino_plan'])}</td><td>{_hx(m['entrada_plan'])}</td>"
            f"<td>{_hx(m['fecha_obj'])}</td><td><span class=\"no\">No cubre</span></td></tr>"
        )
    extras = [k[s]["extra_alm16"] for s in "0ABCD"]
    costos = [k[s]["costo_usd"] for s in "0ABCD"]
    secs = "\n".join(
        f'<section class="sec" id="sec-{sid}">{_sec_esc(sid, payload)}</section>'
        for sid in "ABCD"
    )
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Turnos nocturnos · decisión de planificación</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root {{
  --azul: #1a56db; --azul-claro: #e8f0fe; --azul-borde: #a8c7fa;
  --verde: #137333; --verde-bg: #e6f4ea; --verde-borde: #a8dab5;
  --rojo: #c5221f; --rojo-bg: #fce8e6; --rojo-borde: #f5b5b3;
  --ambar: #b06000; --ambar-bg: #fef7e0;
  --gris: #5f6368; --borde: #dadce0; --texto: #202124; --navy: #12203c;
}}
* {{ box-sizing: border-box; }}
body {{
  margin: 0; background: #f6f7f9; color: var(--texto);
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
  font-size: 15px; line-height: 1.5;
}}
header.top {{ background: var(--navy); color: #fff; padding: 22px 0 0; }}
.wrap {{ max-width: 1280px; margin: 0 auto; padding: 0 20px; }}
header.top h1 {{ margin: 0 0 4px; font-size: 24px; font-weight: 700; }}
header.top h1 em {{ font-style: normal; color: #a8c7fa; }}
header.top .sub {{ color: #b9c4d8; font-size: 14px; margin-bottom: 16px; }}
.chips {{ display: flex; gap: 10px; flex-wrap: wrap; margin-bottom: 18px; }}
.chip {{
  background: rgba(255,255,255,.1); border: 1px solid rgba(255,255,255,.2);
  border-radius: 999px; padding: 5px 14px; font-size: 13px; color: #e8eaed;
}}
.chip b {{ color: #fff; }}
nav.tabs {{ display: flex; gap: 2px; flex-wrap: wrap; }}
nav.tabs button {{
  background: rgba(255,255,255,.08); color: #cdd6e6; border: 0; cursor: pointer;
  padding: 11px 22px; font-size: 14px; font-weight: 600; font-family: inherit;
  border-radius: 8px 8px 0 0;
}}
nav.tabs button:hover {{ background: rgba(255,255,255,.16); color: #fff; }}
nav.tabs button.on {{ background: #f6f7f9; color: var(--navy); }}
.content {{ max-width: 1280px; margin: 0 auto; padding: 24px 20px 60px; }}
.sec {{ display: none; }}
.sec.on {{ display: block; }}
.g {{ display: grid; gap: 14px; margin-bottom: 14px; }}
.g5 {{ grid-template-columns: repeat(5, 1fr); }}
.g4 {{ grid-template-columns: repeat(4, 1fr); }}
.g2 {{ grid-template-columns: 1fr 1fr; }}
.kpi {{ background: #fff; border: 1px solid var(--borde); border-radius: 12px; padding: 14px 16px; }}
.kpi b {{ display: block; font-size: 25px; font-weight: 800; margin-top: 4px; line-height: 1.1; color: var(--navy); }}
.kpi span {{ font-size: 12px; color: var(--gris); text-transform: uppercase; letter-spacing: .4px; font-weight: 600; }}
.card {{ background: #fff; border: 1px solid var(--borde); border-radius: 12px; padding: 18px 20px; margin-bottom: 14px; }}
.card h3 {{ font-size: 15px; margin: 0 0 10px; color: var(--gris); text-transform: uppercase; letter-spacing: .4px; }}
.note {{ background: var(--azul-claro); border: 1px solid var(--azul-borde); border-radius: 12px; padding: 16px 20px; color: var(--texto); line-height: 1.45; font-size: 14.5px; }}
.pro, .con {{ padding: 8px 12px; border-radius: 8px; margin: 6px 0; font-size: 14px; line-height: 1.4; }}
.pro {{ background: var(--verde-bg); border-left: 4px solid var(--verde); }}
.con {{ background: var(--rojo-bg); border-left: 4px solid var(--rojo); }}
.tablebox {{ border: 1px solid var(--borde); border-radius: 12px; overflow: hidden; }}
table {{ width: 100%; border-collapse: collapse; background: #fff; font-size: 14px; }}
th {{
  background: var(--navy); color: #fff; text-align: left; padding: 10px 12px;
  font-size: 12.5px; font-weight: 600; text-transform: uppercase; letter-spacing: .3px;
  position: sticky; top: 0;
}}
td {{ padding: 9px 12px; border-bottom: 1px solid var(--borde); vertical-align: top; }}
tr.lote td {{ background: var(--ambar-bg); }}
.ok {{ color: var(--verde); font-weight: 700; }}
.no {{ color: var(--rojo); font-weight: 700; }}
.badge {{ display: inline-block; padding: 2px 9px; border-radius: 999px; font-size: 12px; font-weight: 700; background: var(--ambar-bg); color: var(--ambar); }}
.wrap {{ overflow: auto; max-height: 520px; }}
.cw {{ height: 280px; position: relative; background: #fff; }}
.foot {{ text-align: center; color: var(--gris); font-size: 12.5px; padding: 18px 20px 28px; }}
.cal-tools {{ display:flex; gap:10px; flex-wrap:wrap; align-items:center; margin-bottom:12px; }}
.cal-tools .lab {{ font-size:12px; font-weight:700; color:var(--gris); text-transform:uppercase; letter-spacing:.3px; }}
.cal-tools button {{
  background:#fff; border:1px solid var(--borde); border-radius:999px; padding:6px 12px;
  font-size:13px; font-weight:600; cursor:pointer; font-family:inherit; color:var(--navy);
}}
.cal-tools button.on {{ background:var(--navy); color:#fff; border-color:var(--navy); }}
.cal-week {{ display:flex; gap:6px; flex-wrap:wrap; }}
table.cal {{ font-size:12.5px; }}
table.cal th {{ text-align:center; }}
table.cal td {{ min-width:110px; vertical-align:top; }}
.cell-mod {{ font-weight:700; color:var(--navy); }}
.cell-qty {{ color:var(--gris); }}
.chip-n {{ display:block; margin-top:6px; padding:4px 6px; border-radius:6px; background:var(--navy); color:#fff; font-size:11px; }}
.chip-r {{ display:inline-block; margin-left:4px; padding:1px 6px; border-radius:999px; background:var(--ambar-bg); color:var(--ambar); font-size:10px; font-weight:700; }}
td.chg {{ background:#e8f0fe; }}
td.rem {{ background:var(--ambar-bg); }}
.legend span {{ display:inline-block; margin-right:12px; font-size:12.5px; color:var(--gris); }}
.sw {{ display:inline-block; width:12px; height:12px; border-radius:3px; margin-right:4px; vertical-align:middle; border:1px solid var(--borde); }}
@media (max-width: 900px) {{ .g5, .g4, .g2 {{ grid-template-columns: 1fr 1fr; }} }}
</style>
</head>
<body>
<header class="top">
  <div class="wrap">
    <h1>Turnos nocturnos · decisión de planificación</h1>
    <div class="sub">Comparación contra el plan con lotes nuevos, sin nocturno · corte almacén 16/11/2026</div>
    <div class="chips">
      <span class="chip">Mar / mié / vie</span>
      <span class="chip">5 líneas · <b>22 personas</b></span>
      <span class="chip">Bono <b>US$ 15</b> / persona / noche</span>
      <span class="chip">Solo Por Hacer</span>
    </div>
    <nav class="tabs">
      <button class="tab on" data-id="res">Resumen</button>
      <button class="tab" data-id="A">Escenario A</button>
      <button class="tab" data-id="B">Escenario B</button>
      <button class="tab" data-id="C">Escenario C</button>
      <button class="tab" data-id="D">Escenario D</button>
      <button class="tab" data-id="cal">Calendario</button>
      <button class="tab" data-id="sup">Supuestos</button>
    </nav>
  </div>
</header>
<div class="content">
  <section class="sec on" id="sec-res">
    <div class="card"><h3>Lectura para gerencia</h3><div class="note">{_hx(payload['lectura'])}</div></div>
    <div class="g g4">
      <div class="kpi"><span>Plan ACTUAL</span><b>{fmt_n(payload['pzas_actual'])} pzas</b></div>
      <div class="kpi"><span>Plan LOTE NUEVO</span><b>{fmt_n(payload['pzas_nuevo'])} pzas</b></div>
      <div class="kpi"><span>Lotes a insertar</span><b>{fmt_n(payload['pzas_lotes'])} pzas</b></div>
      <div class="kpi"><span>Corte almacén</span><b>{_hx(payload['corte_almacen'])}</b></div>
    </div>
    <div class="card"><h3>Comparación de escenarios</h3>
      <div class="cw"><canvas id="ch1"></canvas></div>
      <div class="wrap"><table><thead><tr>
        <th>Esc.</th><th>Descripción</th><th>Noches</th><th>Costo</th><th>Pzas noche</th>
        <th>Extra 16/11</th><th>Lote nuevo 16/11</th><th>Lote nuevo dic</th><th>Lotes adelantados</th><th>Mediana días</th>
      </tr></thead><tbody>{''.join(cmp_rows)}</tbody></table></div>
    </div>
    <div class="card"><h3>Lotes nuevos que hoy no están en el plan ACTUAL</h3>
      <p style="color:var(--gris);font-size:13.5px;margin-bottom:8px">Proyección de tienda / diciembre / tienda nueva. Orden de salida de costura del plan sin nocturnos.</p>
      <div class="wrap"><table><thead><tr><th>Modelo</th><th>Faltante</th><th>Cap/día</th><th>Líneas</th><th>Prioridad</th><th>Salida plan</th><th>Almacén plan</th><th>Obj.</th><th>16/11</th></tr></thead>
      <tbody>{''.join(lote_rows)}</tbody></table></div>
    </div>
    <div class="g g2">
      <div class="card"><h3>Riesgo si no hay nocturnos</h3>
        {_lis(payload['pros_contras']['0']['contras'], 'con')}
      </div>
      <div class="card"><h3>Qué gana el nocturno (todas las líneas)</h3>
        {_lis(payload['pros_contras']['D']['pros'][:4], 'pro')}
        <p style="color:var(--gris);font-size:13px;margin-top:8px">L1 noche = cola de L2. Especiales no se tocan. Cobro 30/09 sin noche.</p>
      </div>
    </div>
  </section>
  {secs}
  <section class="sec" id="sec-cal">
    <div class="card">
      <h3>Calendario de producción · filtro por escenario y semana</h3>
      <p style="color:var(--gris);font-size:13.5px;margin:0 0 10px">
        El plan diurno se recorre en cola: lo que la noche adelanta deja cupo al día siguiente,
        y ese cupo lo toma el siguiente modelo (tablero o remanente no programado).
        Azul = distinto a la base sin nocturnos. Amarillo = remanente. Navy = turno noche.
      </p>
      <div class="cal-tools">
        <span class="lab">Escenario</span>
        <span id="cal-esc"></span>
      </div>
      <div class="cal-tools">
        <span class="lab">Semana</span>
        <span class="cal-week" id="cal-week"></span>
      </div>
      <div class="legend">
        <span><i class="sw" style="background:#fff"></i>Igual a la base</span>
        <span><i class="sw" style="background:#e8f0fe"></i>Modificado vs base</span>
        <span><i class="sw" style="background:#fef7e0"></i>Remanente en cola</span>
        <span><i class="sw" style="background:#12203c"></i>Turno noche</span>
      </div>
    </div>
    <div class="g g4" id="cal-kpis"></div>
    <div class="card">
      <h3 id="cal-title">Semana</h3>
      <div class="wrap" id="cal-grid"></div>
    </div>
  </section>
  <section class="sec" id="sec-sup">
    <div class="card"><h3>Supuestos</h3>
      <table><tbody>
        <tr><td>Calendario</td><td>Primer nocturno martes 29/09/2026. Semana del lunes 28/09.</td></tr>
        <tr><td>Días</td><td>Martes, miércoles y viernes. Sin noche el 15 ni el último de mes (30/09 se salta).</td></tr>
        <tr><td>Líneas</td><td>5 líneas. L1–L3 = 4 personas, L4–L5 = 5. Bono US$ 15 / persona / noche.</td></tr>
        <tr><td>L1 noche</td><td>100% del cupo nocturno sobre la cola de L2 (RIO), no sobre la cola diurna de L1.</td></tr>
        <tr><td>Capacidad</td><td>50% de la cap diurna actual del modelo al frente. Máquinas nuevas aún no a 130 pzas/día.</td></tr>
        <tr><td>Alcance</td><td>Solo Por Hacer. Por Hacer – Especial queda en el turno diurno.</td></tr>
        <tr><td>Almacén</td><td>Salida costura + 4 días hábiles. 16/11 exige costura el 10/11. Diciembre se mide con costura al 25/11.</td></tr>
        <tr><td>Simulación</td><td>Cola cronológica del tablero 12 semanas. La noche consume el frente; los días siguientes fabrican lo que sigue en cola (el plan se corre hacia adelante).</td></tr>
        <tr><td>Cupos diurnos</td><td>El cupo que deja el nocturno se rellena con el siguiente modelo de la línea y con el faltante no programado (remanente). No se deja ocioso a propósito.</td></tr>
        <tr><td>Calendario</td><td>Pestaña Calendario: filtrar escenario (Base/A–D) y semana 1–12. Cada celda muestra día y noche contra la base.</td></tr>
        <tr><td>Listas</td><td>Las tablas de modelos van por días ganados (mayor impacto primero), no por calendario ni alfabético.</td></tr>
        <tr><td>Fuente</td><td>Planificacion Produccion ACTUAL.xlsx y LOTE NUEVO.xlsx.</td></tr>
      </tbody></table>
    </div>
  </section>
</div>
<footer class="foot">Informe generado del plan ACTUAL vs LOTE NUEVO. Cap nocturna = 50% de la cap actual (máquinas nuevas aún no a régimen). Tablas de modelos ordenadas por días ganados.</footer>
<script>
const EXTRA = {extras};
const COSTO = {costos};
const CAL = null;/*CAL_JSON*/
document.querySelectorAll(".tab").forEach(function(b){{
  b.onclick = function(){{
    document.querySelectorAll(".tab").forEach(function(x){{x.classList.remove("on");}});
    document.querySelectorAll(".sec").forEach(function(x){{x.classList.remove("on");}});
    b.classList.add("on");
    document.getElementById("sec-"+b.dataset.id).classList.add("on");
  }};
}});
(function(){{
  var ctx = document.getElementById("ch1");
  if(!ctx || !window.Chart) return;
  new Chart(ctx, {{
    type: "bar",
    data: {{
      labels: ["Base","A","B","C","D"],
      datasets: [
        {{label:"Pzas extra p/ 16/11", data: EXTRA, backgroundColor:"#1a56db", yAxisID:"y"}},
        {{label:"Costo US$", data: COSTO, backgroundColor:"#9aa7bd", yAxisID:"y2"}}
      ]
    }},
    options: {{
      responsive: true,
      maintainAspectRatio: false,
      plugins: {{legend: {{labels: {{color:"#202124", font: {{size: 13}}}}}}}},
      scales: {{
        x: {{ticks: {{color:"#5f6368"}}, grid: {{color:"#eceff3"}}}},
        y: {{ticks: {{color:"#5f6368"}}, grid: {{color:"#eceff3"}}, position:"left"}},
        y2: {{ticks: {{color:"#5f6368"}}, grid: {{display:false}}, position:"right"}}
      }}
    }}
  }});
}})();
(function(){{
  if(!CAL || !CAL.celdas) return;
  var esc = "D";
  var week = 1;
  var escBox = document.getElementById("cal-esc");
  var weekBox = document.getElementById("cal-week");
  var labels = {{"0":"Base","A":"A","B":"B","C":"C","D":"D"}};
  ["0","A","B","C","D"].forEach(function(id){{
    var b = document.createElement("button");
    b.textContent = labels[id];
    b.dataset.id = id;
    if(id==="D") b.className = "on";
    b.onclick = function(){{ esc = id; mark(escBox, b); render(); }};
    escBox.appendChild(b);
  }});
  (CAL.etiquetas || []).forEach(function(e){{
    var b = document.createElement("button");
    b.textContent = "S"+e.semana;
    b.title = e.label;
    b.dataset.w = e.semana;
    if(e.semana===1) b.className = "on";
    b.onclick = function(){{ week = e.semana; mark(weekBox, b); render(); }};
    weekBox.appendChild(b);
  }});
  function mark(box, b){{
    box.querySelectorAll("button").forEach(function(x){{x.classList.remove("on");}});
    b.classList.add("on");
  }}
  function idx(rows){{
    var m = {{}};
    (rows||[]).forEach(function(c){{
      var k = c.fecha_iso+"|"+c.linea+"|"+c.turno;
      if(!m[k]) m[k] = [];
      m[k].push(c);
    }});
    return m;
  }}
  function fmt(n){{
    return Math.round(n).toString().replace(/\\B(?=(\\d{{3}})+(?!\\d))/g, ".");
  }}
  function render(){{
    var et = (CAL.etiquetas||[]).filter(function(e){{return e.semana===week;}})[0];
    document.getElementById("cal-title").textContent = et ? et.label : ("Semana "+week);
    var days = [];
    if(et && et.lunes_iso){{
      var d0 = new Date(et.lunes_iso+"T00:00:00Z");
      for(var i=0;i<5;i++){{
        var dx = new Date(d0.getTime()+i*86400000);
        var iso = dx.toISOString().slice(0,10);
        var dd = iso.slice(8,10)+"/"+iso.slice(5,7);
        days.push({{iso:iso, label:["Lun","Mar","Mié","Jue","Vie"][i]+" "+dd}});
      }}
    }}
    var base = idx(CAL.celdas["0"]);
    var escM = idx(CAL.celdas[esc]||[]);
    var kpis = (CAL.resumen[esc]||[]).filter(function(w){{return w.semana===week;}})[0] || {{pzas:0,pzas_noche:0,pzas_remanente:0,n_modelos:0}};
    var kb = (CAL.resumen["0"]||[]).filter(function(w){{return w.semana===week;}})[0] || {{pzas:0}};
    document.getElementById("cal-kpis").innerHTML =
      '<div class="kpi"><span>Pzas semana</span><b>'+fmt(kpis.pzas||0)+'</b></div>'+
      '<div class="kpi"><span>Pzas noche</span><b>'+fmt(kpis.pzas_noche||0)+'</b></div>'+
      '<div class="kpi"><span>Remanente</span><b>'+fmt(kpis.pzas_remanente||0)+'</b></div>'+
      '<div class="kpi"><span>Delta vs base</span><b>'+fmt((kpis.pzas||0)-(kb.pzas||0))+'</b></div>';
    var html = '<table class="cal"><thead><tr><th>Línea</th>';
    days.forEach(function(d){{ html += "<th>"+d.label+"</th>"; }});
    html += "</tr></thead><tbody>";
    ["1","2","3","4","5"].forEach(function(lin){{
      html += "<tr><td><b>Línea "+lin+"</b></td>";
      days.forEach(function(d){{
        var kd = d.iso+"|"+lin+"|dia";
        var kn = d.iso+"|"+lin+"|noche";
        var dia = escM[kd]||[];
        var noche = escM[kn]||[];
        var bDia = JSON.stringify((base[kd]||[]).map(function(x){{return x.modelo+"|"+x.qty;}}));
        var eDia = JSON.stringify(dia.map(function(x){{return x.modelo+"|"+x.qty;}}));
        var cls = "";
        if(dia.some(function(x){{return x.remanente;}})) cls = "rem";
        else if(bDia !== eDia) cls = "chg";
        html += "<td class='"+cls+"'>";
        if(!dia.length && !noche.length) html += "<span class='cell-qty'>—</span>";
        dia.forEach(function(x){{
          html += "<div><span class='cell-mod'>"+x.modelo+"</span> <span class='cell-qty'>"+fmt(x.qty)+"</span>";
          if(x.remanente) html += "<span class='chip-r'>cola</span>";
          html += "</div>";
        }});
        noche.forEach(function(x){{
          html += "<span class='chip-n'>Noche · "+x.modelo+" "+fmt(x.qty)+"</span>";
        }});
        html += "</td>";
      }});
      html += "</tr>";
    }});
    html += "</tbody></table>";
    document.getElementById("cal-grid").innerHTML = html;
  }}
  render();
}})();
</script>
</body></html>
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)

def main():
    plan_nuevo = cargar_plan(NUEVO_XLSX)
    plan_actual = cargar_plan(ACTUAL_XLSX)
    payload = construir_payload(plan_nuevo, plan_actual)
    os.makedirs(OUT_DIR, exist_ok=True)
    escribir_excel(payload, OUT_XLSX)
    escribir_html(payload, OUT_HTML)
    with open(OUT_HTML, "r", encoding="utf-8") as f:
        html = f.read()
    cal_payload = {
        "celdas": payload.get("calendario") or {},
        "etiquetas": payload.get("etiquetas_semana") or [],
        "resumen": payload.get("resumen_cal") or {},
    }
    html = html.replace("null;/*CAL_JSON*/", json.dumps(cal_payload, ensure_ascii=False) + ";")
    with open(OUT_HTML, "w", encoding="utf-8") as f:
        f.write(html)
    json_path = os.path.join(OUT_DIR, "informe_nocturnos.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    if os.path.isdir("/opt/cursor"):
        os.makedirs(ARTIFACT_DIR, exist_ok=True)
        for src in (OUT_XLSX, OUT_HTML, json_path):
            shutil.copy2(src, os.path.join(ARTIFACT_DIR, os.path.basename(src)))
    print("Excel:", OUT_XLSX)
    print("HTML:", OUT_HTML)
    print("JSON:", json_path)
    for sid in "0ABCD":
        k = payload["kpis"][sid]
        print(sid, k["n_noches"], k["costo_usd"], k["pzas_nocturno"], k["extra_alm16"], k["pzas_lotes_alm16_esc"])


if __name__ == "__main__":
    main()
