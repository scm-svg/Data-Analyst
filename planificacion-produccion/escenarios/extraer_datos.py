"""Extrae de los Excel de planificación los datos que necesita el informe de escenarios.

Lee el plan de referencia (sin el pedido especial) y los cuatro escenarios, y vuelca
un único JSON con, por cada libro:

  semanas["S1".."S12"]  tablero de costura por línea y día (S1 = hoja «Planificacion»,
                        S2..S12 = hojas «Semana 2».. «Semana 12»)
  proyeccion            hoja «Proyeccion» (fecha estimada de término por modelo)
  proy_headers          cabeceras de esa hoja, en orden
  almacen               hoja «Entrada de Almacen Modelo» (salida de costura y entrada a almacén)
  especial              hoja «Por Hacer - Especial» (SKUs del pedido, líneas asignadas,
                        capacidad diaria y día de inicio)

Uso:
    python3 extraer_datos.py [carpeta_con_los_xlsx] [salida.json]

Por defecto busca los .xlsx en la misma carpeta del script y escribe data.json ahí mismo.
"""
import datetime
import json
import os
import sys

import openpyxl

# Libros de entrada. No se versionan en el repo porque pesan ~2,5 MB cada uno.
LIBROS = {
    "BASE": "Planificacion_sin_pedido_d104.xlsx",
    "A": "Escenario_A_8fee.xlsx",
    "B": "Escenario_B_499e.xlsx",
    "C": "Escenario_C_62ac.xlsx",
    "D": "Escenario_D_e947.xlsx",
}

AQUI = os.path.dirname(os.path.abspath(__file__))
CARPETA = sys.argv[1] if len(sys.argv) > 1 else AQUI
SALIDA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(AQUI, "data.json")

SEMANAS = [("S1", "Planificacion")] + [("S%d" % i, "Semana %d" % i) for i in range(2, 13)]


def iso(v):
    """Fecha -> 'AAAA-MM-DD'. Cualquier otra cosa se devuelve tal cual."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return v


def num(v):
    """Celda de cantidad -> float. '--', vacío y texto no numérico cuentan como 0.

    Algunas hojas guardan las cantidades como texto, así que se intenta convertir.
    """
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def filas(ws, min_row):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        yield list(row) + [None] * 24


def tablero(wb, hoja):
    """Tablero semanal: fila 2 las fechas, fila 3 las cabeceras, datos desde la 4.

    La columna «Linea» solo trae valor en la primera fila de cada línea, así que se
    arrastra hacia abajo. La fila de totales (sin modelo) cierra el tablero.
    """
    if hoja not in wb.sheetnames:
        return None
    ws = wb[hoja]
    cabecera = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    fechas = [iso(cabecera[i]) for i in range(5, 10)]
    if not all(isinstance(f, str) for f in fechas):
        return None

    out = []
    linea = None
    for row in filas(ws, 4):
        if row[1] not in (None, ""):
            linea = str(int(row[1])) if isinstance(row[1], (int, float)) else str(row[1]).strip()
        modelo = row[2]
        if modelo in (None, ""):
            break
        out.append({
            "linea": linea,
            "modelo": str(modelo).strip(),
            "mos": num(row[3]),
            "solicitada": num(row[4]),
            "dias": [num(row[i]) for i in range(5, 10)],
            "total": num(row[10]),
        })
    return {"fechas": fechas, "filas": out}


def proyeccion(wb):
    """Hoja «Proyeccion»: cabeceras en la fila 2, datos desde la 3."""
    ws = wb["Proyeccion"]
    cab = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    headers = []
    for v in cab[1:]:
        if v in (None, ""):
            break
        headers.append(str(v).strip())

    out = []
    for row in filas(ws, 3):
        if row[1] in (None, ""):
            break
        fila = {"modelo": str(row[1]).strip()}
        for j, h in enumerate(headers):
            fila[h] = iso(row[1 + j])
        out.append(fila)
    return headers, out


def almacen(wb):
    """Hoja «Entrada de Almacen Modelo»: datos desde la fila 3."""
    ws = wb["Entrada de Almacen Modelo"]
    out = []
    for row in filas(ws, 3):
        if row[2] in (None, ""):
            break
        out.append({
            "mos": num(row[1]),
            "modelo": str(row[2]).strip(),
            "cantidad": num(row[3]),
            "salida": iso(row[4]),
            "entrada": iso(row[5]),
        })
    return out


def especial(wb):
    """Hoja «Por Hacer - Especial»: un SKU del pedido por fila, desde la 3."""
    if "Por Hacer - Especial" not in wb.sheetnames:
        return []
    ws = wb["Por Hacer - Especial"]
    out = []
    for row in filas(ws, 3):
        if row[4] in (None, ""):
            break
        out.append({
            "mo": str(row[4]).strip(),
            "sku": str(row[5]).strip(),
            "producto": str(row[6]).strip() if row[6] else "",
            "genero": str(row[7]).strip() if row[7] else "",
            "talla": str(row[9]).strip() if row[9] else "",
            "lineas": str(row[10]).strip(),
            "cant": num(row[11]),
            "cap": num(row[14]),
            "inicio": iso(row[15]),
        })
    return out


def leer(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    try:
        semanas = {}
        for clave, hoja in SEMANAS:
            t = tablero(wb, hoja)
            if t:
                semanas[clave] = t
        headers, proy = proyeccion(wb)
        return {
            "semanas": semanas,
            "proyeccion": proy,
            "almacen": almacen(wb),
            "especial": especial(wb),
            "proy_headers": headers,
        }
    finally:
        wb.close()


datos = {}
for clave, nombre in LIBROS.items():
    ruta = os.path.join(CARPETA, nombre)
    datos[clave] = leer(ruta)
    d = datos[clave]
    print("%-4s %-38s semanas=%2d modelos_proyeccion=%3d almacen=%3d skus_pedido=%2d" % (
        clave, nombre, len(d["semanas"]), len(d["proyeccion"]), len(d["almacen"]), len(d["especial"])))

json.dump(datos, open(SALIDA, "w"), ensure_ascii=False)
print("escrito:", SALIDA, os.path.getsize(SALIDA), "bytes")
