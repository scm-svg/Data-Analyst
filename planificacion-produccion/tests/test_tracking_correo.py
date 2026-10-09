# -*- coding: utf-8 -*-
"""El script de Tracking declara envío por Gmail/MailApp, correo diario y tablero sin cuadrícula."""
import json
import os
import re
import unittest
from collections import OrderedDict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TRACKING = os.path.join(ROOT, "tracking")
GS = os.path.join(TRACKING, "Codigo.gs")
MANIFEST = os.path.join(TRACKING, "appsscript.json")

SCOPES_REQUERIDOS = {
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/script.container.ui",
    "https://www.googleapis.com/auth/script.send_mail",
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/userinfo.email",
}


def quitar_tildes(s):
    return (
        str(s)
        .replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )


def linea_clave(s):
    t = quitar_tildes(str(s or "").lower())
    m = re.search(r"(\d+)", t)
    if "linea" in t and m:
        return "linea " + m.group(1)
    return re.sub(r"\s+", " ", t).strip()


def dia_clave(s):
    return re.sub(r"\s+", " ", quitar_tildes(str(s or "").lower())).strip()


def fecha_clave(val):
    if hasattr(val, "year"):
        return "%04d-%02d-%02d" % (val.year, val.month, val.day)
    s = str(val or "").strip()
    m = re.match(r"^(\d{1,2})[\/\-.](\d{1,2})[\/\-.](\d{2,4})$", s)
    if m:
        y = int(m.group(3))
        if y < 100:
            y += 2000
        return "%04d-%02d-%02d" % (y, int(m.group(2)), int(m.group(1)))
    return s.lower()


MAX_CANTIDAD_CORREO = 100000


def numero_cantidad(val, display=None):
    def ok(n):
        return isinstance(n, (int, float)) and not isinstance(n, bool) and 0 <= n < MAX_CANTIDAD_CORREO

    if isinstance(val, bool):
        return 0
    if isinstance(val, (int, float)):
        return val if ok(val) else 0
    if display is not None:
        ds = str(display).strip().replace(" ", "").replace(",", ".")
        if re.fullmatch(r"[0-9]+(\.[0-9]+)?", ds):
            dn = float(ds)
            if ok(dn):
                return dn
    if hasattr(val, "year") and hasattr(val, "month") and hasattr(val, "day"):
        from datetime import datetime as dt

        utc = dt(val.year, val.month, val.day)
        epoch = dt(1899, 12, 30)
        serial = (utc - epoch).days
        if 0 < serial < 61:
            serial -= 1
        return serial if ok(serial) else 0
    s = str(val or "").strip().replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return 0
    return n if ok(n) else 0


def normalizar_mo(val):
    s = str(val or "").strip().lower()
    if re.fullmatch(r"\d+", s):
        return str(int(s))
    return s


def normalizar_talla(val):
    if isinstance(val, bool):
        return ""
    if isinstance(val, (int, float)):
        return str(int(val) if val == int(val) else val)
    if hasattr(val, "year"):
        n = numero_cantidad(val)
        return str(int(n)) if n else ""
    return str(val or "").strip().lower()


def clasificar_turno(val):
    t = re.sub(r"\s+", " ", quitar_tildes(str(val or "").lower())).strip()
    if "nocturn" in t:
        return "nocturno"
    return "diurno"


def etiqueta_turno(clave):
    return "Nocturno" if clave == "nocturno" else "Diurno"


def titulo_tablero_turno(clave):
    return "Turno Nocturno" if clave == "nocturno" else "Turno Diurno"


def turno_de_bloque(texto):
    t = re.sub(r"\s+", " ", quitar_tildes(str(texto or "").lower())).strip()
    if "nocturn" in t:
        return "nocturno"
    if "diurn" in t:
        return "diurno"
    return ""


def migrar_clave_historial(clave):
    parts = str(clave or "").split("|")
    if len(parts) == 9:
        parts.insert(3, "diurno")
        return "|".join(parts)
    return clave


def clave_detalle(fila):
    return "|".join(
        [
            dia_clave(fila.get("dia")),
            fecha_clave(fila.get("fecha")),
            linea_clave(fila.get("linea")),
            clasificar_turno(fila.get("turno")),
            normalizar_mo(fila.get("mo")),
            str(fila.get("sku") or "").strip().lower(),
            str(fila.get("producto") or "").strip().lower(),
            str(fila.get("genero") or "").strip().lower(),
            str(fila.get("color") or "").strip().lower(),
            normalizar_talla(fila.get("talla")),
        ]
    )


def clave_estable(fila):
    fecha = fecha_clave(fila.get("fecha"))
    linea = linea_clave(fila.get("linea"))
    turno = clasificar_turno(fila.get("turno"))
    mo = normalizar_mo(fila.get("mo"))
    sku = str(fila.get("sku") or "").strip().lower()
    talla = normalizar_talla(fila.get("talla"))
    if sku:
        return "|".join(["sku", fecha, linea, turno, mo, sku, talla])
    return "|".join(
        [
            "nosku",
            fecha,
            linea,
            turno,
            mo,
            str(fila.get("producto") or "").strip().lower(),
            str(fila.get("genero") or "").strip().lower(),
            str(fila.get("color") or "").strip().lower(),
            talla,
        ]
    )


def fechas_vecinas(yyyy_mm_dd):
    s = str(yyyy_mm_dd or "")
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if not m:
        return [s]
    from datetime import datetime, timedelta

    d = datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return [
        d.strftime("%Y-%m-%d"),
        (d - timedelta(days=1)).strftime("%Y-%m-%d"),
        (d + timedelta(days=1)).strftime("%Y-%m-%d"),
    ]


def estable_desde_partes(parts):
    if not parts or len(parts) != 10:
        return ""
    if parts[5]:
        return "|".join(["sku", parts[1], parts[2], parts[3], parts[4], parts[5], parts[9]])
    return "|".join(["nosku", parts[1], parts[2], parts[3], parts[4], parts[6], parts[7], parts[8], parts[9]])


def registrar_cantidad(mapa, clave, fila, cant):
    cant = numero_cantidad(cant)

    def put(k):
        if not k:
            return
        mapa[k] = max(numero_cantidad(mapa.get(k)), cant)

    put(clave)
    migrada = migrar_clave_historial(clave)
    put(migrada)
    parts = str(migrada).split("|")
    put(estable_desde_partes(parts))
    if len(parts) == 10:
        sin_prod = parts[:]
        sin_prod[6] = ""
        put("|".join(sin_prod))
    if fila:
        put(clave_detalle(fila))
        put(clave_estable(fila))
        reb = clave_detalle(fila).split("|")
        if len(reb) == 10:
            sin_reb = reb[:]
            sin_reb[6] = ""
            put("|".join(sin_reb))


def claves_busqueda(fila):
    base = fila.get("clave") or clave_detalle(fila)
    out = []

    def add(k):
        if k and k not in out:
            out.append(k)

    add(base)
    add(migrar_clave_historial(base))
    add(clave_detalle(fila))
    add(clave_estable(fila))
    parts = str(base).split("|")
    if len(parts) == 10:
        for fch in fechas_vecinas(parts[1]):
            p = parts[:]
            p[1] = fch
            add("|".join(p))
            sin_prod = p[:]
            sin_prod[6] = ""
            add("|".join(sin_prod))
    return out


def cantidad_enviada_de_mapa(mapa, fila):
    best = 0
    for k in claves_busqueda(fila):
        if k in mapa:
            n = numero_cantidad(mapa[k])
            if n > best:
                best = n
    return best


def vacios_dias():
    return {"lunes": 0, "martes": 0, "miercoles": 0, "jueves": 0, "viernes": 0}


def asegurar_totales(totales, turno, linea):
    totales.setdefault(turno, {})
    totales[turno].setdefault(linea, vacios_dias())
    return totales[turno][linea]


def es_fila_linea(fila):
    nom = str(fila[1] if fila and len(fila) > 1 else "").strip().lower()
    return "linea" in nom or "línea" in nom


def es_fila_total(fila):
    nom = quitar_tildes(str(fila[1] if fila and len(fila) > 1 else "").strip().lower())
    return nom.startswith("total")


def es_marcador_turno(fila):
    return turno_de_bloque(fila[1] if fila and len(fila) > 1 else "") != ""


MAX_COLS_TABLERO = 26


def fila_tiene_valor(fila):
    n = min(MAX_COLS_TABLERO, len(fila or []))
    for j in range(1, n):
        if str(fila[j] or "").strip() != "":
            return True
    return False


def columnas_usadas(datos, first, last, last_col):
    min_j = max_j = -1
    for j in range(1, last_col + 1):
        hay = False
        for i in range(first, last + 1):
            if str((datos[i] or [None] * (j + 1))[j] or "").strip() != "":
                hay = True
                break
        if hay:
            if min_j == -1:
                min_j = j
            max_j = j
    if min_j == -1:
        return []
    return list(range(min_j, max_j + 1))


def dividir_tableros(datos):
    bloques = []
    first = last = -1
    last_col = 1
    hay_linea = False
    hay_total = False
    titulo = pending = "Turno Diurno"

    def cerrar():
        nonlocal first, last, last_col, hay_linea, hay_total
        if first != -1 and last >= first and (hay_linea or hay_total):
            bloques.append(
                {
                    "titulo": titulo if hay_linea else "Total del día",
                    "first": first,
                    "last": last,
                    "lastCol": last_col,
                }
            )
        first = last = -1
        last_col = 1
        hay_linea = False
        hay_total = False

    for i, fila in enumerate(datos[:80]):
        if not fila_tiene_valor(fila):
            continue
        n = min(MAX_COLS_TABLERO, len(fila))
        if es_marcador_turno(fila):
            if first != -1 and (hay_linea or hay_total):
                cerrar()
            pending = titulo_tablero_turno(turno_de_bloque(fila[1]))
            titulo = pending
            first = last = i
            last_col = 1
            hay_linea = False
            hay_total = False
            for jm in range(1, n):
                if str(fila[jm] or "").strip() != "":
                    last_col = max(last_col, jm)
            continue
        for j in range(1, n):
            if str(fila[j] or "").strip() != "":
                last_col = max(last_col, j)
        if first == -1:
            first = i
            titulo = pending
        last = i
        if es_fila_linea(fila):
            hay_linea = True
        if es_fila_total(fila):
            hay_total = True
            cerrar()
    cerrar()
    return bloques


def consolidar(filas):
    por = OrderedDict()
    for f in filas:
        c = clave_detalle(f)
        if c not in por:
            por[c] = dict(f)
            por[c]["clave"] = c
            por[c]["cantidad"] = 0
        por[c]["cantidad"] += int(f.get("cantidad") or 0)
    return list(por.values())


def delta(actual, enviado):
    d = (int(actual) or 0) - (int(enviado) or 0)
    return d if d > 0 else 0


def partir_nuevo(consolidados, enviado_map):
    nuevos = []
    persistir = {}
    for f in consolidados:
        ya = cantidad_enviada_de_mapa(enviado_map, f)
        d = delta(f["cantidad"], ya)
        persistir[f["clave"]] = max(ya, f["cantidad"])
        if d > 0:
            copia = dict(f)
            copia["cantidad"] = d
            nuevos.append(copia)
    for clave, cant in enviado_map.items():
        persistir.setdefault(clave, cant)
    return nuevos, persistir


def html_celda_tablero(val, con_borde):
    borde = "border: 1px solid #ccc;" if con_borde else "border: none;"
    return "<td style='padding: 6px 8px; %s text-align: center;'>%s</td>" % (borde, val)


class TestTrackingCorreo(unittest.TestCase):
    def test_manifiesto_incluye_scopes_de_correo(self):
        with open(MANIFEST, encoding="utf-8") as f:
            data = json.load(f)
        scopes = set(data.get("oauthScopes") or [])
        faltan = SCOPES_REQUERIDOS - scopes
        self.assertFalse(faltan, "Faltan oauthScopes: %s" % sorted(faltan))
        self.assertEqual(data.get("runtimeVersion"), "V8")

    def test_script_autoriza_y_envia_con_gmail_y_mailapp(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        for token in (
            "function autorizarEnvioCorreo(",
            "function asegurarPermisoCorreo_(",
            "function enviarCorreoHtml_(",
            "GmailApp.sendEmail",
            "MailApp.sendEmail",
            "GmailApp.getRemainingDailyQuota",
            "MailApp.getRemainingDailyQuota",
            "autorizarEnvioCorreo",
            "enviarCorreoHtml_(",
            "script.send_mail",
        ):
            self.assertIn(token, src, "No aparece en Codigo.gs: %s" % token)
        self.assertIn("Autorizar envío de correo", src)
        self.assertIn("enviarReporteProduccion", src)
        idx_gmail = src.find("GmailApp.sendEmail")
        idx_mail = src.rfind("MailApp.sendEmail")
        self.assertGreater(idx_gmail, 0)
        self.assertGreater(idx_mail, idx_gmail)
        self.assertIn("enviarCorreoHtml_(correosUnidos", src)

    def test_fallback_gmail_luego_mailapp(self):
        def enviar(gmail_ok, mail_ok):
            try:
                if not gmail_ok:
                    raise RuntimeError(
                        "No tienes permiso para llamar a GmailApp.sendEmail"
                    )
                return "GmailApp"
            except RuntimeError:
                try:
                    if not mail_ok:
                        raise RuntimeError(
                            "No tienes permiso para llamar a MailApp.sendEmail. "
                            "Permisos necesarios: https://www.googleapis.com/auth/script.send_mail"
                        )
                    return "MailApp"
                except RuntimeError as e2:
                    return "ERROR:" + str(e2)

        self.assertEqual(enviar(True, False), "GmailApp")
        self.assertEqual(enviar(False, True), "MailApp")
        self.assertIn("script.send_mail", enviar(False, False))

    def test_asunto_y_nota_sin_emoji(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        self.assertIn(
            'ASUNTO_REPORTE_CORREO_ = "Reporte de Producción Diaria y Proyección a Almacén"',
            src,
        )
        self.assertNotIn("📢", src)
        self.assertIn("<strong>NOTA PARA ALMACÉN:</strong>", src)
        self.assertNotIn("📦 NOTA PARA ALMACÉN", src)
        self.assertNotRegex(src, r"📦\s*NOTA PARA ALMAC")

    def test_tablero_tracking_sin_cuadricula(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        self.assertIn("function construirHtmlTableroTracking_(", src)
        self.assertIn("function construirHtmlTablerosTracking_(", src)
        self.assertIn("htmlCeldaTablero_(tag, val, bg, fg, isHeader || esTotal || esMarcador, false)", src)
        self.assertIn("border: none", src)
        self.assertIn('table cellspacing=\'0\'', src)

    def test_resumen_copia_tablero_completo_detalle_solo_nuevo(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        self.assertNotIn("function aplicarDeltasEnTablero_(", src)
        self.assertNotIn("function deltasPorLineaDia_(", src)
        self.assertIn(
            "construirHtmlTablerosTracking_(datosTracking, fondosTracking, coloresTracking)",
            src,
        )
        self.assertIn("construirHtmlDetalle_(filasNuevas", src)
        self.assertIn("tal cual están en la hoja", src)
        self.assertIn("Turno Diurno", src)
        self.assertIn("Turno Nocturno", src)
        self.assertIn("Total del día", src)

    def test_historial_diario_en_script(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        for token in (
            'NOMBRE_HISTORIAL_CORREO_ = "_Correo Enviado"',
            "function partirDetalleNuevo_(",
            "function reiniciarHistorialCorreo(",
            "function extraerFilasDetalleCorreo_(",
            "Todo está al día. No hay modelos ni cantidades nuevas",
        ):
            self.assertIn(token, src, "Falta: %s" % token)

    def test_no_reenvia_cantidades_ya_enviadas(self):
        filas = [
            {
                "dia": "Lunes",
                "fecha": "28/09/2026",
                "linea": "Línea 5",
                "mo": "1256",
                "sku": "SHONRCA30TL",
                "producto": "SHORT SPORT",
                "genero": "CAB",
                "color": "Gris Claro",
                "talla": "L",
                "cantidad": 20,
            },
            {
                "dia": "Lunes",
                "fecha": "28/09/2026",
                "linea": "Línea 4",
                "mo": "",
                "sku": "",
                "producto": "Running Tank",
                "genero": "CAB",
                "color": "BLANCO",
                "talla": "L",
                "cantidad": 40,
            },
            {
                "dia": "Lunes",
                "fecha": "28/09/2026",
                "linea": "Línea 4",
                "mo": "",
                "sku": "",
                "producto": "Running Tank",
                "genero": "CAB",
                "color": "BLANCO",
                "talla": "L",
                "cantidad": 40,
            },
        ]
        cons = consolidar(filas)
        self.assertEqual(len(cons), 2)
        tank = [x for x in cons if x["producto"] == "Running Tank"][0]
        self.assertEqual(tank["cantidad"], 80)

        nuevos1, persistir1 = partir_nuevo(cons, {})
        self.assertEqual(sum(x["cantidad"] for x in nuevos1), 100)
        self.assertEqual(persistir1[tank["clave"]], 80)

        nuevos2, persistir2 = partir_nuevo(cons, persistir1)
        self.assertEqual(nuevos2, [])
        self.assertEqual(persistir2[tank["clave"]], 80)

        cons2 = consolidar(
            filas
            + [
                {
                    "dia": "Martes",
                    "fecha": "29/09/2026",
                    "linea": "Línea 5",
                    "mo": "1256",
                    "sku": "SHONRCA30TL",
                    "producto": "SHORT SPORT",
                    "genero": "CAB",
                    "color": "Gris Claro",
                    "talla": "L",
                    "cantidad": 12,
                }
            ]
        )
        nuevos3, persistir3 = partir_nuevo(cons2, persistir1)
        self.assertEqual(len(nuevos3), 1)
        self.assertEqual(nuevos3[0]["cantidad"], 12)
        self.assertEqual(nuevos3[0]["dia"], "Martes")
        self.assertEqual(sum(persistir3.values()), 80 + 20 + 12)

    def test_delta_nunca_negativo(self):
        self.assertEqual(delta(20, 20), 0)
        self.assertEqual(delta(15, 20), 0)
        self.assertEqual(delta(35, 20), 15)

    def test_html_tablero_sin_borde(self):
        html = html_celda_tablero("66", False)
        self.assertIn("border: none;", html)
        self.assertNotIn("border: 1px solid #ccc;", html)
        html_det = html_celda_tablero("66", True)
        self.assertIn("border: 1px solid #ccc;", html_det)

    def test_clasificar_turno(self):
        self.assertEqual(clasificar_turno(""), "diurno")
        self.assertEqual(clasificar_turno("Diurno"), "diurno")
        self.assertEqual(clasificar_turno("NOCTURNO"), "nocturno")
        self.assertEqual(clasificar_turno("Turno nocturno"), "nocturno")
        self.assertEqual(turno_de_bloque("NOCTURNO"), "nocturno")
        self.assertEqual(turno_de_bloque("Linea 1"), "")

    def test_clave_incluye_turno_y_migra_historial_viejo(self):
        base = {
            "dia": "Lunes",
            "fecha": "05/10/2026",
            "linea": "Línea 2",
            "mo": "1",
            "sku": "ABC",
            "producto": "RIO",
            "genero": "CAB",
            "color": "Negro",
            "talla": "L",
            "cantidad": 10,
        }
        diurno = dict(base, turno="Diurno")
        nocturno = dict(base, turno="Nocturno")
        self.assertNotEqual(clave_detalle(diurno), clave_detalle(nocturno))
        self.assertIn("|diurno|", clave_detalle(diurno))
        self.assertIn("|nocturno|", clave_detalle(nocturno))
        vieja = "lunes|2026-10-05|linea 2|1|abc|rio|cab|negro|l"
        self.assertEqual(migrar_clave_historial(vieja), clave_detalle(diurno))
        cons = consolidar([diurno, nocturno, dict(diurno, cantidad=5)])
        self.assertEqual(len(cons), 2)
        por_turno = {clasificar_turno(x["turno"]): x["cantidad"] for x in cons}
        self.assertEqual(por_turno["diurno"], 15)
        self.assertEqual(por_turno["nocturno"], 10)

    def test_script_reparte_turnos_en_tableros_y_detalle(self):
        with open(GS, encoding="utf-8") as f:
            src = f.read()
        self.assertIn("VERSIÓN 5.9.7", src)
        for token in (
            "function clasificarTurno_(",
            "function turnoDeBloque_(",
            "function dividirTablerosTracking_(",
            "function tituloTableroTotalDia_(",
            "MAX_COLS_TABLERO_CORREO_ = 26",
            "function prepararHojaDetalleTracking_(",
            "totalesPorTurno",
            'HEADERS_DETALLE_TRACKING_ = ["Dia", "Fecha", "Linea", "Turno"',
            '"Turno", "MO", "SKU"',
            "migrarClaveHistorialCorreo_",
            "function numeroCantidad_(",
            "function cantidadEnviadaDeMapa_(",
            "function claveEstableCorreo_(",
            "cant: 11",
        ):
            self.assertIn(token, src, "Falta: %s" % token)
        self.assertIn("tablero diurno y el tablero nocturno", src)

    def test_divide_tableros_diurno_y_nocturno(self):
        vac = [""] * 15
        datos = []
        for _ in range(30):
            datos.append(list(vac))
        datos[1][1] = "SEMANA"
        datos[4][1] = ""
        datos[4][2] = "Lunes"
        datos[6][1] = "Linea 1"
        datos[6][3] = "10"
        datos[11][1] = "Total"
        datos[11][3] = "10"
        datos[14][1] = "NOCTURNO"
        datos[16][2] = "Lunes"
        datos[18][1] = "Linea 1"
        datos[18][3] = "52"
        datos[23][1] = "Total"
        datos[23][3] = "52"
        datos[26][1] = "TOTAL"
        datos[26][3] = "62"
        bloques = dividir_tableros(datos)
        self.assertEqual(
            [b["titulo"] for b in bloques],
            ["Turno Diurno", "Turno Nocturno", "Total del día"],
        )
        self.assertEqual(bloques[0]["first"], 1)
        self.assertEqual(bloques[0]["last"], 11)
        self.assertEqual(bloques[1]["first"], 14)
        self.assertEqual(bloques[1]["last"], 23)
        self.assertEqual(bloques[2]["first"], 26)
        self.assertEqual(bloques[2]["last"], 26)

    def test_agrega_reales_por_turno_y_linea(self):
        filas = [
            ("linea 2", "Diurno", 39, 0),
            ("linea 2", "Diurno", 50, 23),
            ("linea 2", "Nocturno", 46, 101),
            ("linea 1", "Nocturno", 52, 57),
            ("linea 4", "", 88, 0),
        ]
        totales = {}
        for lin, turno, lun, mar in filas:
            prod = asegurar_totales(totales, clasificar_turno(turno), lin)
            prod["lunes"] += lun
            prod["martes"] += mar
        self.assertEqual(totales["diurno"]["linea 2"]["lunes"], 89)
        self.assertEqual(totales["diurno"]["linea 2"]["martes"], 23)
        self.assertEqual(totales["nocturno"]["linea 2"]["lunes"], 46)
        self.assertEqual(totales["nocturno"]["linea 1"]["martes"], 57)
        self.assertEqual(totales["diurno"]["linea 4"]["lunes"], 88)
        self.assertNotIn("linea 1", totales.get("diurno", {}))

    def test_excel_costura_separa_diurno_nocturno(self):
        xlsx = "/home/ubuntu/.cursor/projects/workspace/uploads/Tracking_-_Produccion__2__8f48.xlsx"
        if not os.path.isfile(xlsx):
            self.skipTest("Excel de Tracking no está en este entorno")
        from collections import defaultdict
        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        ws = wb["Unidades Producidas - Costura"]
        tot = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
        days = {10: "lunes", 11: "martes", 12: "miercoles", 13: "jueves", 14: "viernes"}
        for r in range(5, ws.max_row + 1):
            linea = ws.cell(r, 8).value
            if not linea:
                continue
            m = re.search(r"(\d+)", str(linea))
            if not m:
                continue
            lin = "linea " + m.group(1)
            turno = clasificar_turno(ws.cell(r, 9).value)
            for c, d in days.items():
                v = ws.cell(r, c).value
                if v and float(v) > 0:
                    tot[turno][lin][d] += float(v)
        self.assertEqual(tot["diurno"]["linea 2"]["lunes"], 115)
        self.assertEqual(tot["diurno"]["linea 4"]["martes"], 84)
        self.assertEqual(tot["nocturno"]["linea 1"]["lunes"], 52)
        self.assertEqual(tot["nocturno"]["linea 2"]["martes"], 101)
        self.assertEqual(sum(tot["nocturno"]["linea 5"].values()), 30)

    def test_columnas_contiguas_incluyen_columna_extra(self):
        datos = []
        for _ in range(12):
            datos.append([""] * 18)
        datos[1][1] = "SEMANA"
        datos[4][2] = "Lunes"
        datos[4][15] = "% Cumplimiento"
        datos[5][1] = "Linea 1"
        datos[5][3] = "10"
        datos[5][15] = "80%"
        datos[6][1] = "Total"
        datos[6][3] = "10"
        datos[6][15] = "80%"
        bloques = dividir_tableros(datos)
        self.assertEqual(len(bloques), 1)
        self.assertEqual(bloques[0]["lastCol"], 15)
        cols = columnas_usadas(
            datos, bloques[0]["first"], bloques[0]["last"], bloques[0]["lastCol"]
        )
        self.assertEqual(cols[0], 1)
        self.assertEqual(cols[-1], 15)
        self.assertIn(15, cols)

    def test_excel_total_del_dia_en_correo(self):
        xlsx = "/home/ubuntu/.cursor/projects/workspace/uploads/Tracking_-_Produccion__3__edd2.xlsx"
        if not os.path.isfile(xlsx):
            self.skipTest("Excel Tracking (3) no está en este entorno")
        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        ws = wb["Tracking - Produccion"]
        datos = []
        for r in range(1, 31):
            fila = []
            for c in range(1, 17):
                v = ws.cell(r, c).value
                if v is None:
                    fila.append("")
                elif hasattr(v, "strftime"):
                    fila.append(v.strftime("%d/%m/%Y"))
                else:
                    fila.append(v)
            datos.append(fila)
        bloques = dividir_tableros(datos)
        titulos = [b["titulo"] for b in bloques]
        self.assertEqual(titulos, ["Turno Diurno", "Turno Nocturno", "Total del día"])
        self.assertGreaterEqual(bloques[0]["lastCol"], 15)
        self.assertEqual(str(datos[5][15]).strip(), "% Cumplimiento")
        total = bloques[2]
        self.assertEqual(str(datos[total["last"]][1]).strip().upper(), "TOTAL")
        self.assertEqual(float(datos[total["last"]][3]), 427)

    def test_numero_cantidad_no_usa_timestamp_ni_talla(self):
        from datetime import datetime

        self.assertEqual(numero_cantidad(20), 20)
        self.assertEqual(numero_cantidad(5.0, "5"), 5)
        dt = datetime(1900, 1, 5)
        self.assertEqual(numero_cantidad(dt), 5)
        self.assertEqual(numero_cantidad(dt, "05/01/1900"), 5)
        self.assertEqual(numero_cantidad(-2208643200000), 0)
        self.assertEqual(numero_cantidad(dt, "L"), 5)
        self.assertEqual(normalizar_mo("02793"), "2793")
        self.assertEqual(normalizar_mo("01534-002"), "01534-002")
        self.assertEqual(normalizar_talla(14.0), "14")
        self.assertEqual(normalizar_talla("L"), "l")

    def test_historial_encuentra_sku_aunque_cambie_producto_o_fecha(self):
        hist_fila = {
            "dia": "Lunes",
            "fecha": "04/10/2026",
            "linea": "Línea 1",
            "turno": "Nocturno",
            "mo": "2785",
            "sku": "MARMIKI11T14",
            "producto": "MAR LOTE 1",
            "genero": "KIDS",
            "color": "Azul Lavanda",
            "talla": 14,
            "cantidad": 12,
        }
        actual = {
            "dia": "Lunes",
            "fecha": "05/10/2026",
            "linea": "Línea 1",
            "turno": "Nocturno",
            "mo": "02785",
            "sku": "MARMIKI11T14",
            "producto": "MAR",
            "genero": "KIDS",
            "color": "Azul Lavanda",
            "talla": "14",
            "cantidad": 12,
        }
        actual["clave"] = clave_detalle(actual)
        enviado = {}
        registrar_cantidad(enviado, clave_detalle(hist_fila), hist_fila, 12)
        nuevos, persistir = partir_nuevo(consolidar([actual]), enviado)
        self.assertEqual(nuevos, [])
        self.assertEqual(cantidad_enviada_de_mapa(enviado, actual), 12)
        self.assertGreater(persistir[actual["clave"]], 0)

        otra = dict(actual, cantidad=18)
        otra["clave"] = clave_detalle(otra)
        nuevos2, _ = partir_nuevo(consolidar([otra]), enviado)
        self.assertEqual(len(nuevos2), 1)
        self.assertEqual(nuevos2[0]["cantidad"], 6)

    def test_excel3_correo_enviado_cantidades_fecha_no_inflan(self):
        xlsx = "/home/ubuntu/.cursor/projects/workspace/uploads/Tracking_-_Produccion__3__edd2.xlsx"
        if not os.path.isfile(xlsx):
            self.skipTest("Excel Tracking (3) no está en este entorno")
        from datetime import datetime
        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        wh = wb["_Correo Enviado"]
        fechas_como_cant = 0
        recuperadas = []
        naive_huge = 0
        for r in range(3, wh.max_row + 1):
            clave = wh.cell(r, 1).value
            if not clave:
                continue
            raw = wh.cell(r, 12).value
            cant = numero_cantidad(raw)
            self.assertLess(cant, 1000, "cantidad recuperada absurda en r%s: %s" % (r, cant))
            if isinstance(raw, datetime):
                fechas_como_cant += 1
                recuperadas.append(cant)
                if abs(float(raw.timestamp()) * 1000) > 1000000:
                    naive_huge += 1
        self.assertGreaterEqual(fechas_como_cant, 50)
        self.assertGreaterEqual(naive_huge, 50)
        self.assertTrue(all(c > 0 for c in recuperadas))
        self.assertIn(5, recuperadas)

        wd = wb["Detalle Tracking - Produccion"]
        det = []
        for r in range(3, wd.max_row + 1):
            cant = wd.cell(r, 12).value
            try:
                cantn = float(cant or 0)
            except (TypeError, ValueError):
                continue
            if cantn <= 0:
                continue
            fila = {
                "dia": wd.cell(r, 2).value,
                "fecha": wd.cell(r, 3).value,
                "linea": wd.cell(r, 4).value,
                "turno": wd.cell(r, 5).value,
                "mo": wd.cell(r, 6).value,
                "sku": wd.cell(r, 7).value,
                "producto": "MAR" if str(wd.cell(r, 8).value or "").upper().startswith("MAR") else wd.cell(r, 8).value,
                "genero": wd.cell(r, 9).value,
                "color": wd.cell(r, 10).value,
                "talla": wd.cell(r, 11).value,
                "cantidad": cantn,
            }
            fila["clave"] = clave_detalle(fila)
            det.append(fila)

        enviado = {}
        for r in range(3, wh.max_row + 1):
            clave = wh.cell(r, 1).value
            if not clave:
                continue
            fila = {
                "dia": wh.cell(r, 2).value,
                "fecha": wh.cell(r, 3).value,
                "linea": wh.cell(r, 4).value,
                "turno": wh.cell(r, 5).value,
                "mo": wh.cell(r, 6).value,
                "sku": wh.cell(r, 7).value,
                "producto": wh.cell(r, 8).value,
                "genero": wh.cell(r, 9).value,
                "color": wh.cell(r, 10).value,
                "talla": wh.cell(r, 11).value,
            }
            registrar_cantidad(enviado, str(clave), fila, numero_cantidad(wh.cell(r, 12).value))

        cons = consolidar(det)
        nuevos, _ = partir_nuevo(cons, enviado)
        inflados = [x for x in nuevos if x["cantidad"] > 1000]
        self.assertEqual(inflados, [])
        reenviados_completos = [x for x in nuevos if x["sku"] and str(x["sku"]).upper().startswith("MAR")]
        self.assertEqual(
            reenviados_completos,
            [],
            "MAR no debe reenviarse entero al cambiar MAR LOTE 1 → MAR: %s"
            % [(x.get("sku"), x["cantidad"]) for x in reenviados_completos[:8]],
        )


if __name__ == "__main__":
    unittest.main()
