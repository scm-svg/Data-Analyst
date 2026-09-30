#!/usr/bin/env python3
"""Lista fija de Impresión Digital: prioridad, L1-4 vs L5, PD y checks MO+SKU."""
from __future__ import annotations

import json
import math
import re
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

PRIO_LABEL = {0: "Especial", 1: "Urgente", 2: "Alta", 3: "Media", 4: "Baja", 5: "Sin Asignar"}
GRUPO_L14 = "lineas14"
GRUPO_L5 = "linea5"


def norm(s):
    return str("" if s is None else s).strip()


def norm_up(s):
    return norm(s).upper()


def quitar_tildes(s):
    table = str.maketrans("áéíóúñÁÉÍÓÚÑ", "aeiounAEIOUN")
    return str(s).translate(table)


def prioridad_num(txt):
    p = quitar_tildes(norm(txt)).lower()
    if "especial" in p:
        return 0
    if "urgente" in p:
        return 1
    if "alta" in p:
        return 2
    if "media" in p:
        return 3
    if "baja" in p:
        return 4
    return 5


def prioridad_label(n, es_especial=False):
    if es_especial:
        return "Especial"
    return PRIO_LABEL.get(int(n) if n is not None else 5, "Sin Asignar")


def parsear_lineas(val):
    if val is None or val == "":
        return []
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        if not math.isfinite(float(val)):
            return []
        n = int(val)
        return [str(n)] if n > 0 else []
    out = []
    seen = set()
    for part in re.split(r"[,;/|]+", str(val)):
        token = re.sub(r"(?i)linea|línea", "", part).strip()
        if token == "":
            continue
        try:
            n = int(float(token.replace(",", ".")))
        except (TypeError, ValueError):
            continue
        if n <= 0:
            continue
        s = str(n)
        if s not in seen:
            seen.add(s)
            out.append(s)
    return out


def es_pd(mo):
    s = norm_up(mo)
    if not s:
        return False
    return bool(re.search(r"(^|[^A-Z0-9])PD($|[^A-Z0-9])", s)) or s.startswith("PD")


def grupo_linea(lineas):
    ls = parsear_lineas(lineas) if isinstance(lineas, (str, int, float)) else list(lineas or [])
    if "5" in ls:
        return GRUPO_L5
    return GRUPO_L14


def clave_check(mo, sku):
    m = norm_up(mo)
    s = norm_up(sku)
    if m:
        return "M|%s|%s" % (m, s)
    return "S|%s" % s


def rango_color(color):
    c = quitar_tildes(norm(color)).lower()
    if not c:
        return 50
    if c == "negro" or c.startswith("negro"):
        return 0
    if c == "blanco" or c.startswith("blanco") or c in ("ivory", "blanco hueso"):
        return 1
    if c.startswith("azul marino") or c.startswith("navy") or c == "marino":
        return 2
    if "negro" in c:
        return 3
    if "blanco" in c or "ivory" in c:
        return 4
    if "marino" in c or "navy" in c:
        return 5
    return 50


def orden_talla(talla):
    s = quitar_tildes(norm(talla)).upper().replace(" ", "")
    mapa = {
        "XXS": 0, "XS": 1, "S": 2, "M": 3, "L": 4, "XL": 5,
        "XXL": 6, "2XL": 6, "XXXL": 7, "3XL": 7, "XXXXL": 8, "4XL": 8,
    }
    if s in mapa:
        return mapa[s]
    m = re.search(r"(\d+)", s)
    if m:
        return 20 + int(m.group(1))
    return 50


def clave_color(color):
    return quitar_tildes(norm(color)).lower()


def cant_sku(s):
    for k in ("faltante", "cant", "solicitada", "cantidad"):
        n = Numberish(s.get(k) if isinstance(s, dict) else None)
        if n > 0:
            return n
    return 0


def Numberish(v):
    try:
        n = float(v)
        return n if math.isfinite(n) else 0.0
    except (TypeError, ValueError):
        return 0.0


def ordenar_skus(filas):
    arr = list(filas or [])
    vol = defaultdict(float)
    for s in arr:
        vol[clave_color(s.get("color"))] += cant_sku(s)

    def key(s):
        ra = rango_color(s.get("color"))
        core = 0 if ra <= 5 else 1
        vol_c = -vol[clave_color(s.get("color"))] if ra > 5 else 0
        return (
            0 if s.get("esSkuPrio") else 1,
            int(s.get("skuPrioOrden") or 0) if s.get("esSkuPrio") else 0,
            core,
            ra if ra <= 5 else 50,
            vol_c,
            str(s.get("color") or ""),
            orden_talla(s.get("talla")),
            -cant_sku(s),
            str(s.get("sku") or ""),
        )

    return sorted(arr, key=key)


def fecha_ms(v):
    if v is None or v == "":
        return float("inf")
    if isinstance(v, datetime):
        return datetime(v.year, v.month, v.day).timestamp() * 1000
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day).timestamp() * 1000
    if isinstance(v, (int, float)) and math.isfinite(v):
        if v > 1e11:
            return float(v)
        if 20000 < v < 80000:
            utc = datetime.utcfromtimestamp(round((v - 25569) * 86400))
            return datetime(utc.year, utc.month, utc.day).timestamp() * 1000
    s = norm(v)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", s)
    if m:
        y = int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return datetime(y, int(m.group(2)), int(m.group(1))).timestamp() * 1000
        except ValueError:
            return float("inf")
    return float("inf")


def fmt_fecha(v):
    ms = fecha_ms(v)
    if not math.isfinite(ms):
        return ""
    d = datetime.fromtimestamp(ms / 1000)
    return d.strftime("%d/%m/%Y")


def agrupar_impresion(backlog):
    """Modelos fijos por prioridad, partidos en L1-4 vs L5. No usa semanas del plan."""
    grupos = {}
    orden = []
    for row in backlog or []:
        modelo = norm(row.get("modelo"))
        if not modelo:
            continue
        if Numberish(row.get("faltante")) <= 0:
            continue
        if modelo not in grupos:
            grupos[modelo] = {
                "modelo": modelo,
                "tipo": "Especial" if row.get("especial") or norm(row.get("tipo")).lower() == "especial" else (row.get("tipo") or "Producción"),
                "especial": bool(row.get("especial") or norm(row.get("tipo")).lower() == "especial"),
                "prioridad": row.get("prioridad") or "Sin Asignar",
                "prioridadNum": 5,
                "fechaMs": float("inf"),
                "fecha": "",
                "lineas": [],
                "skus": [],
                "solicitada": 0.0,
                "producida": 0.0,
                "faltante": 0.0,
                "mos": set(),
                "nPd": 0,
            }
            orden.append(modelo)
        g = grupos[modelo]
        if g["especial"] or row.get("especial"):
            g["especial"] = True
            g["tipo"] = "Especial"
        prio = row.get("prioridad") or g["prioridad"]
        pn = prioridad_num("Especial" if g["especial"] else prio)
        if g["especial"]:
            pn = 0
        if pn < g["prioridadNum"]:
            g["prioridadNum"] = pn
            g["prioridad"] = "Especial" if g["especial"] else (prio or "Sin Asignar")
        elif not g["especial"] and prio and g["prioridad"] in ("", "Sin Asignar"):
            g["prioridad"] = prio
        fms = fecha_ms(row.get("fechaSalida") or row.get("fecha"))
        if fms < g["fechaMs"]:
            g["fechaMs"] = fms
            g["fecha"] = fmt_fecha(row.get("fechaSalida") or row.get("fecha")) or row.get("fecha") or ""
        for lin in parsear_lineas(row.get("linea") or row.get("lineas")):
            if lin not in g["lineas"]:
                g["lineas"].append(lin)
        g["lineas"].sort(key=lambda x: int(x) if str(x).isdigit() else 99)
        sku = {
            "mo": norm(row.get("mo")),
            "sku": norm(row.get("sku")),
            "detalle": norm(row.get("detalle")),
            "color": norm(row.get("color")),
            "talla": norm(row.get("talla")),
            "genero": norm(row.get("genero")),
            "cant": Numberish(row.get("faltante")),
            "solicitada": Numberish(row.get("solicitada")),
            "producida": Numberish(row.get("producida")),
            "faltante": Numberish(row.get("faltante")),
            "moStatus": norm(row.get("moStatus") or row.get("status")),
            "pd": es_pd(row.get("mo")),
            "linea": norm(row.get("linea")),
            "esSkuPrio": bool(row.get("esSkuPrio")),
            "skuPrioOrden": int(row.get("skuPrioOrden") or 9999),
        }
        if not sku["detalle"]:
            bits = [norm(row.get("producto")) or modelo.replace(" (Especial)", "")]
            if sku["genero"] and sku["genero"] != "--":
                bits.append(sku["genero"])
            if sku["color"] and sku["color"] != "--":
                bits.append(sku["color"])
            if sku["talla"] and sku["talla"] != "--":
                bits.append(sku["talla"])
            sku["detalle"] = " - ".join([b for b in bits if b])
            if g["especial"] and "(Especial)" not in sku["detalle"]:
                sku["detalle"] += " (Especial)"
        g["skus"].append(sku)
        g["solicitada"] += sku["solicitada"] or sku["faltante"]
        g["producida"] += sku["producida"]
        g["faltante"] += sku["faltante"]
        if sku["mo"]:
            g["mos"].add(sku["mo"])
        if sku["pd"]:
            g["nPd"] += 1

    modelos = []
    for nombre in orden:
        g = grupos[nombre]
        g["skus"] = ordenar_skus(g["skus"])
        g["nMos"] = len(g["mos"])
        g["nSkus"] = len(g["skus"])
        g["grupo"] = grupo_linea(g["lineas"])
        g["prioridad"] = prioridad_label(g["prioridadNum"], g["especial"])
        g.pop("mos", None)
        modelos.append(g)

    def sort_key(m):
        fec = m["fechaMs"] if math.isfinite(m["fechaMs"]) else float("inf")
        return (m["prioridadNum"], fec, -m["faltante"], m["modelo"])

    l14 = sorted([m for m in modelos if m["grupo"] == GRUPO_L14], key=sort_key)
    l5 = sorted([m for m in modelos if m["grupo"] == GRUPO_L5], key=sort_key)
    return {"lineas14": l14, "linea5": l5, "modelos": l14 + l5}


def _idx_headers(headers):
    low = [quitar_tildes(norm(h)).lower() for h in headers]

    def find(*frags, exact=None):
        if exact:
            alvo = quitar_tildes(exact).lower()
            for i, h in enumerate(low):
                if h == alvo:
                    return i
        for i, h in enumerate(low):
            if h and all(f in h for f in frags):
                return i
        return -1

    return {
        "sku": find("sku", exact="sku"),
        "mo": find(exact="mo") if find(exact="mo") != -1 else find("mo"),
        "prod": find("producto") if find("producto") != -1 else find("modelo"),
        "gen": find("genero"),
        "col": find("color"),
        "tal": find("talla"),
        "lin": find("linea"),
        "sol": find("cantidad", "solicitada"),
        "prodq": find("producida"),
        "falt": find("faltante"),
        "status": find("mo", "status") if find("mo", "status") != -1 else find("status"),
        "prio": find("prioridad"),
        "fecha": find("fecha", "salida") if find("fecha", "salida") != -1 else find("fecha"),
    }


def _faltante(sol, prod, falt_celda):
    if falt_celda not in (None, ""):
        try:
            return max(0.0, float(falt_celda))
        except (TypeError, ValueError):
            pass
    return max(0.0, Numberish(sol) - Numberish(prod))


def leer_embudo_xlsx(ws, especial, prio_map):
    rows = list(ws.iter_rows(min_row=1, max_col=22, values_only=True))
    if len(rows) < 3:
        return []
    headers = [norm(x) for x in rows[1]]
    idx = _idx_headers(headers)
    if idx["sku"] < 0 or idx["prod"] < 0:
        return []
    out = []
    for raw in rows[2:]:
        sku = norm(raw[idx["sku"]]) if idx["sku"] >= 0 else ""
        prod = norm(raw[idx["prod"]]) if idx["prod"] >= 0 else ""
        if not sku or not prod:
            continue
        status = norm(raw[idx["status"]]) if idx["status"] >= 0 else ""
        if especial and quitar_tildes(status).lower() == "hecho":
            continue
        sol = Numberish(raw[idx["sol"]]) if idx["sol"] >= 0 else 0
        prodq = Numberish(raw[idx["prodq"]]) if idx["prodq"] >= 0 else 0
        falt = _faltante(sol, prodq, raw[idx["falt"]] if idx["falt"] >= 0 else None)
        if falt <= 0:
            continue
        gen = norm(raw[idx["gen"]]) if idx["gen"] >= 0 else ""
        color = norm(raw[idx["col"]]) if idx["col"] >= 0 else ""
        talla = norm(raw[idx["tal"]]) if idx["tal"] >= 0 else ""
        modelo = prod + ((" " + gen) if gen and gen != "--" else "")
        if especial:
            modelo += " (Especial)"
        pinfo = prio_map.get(modelo) or prio_map.get(prod) or {}
        prio_fila = norm(raw[idx["prio"]]) if idx["prio"] >= 0 else ""
        prioridad = prio_fila or pinfo.get("prioridad") or "Sin Asignar"
        fecha = raw[idx["fecha"]] if idx["fecha"] >= 0 else None
        if fecha in (None, "") and pinfo.get("fecha") not in (None, ""):
            fecha = pinfo.get("fecha")
        bits = [prod]
        if gen and gen != "--":
            bits.append(gen)
        if color and color != "--":
            bits.append(color)
        if talla and talla != "--":
            bits.append(talla)
        detalle = " - ".join(bits)
        if especial:
            detalle += " (Especial)"
        out.append({
            "sku": sku,
            "modelo": modelo,
            "producto": prod,
            "detalle": detalle,
            "mo": norm(raw[idx["mo"]]) if idx["mo"] >= 0 else "",
            "genero": gen,
            "color": color,
            "talla": talla,
            "linea": raw[idx["lin"]] if idx["lin"] >= 0 else "",
            "solicitada": sol,
            "producida": prodq,
            "faltante": falt,
            "prioridad": prioridad,
            "fechaSalida": fecha,
            "moStatus": status,
            "especial": especial,
            "tipo": "Especial" if especial else "Producción",
        })
    return out


def leer_priorizacion_xlsx(ws):
    rows = list(ws.iter_rows(min_row=1, max_col=12, values_only=True))
    prio_map = {}
    if len(rows) < 3:
        return prio_map
    headers = [quitar_tildes(norm(x)).lower() for x in rows[1]]
    i_mod = headers.index("modelo") if "modelo" in headers else 0
    i_tipo = headers.index("tipo") if "tipo" in headers else -1
    i_prio = headers.index("prioridad") if "prioridad" in headers else -1
    i_fec = next((i for i, h in enumerate(headers) if "fecha" in h), -1)
    i_lin = next((i for i, h in enumerate(headers) if "linea" in h), -1)
    for raw in rows[2:]:
        mod = norm(raw[i_mod]) if i_mod >= 0 else ""
        if not mod:
            continue
        tipo = norm(raw[i_tipo]) if i_tipo >= 0 else ""
        key = mod + " (Especial)" if tipo.lower() == "especial" else mod
        prio_map[key] = {
            "prioridad": norm(raw[i_prio]) if i_prio >= 0 else "",
            "fecha": raw[i_fec] if i_fec >= 0 else None,
            "lineas": raw[i_lin] if i_lin >= 0 else "",
            "tipo": tipo,
        }
    return prio_map


def leer_checks_xlsx(ws):
    listos = {}
    rows = list(ws.iter_rows(min_row=2, max_col=6, values_only=True))
    for raw in rows:
        clave = norm(raw[0])
        sku = norm(raw[1]) if len(raw) > 1 else ""
        mo = norm(raw[2]) if len(raw) > 2 else ""
        listo = quitar_tildes(norm(raw[4] if len(raw) > 4 else "")).upper()
        if not clave or clave == "CLAVE":
            continue
        if clave.startswith("M|") or clave.startswith("S|"):
            estable = clave
        else:
            estable = clave_check(mo, sku)
        if listo in ("SI", "TRUE", "1") and estable:
            listos[estable] = 1
    return listos


def extraer_desde_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    prio_map = {}
    if "Priorizacion" in wb.sheetnames:
        prio_map = leer_priorizacion_xlsx(wb["Priorizacion"])
    backlog = []
    if "Por Hacer" in wb.sheetnames:
        backlog.extend(leer_embudo_xlsx(wb["Por Hacer"], False, prio_map))
    if "Por Hacer - Especial" in wb.sheetnames:
        backlog.extend(leer_embudo_xlsx(wb["Por Hacer - Especial"], True, prio_map))
    checks = {}
    if "_ImpresionChecks" in wb.sheetnames:
        checks = leer_checks_xlsx(wb["_ImpresionChecks"])
    wb.close()
    lista = agrupar_impresion(backlog)
    return {
        "backlog": backlog,
        "checks": checks,
        "lista": lista,
        "version": "5.9.45-id-fija",
    }


def resumen(lista):
    mods = (lista.get("lineas14") or []) + (lista.get("linea5") or [])
    n_sku = sum(m["nSkus"] for m in mods)
    n_pd = sum(m["nPd"] for m in mods)
    pzas = sum(m["faltante"] for m in mods)
    return {
        "modelos": len(mods),
        "l14": len(lista.get("lineas14") or []),
        "l5": len(lista.get("linea5") or []),
        "skus": n_sku,
        "pd": n_pd,
        "pzas": pzas,
    }


def _json_safe(obj):
    if isinstance(obj, dict):
        return {k: _json_safe(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_json_safe(v) for v in obj]
    if isinstance(obj, tuple):
        return [_json_safe(v) for v in obj]
    if isinstance(obj, set):
        return sorted(_json_safe(v) for v in obj)
    if isinstance(obj, datetime):
        return obj.strftime("%d/%m/%Y")
    if isinstance(obj, date):
        return obj.strftime("%d/%m/%Y")
    if isinstance(obj, float):
        if math.isinf(obj) or math.isnan(obj):
            return None
        if obj.is_integer():
            return int(obj)
        return obj
    return obj


def payload_para_html(extraido):
    lista = extraido["lista"]
    kpis = resumen(lista)
    modelos = sorted({b["modelo"] for b in extraido["backlog"] if b.get("modelo")})
    pris = []
    for lab in ("Especial", "Urgente", "Alta", "Media", "Baja", "Sin Asignar"):
        if any(m["prioridad"] == lab for m in lista.get("modelos") or []):
            pris.append(lab)
    lineas = sorted({
        lin
        for m in (lista.get("modelos") or [])
        for lin in m.get("lineas") or []
    }, key=lambda x: int(x) if str(x).isdigit() else 99)
    return _json_safe({
        "version": extraido.get("version") or "5.9.45-id-fija",
        "publicadoEn": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "kpis": kpis,
        "checks": extraido.get("checks") or {},
        "lista": {
            "lineas14": lista.get("lineas14") or [],
            "linea5": lista.get("linea5") or [],
        },
        "opcionesFiltro": {
            "modelos": modelos,
            "prioridades": pris,
            "lineas": lineas,
        },
    })


def escribir_html(payload, dest, template_path=None):
    dest = Path(dest)
    if template_path is None:
        template_path = Path(__file__).with_name("impresion_digital_template.html")
    tpl = Path(template_path).read_text(encoding="utf-8")
    blob = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    if "__DASH_JSON__" not in tpl:
        raise ValueError("La plantilla no tiene el marcador __DASH_JSON__")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(tpl.replace("__DASH_JSON__", blob), encoding="utf-8")
    return dest


def main(argv=None):
    import argparse
    p = argparse.ArgumentParser(description="Genera el HTML de prueba de Impresión Digital")
    p.add_argument("--xlsx", required=True)
    p.add_argument("--html", required=True)
    p.add_argument("--json", default="")
    args = p.parse_args(argv)
    extraido = extraer_desde_xlsx(args.xlsx)
    payload = payload_para_html(extraido)
    escribir_html(payload, args.html)
    if args.json:
        Path(args.json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    k = payload["kpis"]
    print("HTML:", args.html)
    print("modelos={modelos} L1-4={l14} L5={l5} skus={skus} PD={pd} pzas={pzas}".format(**k))
    print("checks SI:", len(payload.get("checks") or {}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
