#!/usr/bin/env python3
"""Motor de escenarios de turno nocturno sobre un plan ya generado.

Noches: martes, miércoles y viernes. Mitad de la cap diurna.
Línea 1 de noche sigue la cola de Línea 2 (100% del cupo nocturno).
No hay nocturno en día 15 ni último del mes. No aplica a Especial.
"""
from __future__ import annotations

from collections import defaultdict, deque
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook


LUNES_BASE = date(2026, 9, 28)
PRIMER_NOCTURNO = date(2026, 9, 29)
CORTE_ALMACEN = date(2026, 11, 16)
# 16/11 es lunes: 4 días hábiles atrás = martes 10/11.
CORTE_COSTURA_ALMACEN = date(2026, 11, 10)
# Costura al 25/11 → almacén ~01/12 (demanda diciembre / tienda nueva).
CORTE_COSTURA_DICIEMBRE = date(2026, 11, 25)
DIAS_ENTRADA_ALMACEN = 4
FRACCION_NOCTURNO = 0.5
BONO_USD = 15.0
PERSONAS = {"1": 4, "2": 4, "3": 4, "4": 5, "5": 5}
CAP_FALLBACK = {"1": 130.0, "2": 130.0, "3": 130.0, "4": 130.0, "5": 40.0}
HOJAS_SEMANA = ["Planificacion"] + [f"Semana {i}" for i in range(2, 13)]
ESCENARIOS = [
    {"id": "0", "semanas": 0, "titulo": "Sin nocturnos (plan actual)"},
    {"id": "A", "semanas": 4, "titulo": "4 semanas de nocturnos"},
]
SEM_NOCHES_A = 4
LOTES_NUEVOS_CLAVES = ("LOTE NUEVO",)


def as_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def num(v) -> float:
    if v in (None, "", "--"):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def es_cobro(d: date) -> bool:
    if d.day == 15:
        return True
    return (d + timedelta(days=1)).month != d.month


def add_business_days(d: date, days: int = DIAS_ENTRADA_ALMACEN) -> date:
    cur = d
    added = 0
    while added < days:
        cur += timedelta(days=1)
        if cur.weekday() < 5:
            added += 1
    return cur


def fechas_nocturno(n_semanas: int, lunes_base: date = LUNES_BASE) -> List[date]:
    """Martes, miércoles y viernes de las primeras n_semanas, sin días de cobro."""
    out: List[date] = []
    for w in range(max(0, n_semanas)):
        mon = lunes_base + timedelta(days=7 * w)
        for offset in (1, 2, 4):
            d = mon + timedelta(days=offset)
            if not es_cobro(d):
                out.append(d)
    return out


def personas_total() -> int:
    return sum(PERSONAS.values())


def costo_bono(n_noches: int, personas: Optional[int] = None) -> float:
    p = personas_total() if personas is None else personas
    return n_noches * p * BONO_USD


def cap_nocturna(cap_dia: float) -> int:
    return int(float(cap_dia) * FRACCION_NOCTURNO)


def es_especial_nombre(modelo: str) -> bool:
    return "(especial)" in (modelo or "").lower()


def es_lote_nuevo(modelo: str) -> bool:
    u = (modelo or "").upper()
    return any(k in u for k in LOTES_NUEVOS_CLAVES)


def fmt_fecha(d: Optional[date]) -> str:
    if not d:
        return "--"
    return d.strftime("%d/%m/%Y")


def delta_dias(a: Optional[date], b: Optional[date]) -> Optional[int]:
    if not a or not b:
        return None
    return (a - b).days


@dataclass
class EventoLinea:
    fecha: date
    linea: str
    modelo: str
    qty: float
    especial: bool
    seq: int = 0


@dataclass
class ModeloInfo:
    modelo: str
    faltante: float = 0.0
    solicitada: float = 0.0
    producida: float = 0.0
    cap: float = 0.0
    lineas: str = ""
    prioridad: str = ""
    fecha_obj: Optional[date] = None
    skus: int = 0
    lote_nuevo: bool = False
    especial: bool = False
    termino_plan: Optional[date] = None
    entrada_plan: Optional[date] = None
    estado_plan: str = ""
    sin_programar: float = 0.0


@dataclass
class ResultadoModelo:
    modelo: str
    faltante: float
    cap: float
    lineas: str
    prioridad: str
    lote_nuevo: bool
    fecha_obj: Optional[date]
    termino_base: Optional[date]
    entrada_base: Optional[date]
    termino_esc: Optional[date]
    entrada_esc: Optional[date]
    primer_base: Optional[date]
    primer_esc: Optional[date]
    pzas_nocturno: float
    pzas_prog: float
    pzas_alm16_base: float
    pzas_alm16_esc: float
    pzas_dic_base: float
    pzas_dic_esc: float
    dias_ganados: Optional[int]
    a_tiempo_base: Optional[bool]
    a_tiempo_esc: Optional[bool]
    cubre_almacen_base: Optional[bool]
    cubre_almacen_esc: Optional[bool]


def _idx(headers: Sequence, *frags: str) -> int:
    for i, h in enumerate(headers):
        hl = str(h or "").lower()
        if all(f in hl for f in frags):
            return i
    return -1


def leer_por_hacer(wb, especial: bool = False) -> List[dict]:
    name = "Por Hacer - Especial" if especial else "Por Hacer"
    if name not in wb.sheetnames:
        return []
    sh = wb[name]
    rows = list(sh.iter_rows(values_only=True))
    if len(rows) < 3:
        return []
    headers = [str(h).strip() if h else "" for h in rows[1]]
    i_mo = next((i for i, h in enumerate(headers) if h.upper() == "MO"), -1)
    i_sku = next((i for i, h in enumerate(headers) if h.upper() == "SKU"), -1)
    i_prod = _idx(headers, "producto")
    i_gen = _idx(headers, "genero") if _idx(headers, "genero") >= 0 else _idx(headers, "género")
    i_lin = _idx(headers, "linea")
    i_sol = _idx(headers, "cantidad", "solicitada")
    i_pro = next((i for i, h in enumerate(headers) if "producida" in h.lower()), -1)
    i_fal = _idx(headers, "faltante")
    i_cap = _idx(headers, "cap")
    i_pri = _idx(headers, "prioridad")
    i_fec = next(
        (i for i, h in enumerate(headers) if "salida" in h.lower() and "fecha" in h.lower()),
        -1,
    )
    out = []
    for r in rows[2:]:
        sku = str(r[i_sku] or "").strip() if i_sku >= 0 else ""
        if not sku:
            continue
        prod = str(r[i_prod] or "").strip() if i_prod >= 0 else ""
        gen = str(r[i_gen] or "").strip() if i_gen >= 0 else ""
        modelo = prod + ((" " + gen) if gen and gen != "--" else "")
        if especial:
            modelo += " (Especial)"
        sol = num(r[i_sol]) if i_sol >= 0 else 0.0
        prodq = num(r[i_pro]) if i_pro >= 0 else 0.0
        fal = num(r[i_fal]) if i_fal >= 0 and r[i_fal] not in (None, "") else max(0.0, sol - prodq)
        if fal <= 0:
            continue
        out.append(
            {
                "sku": sku,
                "mo": str(r[i_mo] or "").strip() if i_mo >= 0 else "",
                "modelo": modelo,
                "lineas": str(r[i_lin] or "").strip() if i_lin >= 0 else "",
                "solicitada": sol,
                "producida": prodq,
                "faltante": fal,
                "cap": num(r[i_cap]) if i_cap >= 0 else 0.0,
                "prioridad": str(r[i_pri] or "").strip() if i_pri >= 0 else "",
                "fecha_obj": as_date(r[i_fec]) if i_fec >= 0 else None,
                "especial": especial,
            }
        )
    return out


def leer_proyeccion(wb) -> Dict[str, dict]:
    if "Proyeccion" not in wb.sheetnames:
        return {}
    sh = wb["Proyeccion"]
    rows = list(sh.iter_rows(values_only=True))
    out = {}
    for r in rows[2:]:
        modelo = str(r[1] or "").strip()
        if not modelo or modelo.upper() == "TOTAL":
            continue
        out[modelo] = {
            "fecha_obj": as_date(r[2]) if not isinstance(r[2], str) else None,
            "meta": num(r[3]),
            "sin_prog": num(r[16]) if len(r) > 16 and r[16] not in ("--", None, "") else 0.0,
            "termino": as_date(r[17]) if len(r) > 17 else None,
            "estado": str(r[18] or "") if len(r) > 18 else "",
        }
    return out


def leer_almacen(wb) -> Dict[str, dict]:
    name = "Entrada de Almacen Modelo"
    if name not in wb.sheetnames:
        return {}
    sh = wb[name]
    rows = list(sh.iter_rows(values_only=True))
    out = {}
    for r in rows[2:]:
        modelo = str(r[2] or "").strip()
        if not modelo:
            continue
        out[modelo] = {
            "cantidad": num(r[3]),
            "salida_costura": as_date(r[4]),
            "entrada_alm": as_date(r[5]),
        }
    return out


def leer_eventos_tablero(wb) -> List[EventoLinea]:
    eventos: List[EventoLinea] = []
    seq = 0
    for name in HOJAS_SEMANA:
        if name not in wb.sheetnames:
            continue
        sh = wb[name]
        rows = list(sh.iter_rows(max_row=80, max_col=12, values_only=True))
        if len(rows) < 4:
            continue
        fechas = [as_date(v) for v in rows[1][5:10]]
        linea = None
        for r in rows[3:]:
            if r[1] not in (None, ""):
                try:
                    linea = str(int(float(r[1])))
                except (TypeError, ValueError):
                    pass
            modelo = str(r[2] or "").strip()
            if not modelo:
                continue
            up = modelo.upper()
            if up == "MODELO" or "ALERTA" in up or up.startswith("TOTAL"):
                break
            if not linea:
                continue
            for i, d in enumerate(fechas):
                q = num(r[5 + i]) if d else 0.0
                if q > 0 and d:
                    eventos.append(
                        EventoLinea(
                            fecha=d,
                            linea=linea,
                            modelo=modelo,
                            qty=q,
                            especial=es_especial_nombre(modelo),
                            seq=seq,
                        )
                    )
                    seq += 1
    eventos.sort(key=lambda e: (e.linea, e.fecha, e.seq))
    return eventos


def agrupar_modelos(filas: Iterable[dict], proy=None, alm=None) -> Dict[str, ModeloInfo]:
    proy = proy or {}
    alm = alm or {}
    out: Dict[str, ModeloInfo] = {}
    for r in filas:
        m = r["modelo"]
        if m not in out:
            out[m] = ModeloInfo(
                modelo=m,
                lote_nuevo=es_lote_nuevo(m),
                especial=bool(r.get("especial")),
            )
        x = out[m]
        x.faltante += r["faltante"]
        x.solicitada += r["solicitada"]
        x.producida += r["producida"]
        x.cap = max(x.cap, r["cap"])
        if r.get("lineas"):
            x.lineas = r["lineas"] if not x.lineas else x.lineas
        if r.get("prioridad"):
            x.prioridad = r["prioridad"]
        if r.get("fecha_obj") and (x.fecha_obj is None or r["fecha_obj"] < x.fecha_obj):
            x.fecha_obj = r["fecha_obj"]
        x.skus += 1
    for m, x in out.items():
        p = proy.get(m) or {}
        a = alm.get(m) or {}
        x.termino_plan = p.get("termino")
        x.estado_plan = p.get("estado") or ""
        x.sin_programar = p.get("sin_prog") or 0.0
        x.entrada_plan = a.get("entrada_alm")
        if not x.fecha_obj:
            x.fecha_obj = p.get("fecha_obj")
        if x.cap <= 0:
            x.cap = CAP_FALLBACK["1"]
    return out


def cargar_plan(path: str) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ph = leer_por_hacer(wb, False)
        pe = leer_por_hacer(wb, True)
        proy = leer_proyeccion(wb)
        alm = leer_almacen(wb)
        eventos = leer_eventos_tablero(wb)
    finally:
        wb.close()
    modelos = agrupar_modelos(ph, proy, alm)
    especiales = agrupar_modelos(pe, proy, alm)
    return {
        "path": path,
        "modelos": modelos,
        "especiales": especiales,
        "eventos": eventos,
        "proy": proy,
        "alm": alm,
    }


def _cap_modelo(modelos: Dict[str, ModeloInfo], nombre: str, linea: str) -> float:
    m = modelos.get(nombre)
    if m and m.cap > 0:
        return m.cap
    return CAP_FALLBACK.get(str(linea), 130.0)


def _construir_colas(eventos: Sequence[EventoLinea]) -> Dict[str, deque]:
    """Cola cronológica por línea (fragmentos del tablero, sin fusionar modelos)."""
    colas: Dict[str, deque] = {lin: deque() for lin in "12345"}
    regs = [e for e in eventos if not e.especial]
    regs.sort(key=lambda e: (e.linea, e.fecha, e.seq if e.seq else 0))
    for e in regs:
        colas[e.linea].append({"modelo": e.modelo, "qty": e.qty, "fecha_orig": e.fecha})
    return colas


def _lineas_modelo(s: str) -> List[str]:
    out = []
    for part in str(s or "").replace(";", ",").split(","):
        t = part.strip()
        if t.isdigit() and t in "12345":
            out.append(t)
    return out


def asignar_cupo_lotes(idle_dia: Dict[str, float], lotes: Sequence[ModeloInfo]) -> dict:
    """Reparte el cupo diurno que deja el nocturno entre lotes nuevos de la misma línea."""
    by_line: Dict[str, List[ModeloInfo]] = defaultdict(list)
    for m in lotes:
        lins = _lineas_modelo(m.lineas)
        lin = lins[0] if lins else ""
        by_line[lin].append(m)
    rows = []
    usados: Dict[str, float] = defaultdict(float)
    for lin, ms in by_line.items():
        cupo = float(idle_dia.get(lin) or 0.0)
        tot = sum(x.faltante for x in ms) or 1.0
        for m in ms:
            asg = cupo * (m.faltante / tot) if lin else 0.0
            rows.append(
                {
                    "modelo": m.modelo,
                    "lineas": m.lineas,
                    "faltante": round(m.faltante, 0),
                    "cap": m.cap,
                    "prioridad": m.prioridad,
                    "cupo_linea": round(cupo, 1),
                    "asignado": round(asg, 1),
                    "hueco": round(max(0.0, m.faltante - asg), 1),
                    "cobertura": round(100.0 * asg / m.faltante, 1) if m.faltante else 0.0,
                    "dias_cupo": round(asg / m.cap, 1) if m.cap else None,
                }
            )
            usados[lin] += asg
    rows.sort(key=lambda r: (-r["faltante"], r["modelo"]))
    lineas = []
    for lin in "12345":
        ms = by_line.get(lin) or []
        idle = float(idle_dia.get(lin) or 0.0)
        need = sum(m.faltante for m in ms)
        asg = usados.get(lin) or 0.0
        lineas.append(
            {
                "linea": lin,
                "idle": round(idle, 1),
                "n_lotes": len(ms),
                "faltante_lotes": round(need, 0),
                "asignado": round(asg, 1),
                "hueco": round(max(0.0, need - asg), 1),
                "cobertura": round(100.0 * asg / need, 1) if need else None,
            }
        )
    total_falt = sum(m.faltante for m in lotes)
    total_asg = sum(r["asignado"] for r in rows)
    return {
        "por_modelo": rows,
        "por_linea": lineas,
        "faltante_total": round(total_falt, 0),
        "cupo_total": round(sum(float(v) for v in (idle_dia or {}).values()), 1),
        "cupo_usable": round(total_asg, 1),
        "hueco_total": round(max(0.0, total_falt - total_asg), 1),
        "cobertura": round(100.0 * total_asg / total_falt, 1) if total_falt else 0.0,
    }


def _agregar_remanente(colas: Dict[str, deque], eventos: Sequence[EventoLinea], modelos: Dict[str, ModeloInfo]) -> float:
    """El faltante que no está en el tablero sigue al final de la cola de su línea."""
    prog = defaultdict(lambda: defaultdict(float))
    for e in eventos:
        if e.especial:
            continue
        prog[e.modelo][e.linea] += e.qty
    total = 0.0
    for nom, info in modelos.items():
        if info.especial:
            continue
        hecho = sum(prog[nom].values())
        rest = info.faltante - hecho
        if rest <= 0.001:
            continue
        shares = dict(prog[nom])
        if not shares:
            lins = _lineas_modelo(info.lineas)
            if not lins:
                continue
            each = rest / len(lins)
            for lin in lins:
                colas[lin].append({"modelo": nom, "qty": each, "fecha_orig": None, "remanente": True})
                total += each
            continue
        tot = sum(shares.values()) or 1.0
        for lin, q in shares.items():
            parte = rest * (q / tot)
            if parte > 0.001:
                colas[lin].append({"modelo": nom, "qty": parte, "fecha_orig": None, "remanente": True})
                total += parte
    return total


def _capacidad_diaria_regular(eventos: Sequence[EventoLinea]) -> Dict[Tuple[str, date], float]:
    cap: Dict[Tuple[str, date], float] = defaultdict(float)
    for e in eventos:
        if e.especial:
            continue
        cap[(e.linea, e.fecha)] += e.qty
    return cap


def semana_plan(d: date, lunes_base: date = LUNES_BASE) -> int:
    return ((d - lunes_base).days // 7) + 1


def _consumir(cola: deque, amount: float) -> List[Tuple[str, float, bool]]:
    hecho: List[Tuple[str, float, bool]] = []
    left = amount
    while left > 0.001 and cola:
        c = cola[0]
        take = min(c["qty"], left)
        if take > 0:
            hecho.append((c["modelo"], take, bool(c.get("remanente"))))
            c["qty"] -= take
            left -= take
        if c["qty"] <= 0.001:
            cola.popleft()
    return hecho


def _acum_hasta(timeline: Dict[str, List[Tuple[date, float]]], corte: date) -> Dict[str, float]:
    out: Dict[str, float] = defaultdict(float)
    for nom, pares in timeline.items():
        for d, q in pares:
            if d <= corte:
                out[nom] += q
    return out


def _primer_ultimo(timeline: Dict[str, List[Tuple[date, float]]]) -> Dict[str, Tuple[Optional[date], Optional[date]]]:
    out = {}
    for nom, pares in timeline.items():
        if not pares:
            out[nom] = (None, None)
            continue
        ds = [d for d, q in pares if q > 0]
        out[nom] = (min(ds) if ds else None, max(ds) if ds else None)
    return out


def timeline_base(eventos: Sequence[EventoLinea]) -> Dict[str, List[Tuple[date, float]]]:
    tl: Dict[str, List[Tuple[date, float]]] = defaultdict(list)
    for e in eventos:
        if e.especial or e.qty <= 0:
            continue
        tl[e.modelo].append((e.fecha, e.qty))
    return tl


def simular_nocturnos(
    plan: dict,
    n_semanas: int,
    lunes_base: date = LUNES_BASE,
    corte: date = CORTE_ALMACEN,
    corte_costura: date = CORTE_COSTURA_ALMACEN,
    corte_dic: date = CORTE_COSTURA_DICIEMBRE,
) -> dict:
    modelos: Dict[str, ModeloInfo] = plan["modelos"]
    eventos: List[EventoLinea] = plan["eventos"]
    noches = fechas_nocturno(n_semanas, lunes_base)
    noche_set = set(noches)

    colas = _construir_colas(eventos)
    remanente_total = _agregar_remanente(colas, eventos, modelos)
    cap_dia = _capacidad_diaria_regular(eventos)
    fechas = sorted({e.fecha for e in eventos if not e.especial})
    if noches:
        fechas = sorted(set(fechas) | set(noches))

    pzas_noct: Dict[str, float] = defaultdict(float)
    timeline: Dict[str, List[Tuple[date, float]]] = defaultdict(list)
    extra_linea: Dict[str, float] = defaultdict(float)
    idle_noche: Dict[str, float] = defaultdict(float)
    idle_dia: Dict[str, float] = defaultdict(float)
    noches_usadas_linea: Dict[str, int] = defaultdict(int)
    noches_idle_linea: Dict[str, int] = defaultdict(int)
    remanente_usado = 0.0
    celdas: Dict[Tuple, dict] = {}

    def registrar(d: date, lin: str, nocturno: bool, nom: str, q: float, rem: bool) -> None:
        nonlocal remanente_usado
        turno = "noche" if nocturno else "dia"
        key = (d, lin, turno, nom, bool(rem))
        if key not in celdas:
            celdas[key] = {
                "fecha": d,
                "semana": semana_plan(d, lunes_base),
                "linea": lin,
                "turno": turno,
                "modelo": nom,
                "qty": 0.0,
                "remanente": False,
            }
        celdas[key]["qty"] += q
        if rem:
            celdas[key]["remanente"] = True
            remanente_usado += q

    def aplicar(hechos: List[Tuple[str, float, bool]], d: date, nocturno: bool, lin: str) -> float:
        total = 0.0
        for nom, q, rem in hechos:
            if q <= 0:
                continue
            total += q
            timeline[nom].append((d, q))
            if nocturno:
                pzas_noct[nom] += q
            registrar(d, lin, nocturno, nom, q, rem)
        return total

    for d in fechas:
        for lin in "12345":
            qty_dia = cap_dia.get((lin, d), 0.0)
            if qty_dia > 0:
                hechos = _consumir(colas[lin], qty_dia)
                usado = aplicar(hechos, d, False, lin)
                if usado + 0.001 < qty_dia:
                    idle_dia[lin] += qty_dia - usado
        if d in noche_set:
            for lin in ("2", "3", "4", "5"):
                if not colas[lin]:
                    idle_noche[lin] += cap_nocturna(CAP_FALLBACK[lin])
                    noches_idle_linea[lin] += 1
                    continue
                front = colas[lin][0]["modelo"]
                extra = cap_nocturna(_cap_modelo(modelos, front, lin))
                usados = aplicar(_consumir(colas[lin], extra), d, True, lin)
                extra_linea[lin] += usados
                if usados > 0.001:
                    noches_usadas_linea[lin] += 1
                else:
                    noches_idle_linea[lin] += 1
                    idle_noche[lin] += extra
                if usados + 0.001 < extra:
                    idle_noche[lin] += extra - usados
            # L1 noche: 100% del cupo nocturno sobre la cola de L2.
            if colas["2"]:
                extra1 = cap_nocturna(_cap_modelo(modelos, colas["2"][0]["modelo"], "2"))
                usados1 = aplicar(_consumir(colas["2"], extra1), d, True, "1")
                extra_linea["1"] += usados1
                if usados1 > 0.001:
                    noches_usadas_linea["1"] += 1
                else:
                    noches_idle_linea["1"] += 1
                    idle_noche["1"] += extra1
                if usados1 + 0.001 < extra1:
                    idle_noche["1"] += extra1 - usados1
            else:
                idle_noche["1"] += cap_nocturna(CAP_FALLBACK["2"])
                noches_idle_linea["1"] += 1

    restantes_cola = {lin: sum(c["qty"] for c in colas[lin]) for lin in "12345"}
    tl_base = timeline_base(eventos)
    pu_base = _primer_ultimo(tl_base)
    pu_esc = _primer_ultimo(timeline)
    alm16_base = _acum_hasta(tl_base, corte_costura)
    alm16_esc = _acum_hasta(timeline, corte_costura)
    dic_base = _acum_hasta(tl_base, corte_dic)
    dic_esc = _acum_hasta(timeline, corte_dic)
    prog = defaultdict(float)
    for e in eventos:
        if not e.especial:
            prog[e.modelo] += e.qty

    resultados: List[ResultadoModelo] = []
    for nom, info in modelos.items():
        t_base = pu_base.get(nom, (None, None))[1] or info.termino_plan
        e_base = info.entrada_plan or (add_business_days(t_base) if t_base else None)
        t_esc = pu_esc.get(nom, (None, None))[1] or t_base
        leftover = 0.0
        for lin, cola in colas.items():
            leftover += sum(c["qty"] for c in cola if c["modelo"] == nom)
        if leftover > 0.001:
            cap = info.cap or 130.0
            extra_days = int((leftover + cap - 1) // cap)
            last = fechas[-1] if fechas else LUNES_BASE
            t_esc = last
            added = 0
            cur = last
            while added < extra_days:
                cur += timedelta(days=1)
                if cur.weekday() < 5:
                    added += 1
                    t_esc = cur
        e_esc = add_business_days(t_esc) if t_esc else None
        dg = delta_dias(t_base, t_esc) if t_base and t_esc else None
        resultados.append(
            ResultadoModelo(
                modelo=nom,
                faltante=info.faltante,
                cap=info.cap,
                lineas=info.lineas,
                prioridad=info.prioridad,
                lote_nuevo=info.lote_nuevo,
                fecha_obj=info.fecha_obj,
                termino_base=t_base,
                entrada_base=e_base,
                termino_esc=t_esc,
                entrada_esc=e_esc,
                primer_base=pu_base.get(nom, (None, None))[0],
                primer_esc=pu_esc.get(nom, (None, None))[0],
                pzas_nocturno=round(pzas_noct.get(nom, 0.0), 1),
                pzas_prog=round(prog.get(nom, 0.0), 1),
                pzas_alm16_base=round(alm16_base.get(nom, 0.0), 1),
                pzas_alm16_esc=round(alm16_esc.get(nom, 0.0), 1),
                pzas_dic_base=round(dic_base.get(nom, 0.0), 1),
                pzas_dic_esc=round(dic_esc.get(nom, 0.0), 1),
                dias_ganados=dg,
                a_tiempo_base=(t_base <= info.fecha_obj) if t_base and info.fecha_obj else None,
                a_tiempo_esc=(t_esc <= info.fecha_obj) if t_esc and info.fecha_obj else None,
                cubre_almacen_base=(e_base <= corte) if e_base else None,
                cubre_almacen_esc=(e_esc <= corte) if e_esc else None,
            )
        )
    resultados.sort(key=lambda r: (r.termino_esc or date.max, r.modelo))

    n_noches = len(noches)
    cal = sorted(celdas.values(), key=lambda c: (c["fecha"], c["linea"], c["turno"], c["modelo"]))
    return {
        "semanas": n_semanas,
        "noches": noches,
        "n_noches": n_noches,
        "noches_teoricas": n_semanas * 3,
        "noches_perdidas_cobro": n_semanas * 3 - n_noches,
        "costo_usd": costo_bono(n_noches),
        "personas": personas_total(),
        "pzas_nocturno": round(sum(pzas_noct.values()), 1),
        "pzas_por_linea": {k: round(v, 1) for k, v in extra_linea.items()},
        "idle_noche": {k: round(v, 1) for k, v in idle_noche.items()},
        "idle_dia": {k: round(v, 1) for k, v in idle_dia.items()},
        "remanente_total": round(remanente_total, 1),
        "remanente_usado": round(remanente_usado, 1),
        "noches_usadas_linea": dict(noches_usadas_linea),
        "noches_idle_linea": dict(noches_idle_linea),
        "restante_lineas": restantes_cola,
        "resultados": resultados,
        "calendario": cal,
        "corte": corte,
        "corte_costura": corte_costura,
        "corte_dic": corte_dic,
    }


def kpis_escenario(sim: dict) -> dict:
    rows: List[ResultadoModelo] = sim["resultados"]
    lotes = [r for r in rows if r.lote_nuevo]
    urgentes = [r for r in rows if "urgente" in (r.prioridad or "").lower()]

    def _n(pred):
        return sum(1 for r in rows if pred(r))

    extra_alm16 = round(sum(r.pzas_alm16_esc - r.pzas_alm16_base for r in rows), 1)
    extra_dic = round(sum(r.pzas_dic_esc - r.pzas_dic_base for r in rows), 1)
    extra_alm16_lotes = round(sum(r.pzas_alm16_esc - r.pzas_alm16_base for r in lotes), 1)
    extra_dic_lotes = round(sum(r.pzas_dic_esc - r.pzas_dic_base for r in lotes), 1)
    return {
        "n_noches": sim["n_noches"],
        "noches_teoricas": sim["noches_teoricas"],
        "noches_perdidas_cobro": sim["noches_perdidas_cobro"],
        "costo_usd": sim["costo_usd"],
        "pzas_nocturno": sim["pzas_nocturno"],
        "usd_por_pza": round(sim["costo_usd"] / sim["pzas_nocturno"], 3) if sim["pzas_nocturno"] else None,
        "modelos": len(rows),
        "lotes_nuevos": len(lotes),
        "pzas_lotes_nuevos": round(sum(r.faltante for r in lotes), 0),
        "cubre_almacen_base": _n(lambda r: r.cubre_almacen_base is True),
        "cubre_almacen_esc": _n(lambda r: r.cubre_almacen_esc is True),
        "lotes_cubre_base": sum(1 for r in lotes if r.cubre_almacen_base is True),
        "lotes_cubre_esc": sum(1 for r in lotes if r.cubre_almacen_esc is True),
        "a_tiempo_obj_base": _n(lambda r: r.a_tiempo_base is True),
        "a_tiempo_obj_esc": _n(lambda r: r.a_tiempo_esc is True),
        "dias_ganados_mediana": _mediana([r.dias_ganados for r in rows if r.dias_ganados is not None]),
        "dias_ganados_lotes": _mediana([r.dias_ganados for r in lotes if r.dias_ganados is not None]),
        "max_entrada_base": max((r.entrada_base for r in rows if r.entrada_base), default=None),
        "max_entrada_esc": max((r.entrada_esc for r in rows if r.entrada_esc), default=None),
        "max_entrada_lotes_base": max((r.entrada_base for r in lotes if r.entrada_base), default=None),
        "max_entrada_lotes_esc": max((r.entrada_esc for r in lotes if r.entrada_esc), default=None),
        "urgentes": len(urgentes),
        "pzas_por_linea": sim["pzas_por_linea"],
        "idle_noche": sim.get("idle_noche") or {},
        "idle_dia": sim.get("idle_dia") or {},
        "pzas_alm16_base": round(sum(r.pzas_alm16_base for r in rows), 1),
        "pzas_alm16_esc": round(sum(r.pzas_alm16_esc for r in rows), 1),
        "pzas_dic_base": round(sum(r.pzas_dic_base for r in rows), 1),
        "pzas_dic_esc": round(sum(r.pzas_dic_esc for r in rows), 1),
        "extra_alm16": extra_alm16,
        "extra_dic": extra_dic,
        "extra_alm16_lotes": extra_alm16_lotes,
        "extra_dic_lotes": extra_dic_lotes,
        "pzas_lotes_alm16_esc": round(sum(r.pzas_alm16_esc for r in lotes), 1),
        "pzas_lotes_dic_esc": round(sum(r.pzas_dic_esc for r in lotes), 1),
        "modelos_adelantados": sum(1 for r in rows if (r.dias_ganados or 0) > 0),
        "lotes_adelantados": sum(1 for r in lotes if (r.dias_ganados or 0) > 0),
        "remanente_total": sim.get("remanente_total") or 0.0,
        "remanente_usado": sim.get("remanente_usado") or 0.0,
    }


def celdas_a_dict(celdas: Sequence[dict]) -> List[dict]:
    out = []
    for c in celdas:
        d = c["fecha"]
        out.append(
            {
                "fecha": fmt_fecha(d),
                "fecha_iso": d.isoformat() if hasattr(d, "isoformat") else str(d),
                "semana": int(c["semana"]),
                "dow": d.weekday() if hasattr(d, "weekday") else 0,
                "linea": str(c["linea"]),
                "turno": c["turno"],
                "modelo": c["modelo"],
                "qty": round(float(c["qty"]), 1),
                "remanente": bool(c.get("remanente")),
            }
        )
    return out


def etiquetas_semana(lunes_base: date = LUNES_BASE, n: int = 12) -> List[dict]:
    out = []
    for i in range(1, n + 1):
        ini = lunes_base + timedelta(days=7 * (i - 1))
        fin = ini + timedelta(days=4)
        out.append(
            {
                "semana": i,
                "lunes": fmt_fecha(ini),
                "viernes": fmt_fecha(fin),
                "lunes_iso": ini.isoformat(),
                "label": f"Semana {i} · {ini.strftime('%d/%m')}–{fin.strftime('%d/%m')}",
            }
        )
    return out


def resumen_semanas(celdas: Sequence[dict]) -> List[dict]:
    weeks: Dict[int, dict] = {}
    for c in celdas:
        w = int(c["semana"])
        if w not in weeks:
            weeks[w] = {
                "semana": w,
                "pzas_dia": 0.0,
                "pzas_noche": 0.0,
                "pzas_remanente": 0.0,
                "modelos": set(),
            }
        if c["turno"] == "noche":
            weeks[w]["pzas_noche"] += c["qty"]
        else:
            weeks[w]["pzas_dia"] += c["qty"]
        if c.get("remanente"):
            weeks[w]["pzas_remanente"] += c["qty"]
        weeks[w]["modelos"].add(c["modelo"])
    out = []
    for w in sorted(weeks):
        x = weeks[w]
        out.append(
            {
                "semana": w,
                "pzas_dia": round(x["pzas_dia"], 1),
                "pzas_noche": round(x["pzas_noche"], 1),
                "pzas_remanente": round(x["pzas_remanente"], 1),
                "pzas": round(x["pzas_dia"] + x["pzas_noche"], 1),
                "n_modelos": len(x["modelos"]),
            }
        )
    return out


def _mediana(vals: List[int]) -> Optional[float]:
    if not vals:
        return None
    s = sorted(vals)
    n = len(s)
    if n % 2:
        return float(s[n // 2])
    return (s[n // 2 - 1] + s[n // 2]) / 2.0


def comparar_lotes(plan_actual: dict, plan_nuevo: dict) -> List[ModeloInfo]:
    ma = set(plan_actual["modelos"])
    out = []
    for nom, info in plan_nuevo["modelos"].items():
        if nom not in ma or info.lote_nuevo:
            out.append(info)
    out.sort(key=lambda x: (x.termino_plan or date.max, x.modelo))
    return out


def modelos_a_dict(rows: Sequence[ResultadoModelo]) -> List[dict]:
    out = []
    for r in rows:
        out.append(
            {
                "modelo": r.modelo,
                "faltante": r.faltante,
                "cap": r.cap,
                "lineas": r.lineas,
                "prioridad": r.prioridad,
                "lote_nuevo": r.lote_nuevo,
                "fecha_obj": fmt_fecha(r.fecha_obj),
                "termino_base": fmt_fecha(r.termino_base),
                "entrada_base": fmt_fecha(r.entrada_base),
                "termino_esc": fmt_fecha(r.termino_esc),
                "entrada_esc": fmt_fecha(r.entrada_esc),
                "primer_base": fmt_fecha(r.primer_base),
                "primer_esc": fmt_fecha(r.primer_esc),
                "pzas_nocturno": r.pzas_nocturno,
                "pzas_prog": r.pzas_prog,
                "pzas_alm16_base": r.pzas_alm16_base,
                "pzas_alm16_esc": r.pzas_alm16_esc,
                "pzas_dic_base": r.pzas_dic_base,
                "pzas_dic_esc": r.pzas_dic_esc,
                "dias_ganados": r.dias_ganados,
                "a_tiempo_base": r.a_tiempo_base,
                "a_tiempo_esc": r.a_tiempo_esc,
                "cubre_almacen_base": r.cubre_almacen_base,
                "cubre_almacen_esc": r.cubre_almacen_esc,
                "termino_esc_iso": r.termino_esc.isoformat() if r.termino_esc else "9999-12-31",
                "entrada_esc_iso": r.entrada_esc.isoformat() if r.entrada_esc else "9999-12-31",
            }
        )
    return out
