#!/usr/bin/env python3
"""Genera proyección de producción LITE PANT DAMA (XS–XL).

Referencia: ventas DAMA de BASIC LINE PANT en Dashboard Basic Line.
Consumo tela/insumos: Ficha técnica LITE PANT DAMA (VIORI).
"""

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
HTML_PATHS = [BASE / "Dashboard_Basic_Line.html"]
OUTPUT_PATH = BASE / "LITE_PANT_RANGO_PRODUCCION.xlsx"

PRODUCTO = "LITE PANT"
REFERENCE_MODEL = "BASIC LINE PANT"
GENERO = "DAMA"
TALLAS = ["XS", "S", "M", "L", "XL"]

# Normalización nombres de tienda (dashboard Basic Line → convención proyección)
STORE_MAP = {
    "SAMBIL VALENCIA": "SAMBIL",
    "SAMBIL CHACAO": "CHACAO",
    "GRANDPLAZ": "GRAND PLAZ",
    "VELA": "LA VELA",
    "GRIE": "GRIETA",
}

# ── Parámetros de proyección ──
HIGH_SEASON_FACTOR = 1.25
SAFETY_STOCK_PCT = 0.15
TELA_SS_PCT = 0.20             # Stock de seguridad tela +20% sobre consumo
COVER_MONTHS_MIN = 2.5
COVER_MONTHS_MAX = 3.5
MAX_PCT_ABOVE_MIN = 1.12
TOLON_VS_CHACAO = 0.85          # Tolón ≈ 85% de Chacao
WEB_VS_CERRO_VERDE = 0.60       # Web ≈ 60% de Cerro Verde
GRAND_PLAZ_BONUS = 1.68         # Grand Plaz +68% hist. (+40% base + 20% adicional)
VELOCITY_MONTHS = ["junio-2026", "julio-2026", "agosto-2026"]

# ── Colores de producción ──
COLORES = ["Negro", "Vinotinto", "Verde Militar"]
COLOR_PCT = {"Negro": 0.40, "Vinotinto": 0.30, "Verde Militar": 0.30}
COLOR_FILL = {
    "Negro": PatternFill("solid", fgColor="424242"),
    "Vinotinto": PatternFill("solid", fgColor="8B2942"),
    "Verde Militar": PatternFill("solid", fgColor="4B5320"),
}
COLOR_FONT = {"Negro": "FFFFFF", "Vinotinto": "FFFFFF", "Verde Militar": "FFFFFF"}

# Estilo curva completa (formato SPOTS)
CURVA_HDR_FILL = PatternFill("solid", fgColor="1F3864")
CURVA_DATA_FILL = PatternFill("solid", fgColor="D6E4BC")
CURVA_TOT_FILL = PatternFill("solid", fgColor="A9D08E")
CURVA_HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
CURVA_NUM_FONT = Font(bold=True, color="375623", size=11)
CURVA_LBL_FONT = Font(bold=True, size=10)

# ── Consumo tela VIORI por talla (ficha técnica LITE PANT DAMA) ──
CONSUMO_MTS = {"XS": 1.45, "S": 1.48, "M": 1.58, "L": 1.63, "XL": 1.66}
CONSUMO_KG = {"XS": 0.165, "S": 0.172, "M": 0.177, "L": 0.187, "XL": 0.206}
ELASTICA_CM = {"XS": 67, "S": 69, "M": 72, "L": 75, "XL": 77}
ETIQUETAS_POR_PIEZA = 1

# ── Tiendas ──
HISTORICAL_STORES = ["SAMBIL", "GRIETA", "CERRO VERDE", "CHACAO", "GRAND PLAZ", "LA VELA"]
NEW_STORE = "BARQUISIMETO"
WEB_STORE = "WEB"
TOLON_STORE = "TOLON"
EXCLUDE_STORES = {"CORPORATIVO", "PEDIDOS"}
DISTRIBUTION_STORES = HISTORICAL_STORES + [TOLON_STORE, WEB_STORE, NEW_STORE]

# ── Estilos Excel ──
title_fill = PatternFill("solid", fgColor="1E3A5F")
sub_fill = PatternFill("solid", fgColor="4A7FB5")
color_fill = PatternFill("solid", fgColor="D6E4F0")
tot_fill = PatternFill("solid", fgColor="A8C8E8")
min_fill = PatternFill("solid", fgColor="FFF9C4")
max_fill = PatternFill("solid", fgColor="FFE0B2")
warn_fill = PatternFill("solid", fgColor="FFCDD2")
white_fill = PatternFill("solid", fgColor="FFFFFF")
adj_fill = PatternFill("solid", fgColor="FFF3E0")
thin = Side(style="thin", color="1E3A5F")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def load_combined_data() -> dict:
    all_rows = []
    meses_order = []
    for path in HTML_PATHS:
        html = path.read_text(encoding="utf-8")
        match = re.search(r"var DATA=(\{.*?\});", html, re.DOTALL)
        if not match:
            raise ValueError(f"No se encontró DATA en {path.name}")
        data = json.loads(match.group(1))
        meses_order.extend(data.get("meses_order", []))
        for row in data["raw_rows"]:
            row = dict(row)
            row["tienda"] = STORE_MAP.get(row.get("tienda", ""), row.get("tienda", ""))
            all_rows.append(row)

    seen = set()
    months_chrono = []
    for m in meses_order:
        if m not in seen:
            seen.add(m)
            months_chrono.append(m)

    return {"raw_rows": all_rows, "meses_order": months_chrono}


def analyze_reference(data: dict) -> dict:
    rows = [
        r for r in data["raw_rows"]
        if r.get("modelo") == REFERENCE_MODEL
        and r["genero"] == GENERO
        and r["tienda"] not in EXCLUDE_STORES
    ]
    months = data["meses_order"]
    months_with_data = [m for m in months if any(r["mes"] == m for r in rows)]
    n_months = len(months_with_data)

    by_month = {m: sum(r["v"] for r in rows if r["mes"] == m) for m in months}
    vel_base = sum(by_month.get(m, 0) for m in VELOCITY_MONTHS) / len(VELOCITY_MONTHS)
    vel_all = sum(by_month.values()) / n_months if n_months else 0
    dec_vel = by_month.get("diciembre-2025", 0)

    talla_tot = {t: 0 for t in TALLAS}
    for r in rows:
        if r["talla"] in talla_tot:
            talla_tot[r["talla"]] += r["v"]
    talla_grand = sum(talla_tot.values())
    talla_pct = {t: talla_tot[t] / talla_grand if talla_grand else 0 for t in TALLAS}

    store_tot = {}
    for s in set(r["tienda"] for r in rows):
        store_tot[s] = sum(r["v"] for r in rows if r["tienda"] == s)
    store_monthly_hist = {s: store_tot.get(s, 0) / n_months for s in store_tot}

    sambil_m = store_monthly_hist.get("SAMBIL", 0)
    chacao_m = store_monthly_hist.get("CHACAO", 0)
    cerro_m = store_monthly_hist.get("CERRO VERDE", 0)
    grieta_m = store_monthly_hist.get("GRIETA", 0)
    grand_hist = store_monthly_hist.get("GRAND PLAZ", 0)
    vela_hist = store_monthly_hist.get("LA VELA", 0)

    tolon_proj = chacao_m * TOLON_VS_CHACAO
    web_proj = cerro_m * WEB_VS_CERRO_VERDE
    barq_proj = (grieta_m + chacao_m + tolon_proj) / 3
    vela_proj = (sambil_m + cerro_m) / 2
    grand_proj = grand_hist * GRAND_PLAZ_BONUS

    store_monthly = {}
    for s in HISTORICAL_STORES:
        if s == "GRAND PLAZ":
            store_monthly[s] = grand_proj
        elif s == "LA VELA":
            store_monthly[s] = vela_proj
        else:
            store_monthly[s] = store_monthly_hist.get(s, 0)
    store_monthly[TOLON_STORE] = tolon_proj
    store_monthly[WEB_STORE] = web_proj
    store_monthly[NEW_STORE] = barq_proj

    total_weight = sum(store_monthly.values())
    store_share = {s: store_monthly[s] / total_weight if total_weight else 0 for s in DISTRIBUTION_STORES}

    return {
        "vel_base": vel_base,
        "vel_all": vel_all,
        "dec_vel": dec_vel,
        "vel_months": VELOCITY_MONTHS,
        "talla_pct": talla_pct,
        "talla_tot": talla_tot,
        "store_monthly": store_monthly,
        "store_monthly_hist": store_monthly_hist,
        "store_share": store_share,
        "tolon_proj": tolon_proj,
        "web_proj": web_proj,
        "barq_proj": barq_proj,
        "vela_proj": vela_proj,
        "grand_proj": grand_proj,
        "grand_hist": grand_hist,
        "vela_hist": vela_hist,
        "sambil_m": sambil_m,
        "chacao_m": chacao_m,
        "cerro_m": cerro_m,
        "months": months,
        "by_month": by_month,
        "sources": [p.name for p in HTML_PATHS],
    }


def calc_production(ref: dict) -> dict:
    vel_adj = ref["vel_base"] * HIGH_SEASON_FACTOR
    barq_add = ref["barq_proj"]
    vel_network = vel_adj + barq_add

    raw_min = round(vel_network * COVER_MONTHS_MIN * (1 + SAFETY_STOCK_PCT))
    raw_max = round(vel_network * COVER_MONTHS_MAX * (1 + SAFETY_STOCK_PCT))

    return {
        "vel_adj": vel_adj,
        "vel_network": vel_network,
        "barq_add": barq_add,
        "raw_min": raw_min,
        "raw_max": raw_max,
        "prod_min": raw_min,
        "prod_max": raw_max,
    }


def distribute_by_talla(total: int, talla_pct: dict) -> dict:
    result = {}
    allocated = 0
    for t in TALLAS[:-1]:
        qty = round(total * talla_pct[t])
        result[t] = qty
        allocated += qty
    result[TALLAS[-1]] = total - allocated
    return result


def distribute_by_store(total: int, ref: dict) -> dict:
    weights = ref["store_monthly"]
    total_w = sum(weights.values())
    store_units = {s: round(total * weights[s] / total_w) for s in DISTRIBUTION_STORES}
    diff = total - sum(store_units.values())
    if diff:
        top = max(DISTRIBUTION_STORES, key=lambda s: store_units[s])
        store_units[top] += diff
    return store_units


def distribute_store_talla(store_units: dict, talla_pct: dict) -> dict:
    return {store: distribute_by_talla(qty, talla_pct) for store, qty in store_units.items()}


def distribute_by_color_talla(total: int, talla_pct: dict) -> tuple[dict, dict]:
    """Distribuye unidades por color (40/30/30) y luego por talla dentro de cada color."""
    matrix = {}
    color_totals = {}
    allocated = 0
    for color in COLORES[:-1]:
        ct = round(total * COLOR_PCT[color])
        color_totals[color] = ct
        matrix[color] = distribute_by_talla(ct, talla_pct)
        allocated += ct
    last = COLORES[-1]
    color_totals[last] = total - allocated
    matrix[last] = distribute_by_talla(color_totals[last], talla_pct)
    return matrix, color_totals


def calc_tela_totals(color_matrix: dict) -> dict:
    """Calcula metros y kg de tela VIORI por color, con compra incluyendo SS."""
    ss_factor = 1 + TELA_SS_PCT
    out = {}
    for color in COLORES:
        mts = sum(color_matrix[color][t] * CONSUMO_MTS[t] for t in TALLAS)
        kg = sum(color_matrix[color][t] * CONSUMO_KG[t] for t in TALLAS)
        und = sum(color_matrix[color][t] for t in TALLAS)
        out[color] = {
            "und": und,
            "mts": round(mts, 2),
            "kg": round(kg, 3),
            "mts_compra": round(mts * ss_factor, 2),
            "kg_compra": round(kg * ss_factor, 3),
        }
    total_mts = sum(v["mts"] for v in out.values())
    total_kg = sum(v["kg"] for v in out.values())
    total_mts_compra = round(total_mts * ss_factor, 2)
    total_kg_compra = round(total_kg * ss_factor, 3)
    total_und = sum(v["und"] for v in out.values())
    return {
        "by_color": out,
        "total_mts": round(total_mts, 2),
        "total_kg": round(total_kg, 3),
        "total_mts_compra": total_mts_compra,
        "total_kg_compra": total_kg_compra,
        "total_und": total_und,
    }


def calc_insumos_totals(talla_qty: dict) -> dict:
    """Totales de insumos (elástica en metros, etiquetas en und)."""
    result = {"elastica_m": 0, "etiquetas": 0, "und": 0}
    for t in TALLAS:
        q = talla_qty[t]
        result["elastica_m"] += q * ELASTICA_CM[t] / 100
        result["etiquetas"] += q * ETIQUETAS_POR_PIEZA
        result["und"] += q
    result["elastica_m"] = round(result["elastica_m"], 2)
    return result


def style_cell(cell, fill=None, bold=False, align=center):
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)
    cell.alignment = align
    cell.border = border


def write_resumen(wb, ref, prod):
    ws = wb.create_sheet("Resumen", 0)
    rows = [
        [f"{PRODUCTO} — PROYECCIÓN DE PRODUCCIÓN (DAMA)"],
        [],
        ["Referencia analítica", f"{REFERENCE_MODEL} — ventas DAMA (Basic Line)"],
        ["Fuentes datos", " + ".join(ref["sources"])],
        ["Ficha técnica", "Ficha_tecnica_LITE_PANT.xlsx"],
        ["Período velocidad base", f"Jun–Jul–Ago 2026 ({', '.join(ref['vel_months'])})"],
        [],
        ["── VELOCIDAD ──"],
        ["Velocidad base (3m combinado)", round(ref["vel_base"], 1), "und/mes"],
        ["Velocidad histórica promedio", round(ref["vel_all"], 1), "und/mes"],
        ["Diciembre 2025 (temporada alta ref.)", ref["dec_vel"], "und"],
        ["Factor temporada alta aplicado", HIGH_SEASON_FACTOR, "×"],
        ["Velocidad ajustada temporada alta", round(prod["vel_adj"], 1), "und/mes"],
        ["Barquisimeto adicional (tienda nueva)", round(prod["barq_add"], 1), "und/mes"],
        ["Velocidad red completa ajustada", round(prod["vel_network"], 1), "und/mes"],
        [],
        ["── COBERTURA Y STOCK DE SEGURIDAD ──"],
        ["Meses cobertura mínimo", COVER_MONTHS_MIN],
        ["Meses cobertura máximo", COVER_MONTHS_MAX],
        ["Stock de seguridad", f"{int(SAFETY_STOCK_PCT * 100)}%"],
        ["Demanda teórica mín", prod["raw_min"], "und"],
        ["Demanda teórica máx", prod["raw_max"], "und"],
        [],
        ["── RANGO DE PRODUCCIÓN (CALCULADO) ──"],
        ["MÍNIMO (compromiso)", prod["prod_min"], "und"],
        ["MÁXIMO", prod["prod_max"], "und"],
        ["Rango de acción", prod["prod_max"] - prod["prod_min"], "und"],
        ["Base del rango", f"{COVER_MONTHS_MIN}–{COVER_MONTHS_MAX} meses + SS {int(SAFETY_STOCK_PCT*100)}%", ""],
        [],
        ["── COLORES DE PRODUCCIÓN ──"],
        ["Negro", f"{int(COLOR_PCT['Negro']*100)}%", "protagonista"],
        ["Vinotinto", f"{int(COLOR_PCT['Vinotinto']*100)}%", ""],
        ["Verde Militar", f"{int(COLOR_PCT['Verde Militar']*100)}%", ""],
        ["Curva completa color × talla", "Hoja Cantidades por Colores", ""],
        ["Detalle color/talla Mín-Máx", "Hoja Producción Color × Talla", ""],
        ["Compra tela VIORI", f"Hoja Compra de Tela (+{int(TELA_SS_PCT*100)}% SS)", ""],
        [],
        ["── AJUSTES DE TIENDA (DISTRIBUCIÓN) ──"],
        ["Tolón histórico", round(ref["store_monthly_hist"].get("TOLON", 0), 1), "und/mes"],
        [f"Tolón proyectado ({int(TOLON_VS_CHACAO*100)}% Chacao)", round(ref["tolon_proj"], 1), "und/mes"],
        ["Web histórica", round(ref["store_monthly_hist"].get("WEB", 0), 1), "und/mes"],
        [f"Web proyectada ({int(WEB_VS_CERRO_VERDE*100)}% Cerro Verde)", round(ref["web_proj"], 1), "und/mes"],
        ["La Vela histórica", round(ref["vela_hist"], 1), "und/mes"],
        ["La Vela proyectada (avg Sambil + Cerro Verde)", round(ref["vela_proj"], 1), "und/mes"],
        ["Grand Plaz histórica", round(ref["grand_hist"], 1), "und/mes"],
        [f"Grand Plaz proyectada (+{round((GRAND_PLAZ_BONUS - 1) * 100)}% hist.)", round(ref["grand_proj"], 1), "und/mes"],
        ["Barquisimeto (avg Grieta+Chacao+Tolón proy.)", round(ref["barq_proj"], 1), "und/mes"],
        ["Corporativo", "EXCLUIDO"],
        [],
        ["Notas"],
        [f"• Referencia: {REFERENCE_MODEL} DAMA del dashboard Basic Line."],
        ["• Producto NUEVO manufacturado — sin stock inicial."],
        ["• Solo género DAMA, tallas XS a XL."],
        ["• Compra tela VIORI incluye +20% stock de seguridad."],
    ]
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = Font(bold=True, size=14, color="FFFFFF")
                cell.fill = title_fill
            elif row and row[0] and str(row[0]).startswith("──"):
                cell.font = Font(bold=True, color="1E3A5F")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 14


def write_tallas_sheet(wb, ref, prod):
    ws = wb.create_sheet("Producción por Talla")
    talla_min = distribute_by_talla(prod["prod_min"], ref["talla_pct"])
    talla_max = distribute_by_talla(prod["prod_max"], ref["talla_pct"])

    ws.merge_cells("A1:E1")
    style_cell(ws.cell(row=1, column=1, value=f"{PRODUCTO} — CANTIDADES POR TALLA (MÍN / MÁX)"), title_fill, bold=True)
    ws.cell(row=2, column=1, value=f"{GENERO} · Curva {REFERENCE_MODEL} DAMA").font = Font(italic=True)

    headers = ["Talla", "Curva %", "Mínimo", "Máximo"]
    for c, h in enumerate(headers, 1):
        style_cell(ws.cell(row=4, column=c, value=h), sub_fill, bold=True)

    t_min_total = t_max_total = 0
    for i, t in enumerate(TALLAS, start=5):
        mn, mx = talla_min[t], talla_max[t]
        style_cell(ws.cell(row=i, column=1, value=t), color_fill, bold=True)
        ws.cell(row=i, column=2, value=f"{ref['talla_pct'][t]*100:.1f}%")
        style_cell(ws.cell(row=i, column=3, value=mn), min_fill, bold=True)
        style_cell(ws.cell(row=i, column=4, value=mx), max_fill, bold=True)
        t_min_total += mn
        t_max_total += mx

    r = len(TALLAS) + 5
    style_cell(ws.cell(row=r, column=1, value="TOTAL"), tot_fill, bold=True, align=left)
    style_cell(ws.cell(row=r, column=3, value=t_min_total), tot_fill, bold=True)
    style_cell(ws.cell(row=r, column=4, value=t_max_total), tot_fill, bold=True)

    for col in "ABCDE":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 18


def write_tiendas_sheet(wb, ref, prod):
    ws = wb.create_sheet("Distribución por Tienda")
    adjusted = {TOLON_STORE, WEB_STORE, NEW_STORE, "LA VELA", "GRAND PLAZ"}
    store_min = distribute_by_store(prod["prod_min"], ref)
    store_max = distribute_by_store(prod["prod_max"], ref)
    matrix_min = distribute_store_talla(store_min, ref["talla_pct"])
    matrix_max = distribute_store_talla(store_max, ref["talla_pct"])

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(TALLAS) * 2)
    style_cell(ws.cell(row=1, column=1, value=f"{PRODUCTO} — DISTRIBUCIÓN POR TIENDA Y TALLA"), title_fill, bold=True)
    ws.cell(row=2, column=1, value="★ = tienda con velocidad proyectada/ajustada").font = Font(italic=True, color="E65100")

    col = 2
    for t in TALLAS:
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 1)
        style_cell(ws.cell(row=3, column=col, value=t), sub_fill, bold=True)
        style_cell(ws.cell(row=4, column=col, value="Mín"), min_fill, bold=True)
        style_cell(ws.cell(row=4, column=col + 1, value="Máx"), max_fill, bold=True)
        col += 2

    style_cell(ws.cell(row=3, column=1, value="Tienda"), sub_fill, bold=True, align=left)

    row = 5
    for store in DISTRIBUTION_STORES:
        is_adj = store in adjusted
        label = store + (" ★" if is_adj else "")
        fill = adj_fill if is_adj else white_fill
        style_cell(ws.cell(row=row, column=1, value=label), fill, bold=is_adj, align=left)
        col = 2
        for t in TALLAS:
            style_cell(ws.cell(row=row, column=col, value=matrix_min[store][t]), min_fill if matrix_min[store][t] else white_fill)
            style_cell(ws.cell(row=row, column=col + 1, value=matrix_max[store][t]), max_fill if matrix_max[store][t] else white_fill)
            col += 2
        row += 1

    style_cell(ws.cell(row=row, column=1, value="TOTAL"), tot_fill, bold=True, align=left)
    col = 2
    for t in TALLAS:
        style_cell(ws.cell(row=row, column=col, value=sum(matrix_min[s][t] for s in DISTRIBUTION_STORES)), tot_fill, bold=True)
        style_cell(ws.cell(row=row, column=col + 1, value=sum(matrix_max[s][t] for s in DISTRIBUTION_STORES)), tot_fill, bold=True)
        col += 2

    row += 2
    ws.cell(row=row, column=1, value="Totales por tienda (Mín / Máx):").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="Tienda").font = Font(bold=True)
    ws.cell(row=row, column=2, value="Mín").font = Font(bold=True)
    ws.cell(row=row, column=3, value="Máx").font = Font(bold=True)
    ws.cell(row=row, column=4, value="Vel. mensual proy.").font = Font(bold=True)
    row += 1
    for store in DISTRIBUTION_STORES:
        ws.cell(row=row, column=1, value=store + (" ★" if store in adjusted else ""))
        ws.cell(row=row, column=2, value=store_min[store])
        ws.cell(row=row, column=3, value=store_max[store])
        ws.cell(row=row, column=4, value=round(ref["store_monthly"][store], 1))
        row += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["D"].width = 16
    for i in range(2, 2 + len(TALLAS) * 2):
        ws.column_dimensions[get_column_letter(i)].width = 7


def _write_curva_block(ws, start_row: int, block_title: str, color_matrix: dict) -> int:
    """Escribe bloque curva color×talla estilo SPOTS. Retorna siguiente fila libre."""
    ncol = len(TALLAS) + 2
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncol)
    c = ws.cell(row=start_row, column=1, value=block_title)
    c.font = Font(bold=True, size=12, color="1F3864")
    c.alignment = left
    row = start_row + 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1)
    style_cell(ws.cell(row=row, column=1, value=f"{PRODUCTO} {GENERO}"), CURVA_HDR_FILL, bold=True, align=left)
    ws.cell(row=row, column=1).font = CURVA_HDR_FONT
    for i, t in enumerate(TALLAS, start=2):
        style_cell(ws.cell(row=row, column=i, value=t), CURVA_HDR_FILL, bold=True)
        ws.cell(row=row, column=i).font = CURVA_HDR_FONT
    style_cell(ws.cell(row=row, column=ncol, value="Tot"), CURVA_HDR_FILL, bold=True)
    ws.cell(row=row, column=ncol).font = CURVA_HDR_FONT
    row += 1

    col_totals = {t: 0 for t in TALLAS}
    grand = 0
    for color in COLORES:
        pct = int(COLOR_PCT[color] * 100)
        style_cell(ws.cell(row=row, column=1, value=f"{color} · {pct}%"), CURVA_DATA_FILL, align=left)
        ws.cell(row=row, column=1).font = CURVA_LBL_FONT
        row_sum = 0
        for i, t in enumerate(TALLAS, start=2):
            val = color_matrix[color][t]
            cell = ws.cell(row=row, column=i, value=val)
            style_cell(cell, CURVA_DATA_FILL)
            cell.font = CURVA_NUM_FONT
            col_totals[t] += val
            row_sum += val
        cell = ws.cell(row=row, column=ncol, value=row_sum)
        style_cell(cell, CURVA_DATA_FILL)
        cell.font = CURVA_NUM_FONT
        grand += row_sum
        row += 1

    style_cell(ws.cell(row=row, column=1, value="TOTAL GENERAL"), CURVA_TOT_FILL, bold=True, align=left)
    ws.cell(row=row, column=1).font = Font(bold=True, size=11)
    for i, t in enumerate(TALLAS, start=2):
        cell = ws.cell(row=row, column=i, value=col_totals[t])
        style_cell(cell, CURVA_TOT_FILL, bold=True)
        cell.font = Font(bold=True, size=11)
    cell = ws.cell(row=row, column=ncol, value=grand)
    style_cell(cell, CURVA_TOT_FILL, bold=True)
    cell.font = Font(bold=True, size=11)
    return row + 2


def write_curva_completa_sheet(wb, ref, prod):
    ws = wb.create_sheet("Cantidades por Colores")
    color_min, _ = distribute_by_color_talla(prod["prod_min"], ref["talla_pct"])
    color_max, _ = distribute_by_color_talla(prod["prod_max"], ref["talla_pct"])
    tela_min = calc_tela_totals(color_min)
    tela_max = calc_tela_totals(color_max)
    talla_min = distribute_by_talla(prod["prod_min"], ref["talla_pct"])
    talla_max = distribute_by_talla(prod["prod_max"], ref["talla_pct"])
    ins_min = calc_insumos_totals(talla_min)
    ins_max = calc_insumos_totals(talla_max)

    ws.merge_cells("A1:G1")
    style_cell(ws.cell(row=1, column=1, value="CANTIDADES POR COLORES"), title_fill, bold=True)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="FFFFFF")
    ws.cell(row=2, column=1, value=f"{PRODUCTO} · Proporción Negro 40% · Vinotinto 30% · Verde Militar 30%").font = Font(italic=True)
    ws.cell(row=3, column=1, value=f"Rango producción: {prod['prod_min']} – {prod['prod_max']} und").font = Font(bold=True)

    row = _write_curva_block(ws, 5, f"MÍNIMO — {prod['prod_min']} und (compromiso)", color_min)
    row = _write_curva_block(ws, row, f"MÁXIMO — {prod['prod_max']} und (techo)", color_max)

    ws.cell(row=row, column=1, value="── INSUMOS PARA COMPRA (Mín / Máx) ──").font = Font(bold=True, size=11, color="1F3864")
    row += 1
    ins_headers = ["Insumo", "Mín", "Máx", "Unidad", "Nota"]
    for c, h in enumerate(ins_headers, 1):
        style_cell(ws.cell(row=row, column=c, value=h), sub_fill, bold=True)
    row += 1

    ins_rows = [
        ("Tela VIORI — Negro (consumo)", tela_min["by_color"]["Negro"]["mts"], tela_max["by_color"]["Negro"]["mts"], "metros", "Sin SS"),
        ("Tela VIORI — Negro (compra +SS)", tela_min["by_color"]["Negro"]["mts_compra"], tela_max["by_color"]["Negro"]["mts_compra"], "metros", f"{tela_max['by_color']['Negro']['kg_compra']} kg al máx"),
        ("Tela VIORI — Vinotinto (compra +SS)", tela_min["by_color"]["Vinotinto"]["mts_compra"], tela_max["by_color"]["Vinotinto"]["mts_compra"], "metros", f"{tela_max['by_color']['Vinotinto']['kg_compra']} kg al máx"),
        ("Tela VIORI — Verde Militar (compra +SS)", tela_min["by_color"]["Verde Militar"]["mts_compra"], tela_max["by_color"]["Verde Militar"]["mts_compra"], "metros", f"{tela_max['by_color']['Verde Militar']['kg_compra']} kg al máx"),
        ("TOTAL tela VIORI (compra +SS)", tela_min["total_mts_compra"], tela_max["total_mts_compra"], "metros", f"{tela_max['total_kg_compra']} kg al máx · SS {int(TELA_SS_PCT*100)}%"),
        ("Elástica 4.5 cm", ins_min["elastica_m"], ins_max["elastica_m"], "metros", "Ficha técnica LITE PANT"),
        ("Etiqueta agua / SENCAMER", ins_min["etiquetas"], ins_max["etiquetas"], "und", "1 und/pieza"),
    ]
    for label, vmin, vmax, unit, note in ins_rows:
        is_sub = label.startswith("  ·")
        fill = white_fill if is_sub else color_fill
        style_cell(ws.cell(row=row, column=1, value=label), fill, bold=not is_sub, align=left)
        style_cell(ws.cell(row=row, column=2, value=vmin), min_fill if not is_sub else white_fill)
        style_cell(ws.cell(row=row, column=3, value=vmax), max_fill if not is_sub else white_fill)
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=note)
        row += 1

    ws.column_dimensions["A"].width = 32
    for col in "BCDE":
        ws.column_dimensions[col].width = 14


def write_colores_sheet(wb, ref, prod):
    ws = wb.create_sheet("Producción Color × Talla")
    color_min, totals_min = distribute_by_color_talla(prod["prod_min"], ref["talla_pct"])
    color_max, totals_max = distribute_by_color_talla(prod["prod_max"], ref["talla_pct"])

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(TALLAS) * 2)
    style_cell(ws.cell(row=1, column=1, value=f"{PRODUCTO} — PRODUCCIÓN POR COLOR Y TALLA"), title_fill, bold=True)
    ws.cell(row=2, column=1, value="Proporción: Negro 40% · Vinotinto 30% · Verde Militar 30%").font = Font(italic=True)
    ws.cell(row=3, column=1, value=f"Rango global: {prod['prod_min']} – {prod['prod_max']} und").font = Font(italic=True)

    row = 5
    for color in COLORES:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2 + len(TALLAS) * 2)
        c = ws.cell(row=row, column=1, value=f"{color} ({int(COLOR_PCT[color]*100)}%)")
        style_cell(c, COLOR_FILL[color], bold=True, align=left)
        c.font = Font(bold=True, color=COLOR_FONT[color])
        row += 1

        col = 2
        for t in TALLAS:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
            style_cell(ws.cell(row=row, column=col, value=t), sub_fill, bold=True)
            style_cell(ws.cell(row=row + 1, column=col, value="Mín"), min_fill, bold=True)
            style_cell(ws.cell(row=row + 1, column=col + 1, value="Máx"), max_fill, bold=True)
            col += 2
        style_cell(ws.cell(row=row, column=1, value="Talla"), sub_fill, bold=True, align=left)
        row += 2

        style_cell(ws.cell(row=row, column=1, value="Cantidad"), color_fill, bold=True, align=left)
        col = 2
        tmin = tmax = 0
        for t in TALLAS:
            mn, mx = color_min[color][t], color_max[color][t]
            style_cell(ws.cell(row=row, column=col, value=mn), min_fill, bold=True)
            style_cell(ws.cell(row=row, column=col + 1, value=mx), max_fill, bold=True)
            tmin += mn
            tmax += mx
            col += 2
        row += 1
        style_cell(ws.cell(row=row, column=1, value="TOTAL color"), tot_fill, bold=True, align=left)
        style_cell(ws.cell(row=row, column=2, value=tmin), tot_fill, bold=True)
        style_cell(ws.cell(row=row, column=3, value=tmax), tot_fill, bold=True)
        row += 2

    # Gran total
    style_cell(ws.cell(row=row, column=1, value="TOTAL GENERAL"), tot_fill, bold=True, align=left)
    style_cell(ws.cell(row=row, column=2, value=prod["prod_min"]), tot_fill, bold=True)
    style_cell(ws.cell(row=row, column=3, value=prod["prod_max"]), tot_fill, bold=True)
    row += 2
    ws.cell(row=row, column=1, value="Resumen por color (Mín / Máx):").font = Font(bold=True)
    row += 1
    for color in COLORES:
        ws.cell(row=row, column=1, value=color)
        ws.cell(row=row, column=2, value=totals_min[color])
        ws.cell(row=row, column=3, value=totals_max[color])
        row += 1

    ws.column_dimensions["A"].width = 22
    for i in range(2, 2 + len(TALLAS) * 2):
        ws.column_dimensions[get_column_letter(i)].width = 7


def write_compra_tela_sheet(wb, ref, prod):
    ws = wb.create_sheet("Compra de Tela VIORI")
    color_min, _ = distribute_by_color_talla(prod["prod_min"], ref["talla_pct"])
    color_max, _ = distribute_by_color_talla(prod["prod_max"], ref["talla_pct"])
    tela_min = calc_tela_totals(color_min)
    tela_max = calc_tela_totals(color_max)
    talla_min = distribute_by_talla(prod["prod_min"], ref["talla_pct"])
    talla_max = distribute_by_talla(prod["prod_max"], ref["talla_pct"])
    ins_min = calc_insumos_totals(talla_min)
    ins_max = calc_insumos_totals(talla_max)

    ws.merge_cells("A1:J1")
    style_cell(ws.cell(row=1, column=1, value=f"{PRODUCTO} — COMPRA DE TELA VIORI POR COLOR"), title_fill, bold=True)
    ws.cell(row=2, column=1, value=f"Consumo ficha técnica DAMA · 3 colores 40/30/30 · SS tela +{int(TELA_SS_PCT*100)}%").font = Font(italic=True)

    headers = [
        "Color", "%", "Und Mín", "Und Máx",
        "Mts consumo Mín", "Mts consumo Máx",
        "Mts compra Mín (+SS)", "Mts compra Máx (+SS)",
        "Kg compra Mín", "Kg compra Máx",
    ]
    for c, h in enumerate(headers, 1):
        style_cell(ws.cell(row=4, column=c, value=h), sub_fill, bold=True)

    row = 5
    for color in COLORES:
        bc = tela_min["by_color"][color]
        bx = tela_max["by_color"][color]
        style_cell(ws.cell(row=row, column=1, value=color), COLOR_FILL[color], bold=True, align=left)
        ws.cell(row=row, column=1).font = Font(bold=True, color=COLOR_FONT[color])
        ws.cell(row=row, column=2, value=f"{int(COLOR_PCT[color]*100)}%")
        ws.cell(row=row, column=3, value=bc["und"])
        ws.cell(row=row, column=4, value=bx["und"])
        ws.cell(row=row, column=5, value=bc["mts"])
        ws.cell(row=row, column=6, value=bx["mts"])
        style_cell(ws.cell(row=row, column=7, value=bc["mts_compra"]), min_fill, bold=True)
        style_cell(ws.cell(row=row, column=8, value=bx["mts_compra"]), max_fill, bold=True)
        style_cell(ws.cell(row=row, column=9, value=bc["kg_compra"]), min_fill)
        style_cell(ws.cell(row=row, column=10, value=bx["kg_compra"]), max_fill)
        row += 1

    style_cell(ws.cell(row=row, column=1, value="TOTAL TELA VIORI"), tot_fill, bold=True, align=left)
    ws.cell(row=row, column=3, value=tela_min["total_und"])
    ws.cell(row=row, column=4, value=tela_max["total_und"])
    ws.cell(row=row, column=5, value=tela_min["total_mts"])
    ws.cell(row=row, column=6, value=tela_max["total_mts"])
    style_cell(ws.cell(row=row, column=7, value=tela_min["total_mts_compra"]), tot_fill, bold=True)
    style_cell(ws.cell(row=row, column=8, value=tela_max["total_mts_compra"]), tot_fill, bold=True)
    style_cell(ws.cell(row=row, column=9, value=tela_min["total_kg_compra"]), tot_fill, bold=True)
    style_cell(ws.cell(row=row, column=10, value=tela_max["total_kg_compra"]), tot_fill, bold=True)
    row += 2

    ws.cell(row=row, column=1, value="── DETALLE POR TALLA Y COLOR (MÁX) ──").font = Font(bold=True)
    row += 1
    hdr = ["Color / Talla"] + TALLAS + ["Total", "Mts consumo", "Mts compra (+SS)", "Kg compra (+SS)"]
    for c, h in enumerate(hdr, 1):
        style_cell(ws.cell(row=row, column=c, value=h), sub_fill, bold=True)
    row += 1
    for color in COLORES:
        style_cell(ws.cell(row=row, column=1, value=color), COLOR_FILL[color], bold=True, align=left)
        ws.cell(row=row, column=1).font = Font(bold=True, color=COLOR_FONT[color])
        for i, t in enumerate(TALLAS, 2):
            ws.cell(row=row, column=i, value=color_max[color][t])
        ws.cell(row=row, column=7, value=tela_max["by_color"][color]["und"])
        ws.cell(row=row, column=8, value=tela_max["by_color"][color]["mts"])
        style_cell(ws.cell(row=row, column=9, value=tela_max["by_color"][color]["mts_compra"]), max_fill, bold=True)
        ws.cell(row=row, column=10, value=tela_max["by_color"][color]["kg_compra"])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="── CONSUMO UNITARIO VIORI (referencia ficha) ──").font = Font(bold=True)
    row += 1
    style_cell(ws.cell(row=row, column=1, value="Talla"), sub_fill, bold=True, align=left)
    for i, t in enumerate(TALLAS, 2):
        style_cell(ws.cell(row=row, column=i, value=t), sub_fill, bold=True)
    row += 1
    ws.cell(row=row, column=1, value="Metros/pieza")
    for i, t in enumerate(TALLAS, 2):
        ws.cell(row=row, column=i, value=CONSUMO_MTS[t])
    row += 1
    ws.cell(row=row, column=1, value="Kg/pieza")
    for i, t in enumerate(TALLAS, 2):
        ws.cell(row=row, column=i, value=CONSUMO_KG[t])

    row += 2
    ws.cell(row=row, column=1, value="── OTROS INSUMOS (total producción Mín / Máx) ──").font = Font(bold=True)
    row += 1
    for label, k in [
        ("Elástica 4.5 cm (metros)", "elastica_m"),
        ("Etiqueta agua / SENCAMER (und)", "etiquetas"),
    ]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=ins_min[k])
        ws.cell(row=row, column=3, value=ins_max[k])
        row += 1

    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 28


def write_insumos_sheet(wb, ref, prod):
    ws = wb.create_sheet("Insumos")
    talla_min = distribute_by_talla(prod["prod_min"], ref["talla_pct"])
    talla_max = distribute_by_talla(prod["prod_max"], ref["talla_pct"])
    ins_min = calc_insumos_totals(talla_min)
    ins_max = calc_insumos_totals(talla_max)

    rows = [
        [f"INSUMOS — {PRODUCTO} {GENERO}"],
        ["Fuente", "Ficha_tecnica_LITE_PANT.xlsx"],
        [],
        ["── ELÁSTICA 4.5 CM ──"],
        ["Talla", "Cm/pieza", "Demanda Mín", "Demanda Máx"],
    ]
    for t in TALLAS:
        rows.append([t, ELASTICA_CM[t], talla_min[t], talla_max[t]])
    rows += [
        [],
        ["Total elástica (metros)", "", ins_min["elastica_m"], ins_max["elastica_m"]],
        [],
        ["── ETIQUETA AGUA / SENCAMER ──"],
        ["Consumo", "1 und/pieza", "Ubicación", "Costado derecho"],
        ["Demanda Mín", ins_min["etiquetas"], "Demanda Máx", ins_max["etiquetas"]],
        [],
        ["── HILOS (por color de tela) ──"],
        ["Negro", "Hilo Negro-N", "Vinotinto", "Hilo Vinotinto"],
        ["Verde Militar", "Hilo Verde Militar", "", ""],
        [],
        [f"Rango producción", f"{prod['prod_min']} – {prod['prod_max']} und"],
    ]
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if r == 1 or (row and row[0] and str(row[0]).startswith("──")):
                cell.font = Font(bold=True)
                if r == 1:
                    cell.fill = title_fill
                    cell.font = Font(bold=True, color="FFFFFF")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16


def write_metodologia(wb, ref, prod):
    ws = wb.create_sheet("Metodología")
    text = [
        f"METODOLOGÍA — PROYECCIÓN {PRODUCTO}",
        "",
        "1. PRODUCTO DE REFERENCIA",
        f"   Fuente: Dashboard_Basic_Line.html — filtro {REFERENCE_MODEL} DAMA.",
        f"   Velocidad base = promedio Jun–Jul–Ago 2026: {ref['vel_base']:.0f} und/mes.",
        "   Ficha técnica: Ficha_tecnica_LITE_PANT.xlsx (consumo VIORI + elástica).",
        "",
        "2. AJUSTE TEMPORADA ALTA",
        f"   Factor ×{HIGH_SEASON_FACTOR} (diciembre). Diciembre 2025 combinado: {ref['dec_vel']} und.",
        "",
        "3. AJUSTES DE TIENDA (según indicación)",
        f"   • Tolón: proyectado al {int(TOLON_VS_CHACAO*100)}% de Chacao ({ref['chacao_m']:.0f} → {ref['tolon_proj']:.0f} und/mes).",
        f"     Histórico Tolón: {ref['store_monthly_hist'].get('TOLON', 0):.0f} und/mes — subestimado por tienda nueva.",
        f"   • Web: proyectada al {int(WEB_VS_CERRO_VERDE*100)}% de Cerro Verde ({ref['cerro_m']:.0f} → {ref['web_proj']:.0f} und/mes).",
        f"     Histórico Web: {ref['store_monthly_hist'].get('WEB', 0):.0f} und/mes.",
        f"   • La Vela (nueva): promedio Sambil + Cerro Verde = {ref['vela_proj']:.0f} und/mes.",
        f"     Histórico La Vela: {ref['vela_hist']:.0f} und/mes.",
        f"   • Grand Plaz: histórico × {GRAND_PLAZ_BONUS} (+{round((GRAND_PLAZ_BONUS-1)*100)}%) = {ref['grand_proj']:.0f} und/mes.",
        f"     Histórico Grand Plaz: {ref['grand_hist']:.0f} und/mes.",
        f"   • Barquisimeto (nueva): promedio Grieta + Chacao + Tolón proyectado = {ref['barq_proj']:.0f} und/mes.",
        "   • Corporativo: EXCLUIDO.",
        "",
        "4. RANGO DE PRODUCCIÓN",
        f"   Mínimo: {prod['prod_min']} und | Máximo: {prod['prod_max']} und.",
        f"   Calculado: {COVER_MONTHS_MIN}–{COVER_MONTHS_MAX} meses cobertura + SS {int(SAFETY_STOCK_PCT*100)}%.",
        f"   Velocidad red ajustada: {prod['vel_network']:.0f} und/mes.",
        "",
        "5. DISTRIBUCIÓN",
        "   Por tienda: pesos mensuales proyectados (Tolón/Web/Barquisimeto ajustados).",
        f"   Por talla: curva {REFERENCE_MODEL} DAMA del dashboard.",
        "",
        "6. COLORES Y COMPRA DE TELA",
        "   Colores producción: Negro 40% · Vinotinto 30% · Verde Militar 30%.",
        "   Dentro de cada color se aplica la misma curva de tallas.",
        "   Consumo VIORI por pieza (ficha LITE PANT DAMA): XS 1.45m · S 1.48m · M 1.58m · L 1.63m · XL 1.66m.",
        f"   Stock de seguridad tela: +{int(TELA_SS_PCT*100)}% sobre consumo (compra = consumo × {1+TELA_SS_PCT}).",
        "   Ver hojas 'Producción Color × Talla' y 'Compra de Tela VIORI'.",
        "",
        "7. OTROS INSUMOS",
        "   Elástica 4.5 cm: consumo por talla según ficha (67–77 cm/pieza).",
        "   Etiqueta agua/SENCAMER: 1 und/pieza · costado derecho.",
        "   Hilos: Negro-N · Vinotinto · Verde Militar (según color de tela).",
    ]
    for r, line in enumerate(text, start=1):
        cell = ws.cell(row=r, column=1, value=line)
        if r == 1:
            cell.font = Font(bold=True, size=13)
        if line and line[0].isdigit():
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 95


def main():
    data = load_combined_data()
    ref = analyze_reference(data)
    prod = calc_production(ref)

    wb = Workbook()
    wb.remove(wb.active)
    write_resumen(wb, ref, prod)
    write_tallas_sheet(wb, ref, prod)
    write_curva_completa_sheet(wb, ref, prod)
    write_colores_sheet(wb, ref, prod)
    write_compra_tela_sheet(wb, ref, prod)
    write_tiendas_sheet(wb, ref, prod)
    write_insumos_sheet(wb, ref, prod)
    write_metodologia(wb, ref, prod)
    wb.save(OUTPUT_PATH)

    print(f"✅ Generado: {OUTPUT_PATH}")
    print(f"   Fuentes: {', '.join(ref['sources'])}")
    print(f"   Velocidad base (Jun-Ago 2026): {ref['vel_base']:.1f} und/mes")
    print(f"   Tolón proy: {ref['tolon_proj']:.1f}/mes ({TOLON_VS_CHACAO*100:.0f}% Chacao)")
    print(f"   Web proy: {ref['web_proj']:.1f}/mes ({WEB_VS_CERRO_VERDE*100:.0f}% Cerro Verde)")
    print(f"   Rango producción: {prod['prod_min']} – {prod['prod_max']} und")
    color_min_m, cm = distribute_by_color_talla(prod["prod_min"], ref["talla_pct"])
    color_max_m, cx = distribute_by_color_talla(prod["prod_max"], ref["talla_pct"])
    tm = calc_tela_totals(color_min_m)
    tx = calc_tela_totals(color_max_m)
    print(f"   Tela VIORI máx consumo: {tx['total_mts']} mts / {tx['total_kg']} kg")
    print(f"   Tela VIORI máx compra (+{int(TELA_SS_PCT*100)}% SS): {tx['total_mts_compra']} mts / {tx['total_kg_compra']} kg")
    print(f"   Colores máx: Negro {cx['Negro']} · Vinotinto {cx['Vinotinto']} · V.Militar {cx['Verde Militar']}")


if __name__ == "__main__":
    main()
