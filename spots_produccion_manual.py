"""Carga producción expansión desde SPOTS_PRODUCCION_EXPANSION.xlsx (CAB/DAMA)."""

from collections import defaultdict
from pathlib import Path

from build_spots_dashboard import (
    BASE_STORE,
    BQT_DISENO_VIRGEN,
    DEFAULT_PROD_MONTHS,
    DECEMBER_HS_FACTOR,
    EXPANSION_CAPS,
    EXPANSION_STORES,
    HIGH_SEASON_FACTOR,
    LEAD_MONTHS,
    MC_MODEL,
    PROD_MONTHS_OPTIONS,
    TALLA_BOOST,
    VELOCITY_MONTHS_COUNT,
    ZONE_ADICIONAL_COLOR,
    _sku_from_alloc,
    _zone_adicional_color,
    compute_fabric_needs,
    velocity_months_label,
)

MANUAL_PRODUCTION_XLSX = Path(__file__).resolve().parent / "SPOTS_PRODUCCION_EXPANSION.xlsx"


def _parse_label(label: str):
    if " · Blanco (" in label:
        diseno, rest = label.split(" · Blanco (")
        return rest.rstrip(")").upper(), diseno, "Blanco"
    if " · " in label:
        color, zona = label.rsplit(" · ", 1)
        cap = EXPANSION_CAPS.get(zona.upper(), {})
        diseno = cap.get("blanco_diseno") or cap.get("adicional_diseno") or color
        if color != "Blanco" and zona.upper() == "BARQUISIMETO":
            diseno = "Ciudad"
        return zona.upper(), diseno, color
    return None, None, None


def parse_consolidated_sheet(ws, genero: str):
    tallas = []
    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if not any(v is not None for v in vals):
            continue
        if vals[0] == f"SPOTS MANGA CORTA {genero}":
            tallas = [v for v in vals[1:-1] if v]
            continue
        if vals[0] in ("CANTIDADES POR COLORES", "TOTAL BLANCO", "TOTAL COLORES", "TOTAL GENERAL"):
            continue
        if not tallas or not vals[0]:
            continue
        zona, diseno, color = _parse_label(str(vals[0]))
        if not zona:
            continue
        qtys = {str(t): int(vals[i + 1] or 0) for i, t in enumerate(tallas)}
        rows.append({
            "label": vals[0],
            "zona": zona,
            "diseno": diseno,
            "color": color,
            "genero": genero,
            "tallas": qtys,
            "total": sum(qtys.values()),
        })
    return tallas, rows


def _hs_for(diseno: str, color: str):
    if diseno == BQT_DISENO_VIRGEN and color == "Blanco":
        return DECEMBER_HS_FACTOR, "Festividad Virgen · rotación dic ×1.4"
    return HIGH_SEASON_FACTOR, ""


def build_expansion_from_rows(cab_rows, dama_rows, vel_months=None, months=DEFAULT_PROD_MONTHS):
    vel_months = vel_months or []
    vel_lbl = ", ".join(velocity_months_label(vel_months)) if vel_months else ""
    all_rows = cab_rows + dama_rows

    by_zone_rows = defaultdict(list)
    for r in all_rows:
        by_zone_rows[r["zona"]].append(r)

    by_store = []
    total_blanco = total_color = 0

    for store in EXPANSION_STORES:
        cap = EXPANSION_CAPS[store]
        rows = by_zone_rows.get(store, [])
        if not rows:
            continue
        skus = []
        designs_meta = []
        design_totals = defaultdict(int)
        store_blanco = store_color = 0

        for r in rows:
            hs, seasonality = _hs_for(r["diseno"], r["color"])
            design_total = 0
            for talla, need_3m in r["tallas"].items():
                if need_3m <= 0:
                    continue
                skus.append(_sku_from_alloc(
                    MC_MODEL, r["genero"], talla,
                    r["diseno"], r["color"], need_3m, hs, seasonality,
                ))
                design_total += need_3m
            key = (r["diseno"], r["color"])
            design_totals[key] += design_total
            if r["color"] == "Blanco":
                store_blanco += design_total
            else:
                store_color += design_total

        for (diseno, color), target in sorted(design_totals.items()):
            designs_meta.append({"diseno": diseno, "color": color, "target_3m": target})

        store_total = store_blanco + store_color
        total_blanco += store_blanco
        total_color += store_color
        adicional_color = _zone_adicional_color(cap, store)

        if store == "BARQUISIMETO":
            ciudad_b = sum(t for (d, c), t in design_totals.items() if d == "Ciudad" and c == "Blanco")
            virgen_b = sum(t for (d, c), t in design_totals.items() if d == BQT_DISENO_VIRGEN and c == "Blanco")
            label = (
                f"Barquisimeto · 1 tienda · MC · Ciudad {ciudad_b} + Virgen {virgen_b} + "
                f"Verde {store_color} · total {store_total}"
            )
        elif store == "CARACAS":
            label = f"Caracas · 4 tiendas CCS · MC · Blanco {store_blanco} + Azul Marino {store_color} · total {store_total}"
        elif store == "VALENCIA":
            label = f"Valencia · 2 tiendas · MC · Blanco {store_blanco} + Vinotinto {store_color} · total {store_total}"
        else:
            label = cap.get("label", store)

        weights = cap.get("tienda_weights")
        by_store.append({
            "store": store,
            "label": label,
            "adicional_color": adicional_color,
            "tiendas": cap.get("tiendas", 1),
            "tienda_weights": weights,
            "effective_tiendas": round(sum(weights), 2) if weights else cap.get("tiendas", 1),
            "target_total_3m": store_total,
            "blanco": store_blanco,
            "adicional": store_color,
            "total": store_total,
            "skus": skus,
            "designs": designs_meta,
        })

    expansion = {
        "base_store": BASE_STORE,
        "stores": EXPANSION_STORES,
        "zone_colors": ZONE_ADICIONAL_COLOR,
        "additional_color_factor": 0.7,
        "high_season_factor": HIGH_SEASON_FACTOR,
        "december_hs_factor": DECEMBER_HS_FACTOR,
        "velocity_months_count": VELOCITY_MONTHS_COUNT,
        "lead_months": LEAD_MONTHS,
        "prod_months_options": PROD_MONTHS_OPTIONS,
        "default_prod_months": DEFAULT_PROD_MONTHS,
        "talla_boost": TALLA_BOOST,
        "vela_exclusive_note": (
            "Diseños VELA (Nueva Esparta, Virgen del Valle, Manga Larga, etc.) "
            "son exclusivos de la zona Margarita"
        ),
        "by_store": by_store,
        "total_blanco": total_blanco,
        "total_adicional": total_color,
        "total_expansion": total_blanco + total_color,
        "source": "manual_xlsx",
        "manual_months": months,
    }
    expansion["by_color"] = [{"color": "Blanco", "need_3m": total_blanco}]
    for s in by_store:
        expansion["by_color"].append({
            "color": s["adicional_color"],
            "need_3m": s["adicional"],
            "zona": s["store"],
        })
    expansion["nota"] = (
        f"Cantidades ajustadas manualmente (SPOTS_PRODUCCION_EXPANSION.xlsx) · "
        f"horizonte {months} meses · colores: Caracas Azul Marino · Valencia Vinotinto · "
        f"Barquisimeto Verde. Total {expansion['total_expansion']} und."
    )
    if vel_lbl:
        expansion["nota"] += f" Base rotación VELA: {vel_lbl}."
    expansion["fabric"] = compute_fabric_needs(expansion)
    return expansion


def load_manual_expansion(path: Path | None = None, vel_months=None, months=DEFAULT_PROD_MONTHS):
    xlsx = Path(path) if path else MANUAL_PRODUCTION_XLSX
    if not xlsx.exists():
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    wb = load_workbook(xlsx, data_only=True)
    if "CAB" not in wb.sheetnames or "DAMA" not in wb.sheetnames:
        return None
    _, cab_rows = parse_consolidated_sheet(wb["CAB"], "CAB")
    _, dama_rows = parse_consolidated_sheet(wb["DAMA"], "DAMA")
    if not cab_rows and not dama_rows:
        return None
    return build_expansion_from_rows(cab_rows, dama_rows, vel_months, months)
