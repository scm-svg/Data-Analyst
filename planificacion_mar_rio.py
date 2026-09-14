#!/usr/bin/env python3
"""
Plan Mar/Rio — Excel editable, prioridad Rio KIDS → Mar KIDS, curvas por talla.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAR_XLSX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/MAR_PROYECCION_CANTIDADES_SUGERIDAS__2__23dd.xlsx")
RIO_XLSX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/RIO_PROYECCION_CANTIDADES_SUGERIDAS__2__cc77.xlsx")
OUT_DIR = Path("/workspace/output")

TELA_KG = {
    "Aguamarina": 348.34, "Amarillo Neón": 57.94, "Azul Lavanda": 221.00,
    "Azul Marino": 689.90, "Azul Rey": 271.46, "Blanco": 162.04,
    "Gris Claro": 47.68, "Lila": 188.44, "Morado": 94.20, "Negro": 657.16,
    "Rojo": 0.18, "Rosado Pastel": 31.44, "Verde Militar": 46.02, "Vinotinto": 175.38,
}

LOTE_CORTADO_COLORES = [
    "Verde Militar", "Vinotinto", "Azul Marino", "Azul Rey", "Rojo", "Gris Claro",
    "Azul Lavanda", "Aguamarina", "Rosado Pastel", "Negro", "Lila",
]
LOTE_CORTADO_TOTAL = 1300
LOTE_A_META_KIDS = {
    "Verde Militar": "Verde Militar", "Vinotinto": None, "Azul Marino": "Azul Marino",
    "Azul Rey": "Azul Rey", "Rojo": "Rojo", "Gris Claro": None,
    "Azul Lavanda": "Azul Lavanda", "Aguamarina": "Aguamarina",
    "Rosado Pastel": "Rosado Pastel", "Negro": "Negro", "Lila": "Lila",
}

MIN_KG = 5.0
PRIORIDAD_ORDEN = ["RIO KIDS", "MAR KIDS", "MAR CAB", "MAR DAMA", "RIO CAB", "RIO DAMA"]

LINEAS_CONFIG = [
    ("RIO", "KIDS", 1, 0.50, 0.32),
    ("MAR", "KIDS", 2, 0.50, 0.26),
    ("MAR", "CAB", 3, 0.50, 0.50),
    ("MAR", "DAMA", 4, 0.50, 0.40),
    ("RIO", "CAB", 5, 0.50, 0.66),
    ("RIO", "DAMA", 6, 0.50, 0.52),
]

COLOR_CANONICAL = {"Púrpura": "Morado", "Morado": "Morado"}
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
KIDS_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_MIN_FILL = PatternFill("solid", fgColor="C6EFCE")
TOTAL_MAX_FILL = PatternFill("solid", fgColor="FFEB9C")
CORTE_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_CORTE_FILL = PatternFill("solid", fgColor="FFD966")
MIN_FONT = Font(color="0563C1")
MAX_FONT = Font(bold=True)
CORTE_FONT = Font(bold=True, color="C65911")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def canonical_color(c: str) -> str:
    return COLOR_CANONICAL.get(c, c)


def tela_stock(color: str) -> float:
    return TELA_KG.get(canonical_color(color), TELA_KG.get(color, 0.0))


def parse_colores(path: Path) -> dict[str, dict[str, int]]:
    df = pd.read_excel(path, sheet_name="Cantidades por Colores", header=None)
    result: dict[str, dict[str, int]] = {"CAB": {}, "DAMA": {}, "KIDS": {}}
    current = None
    for _, row in df.iterrows():
        val0 = str(row[0]) if pd.notna(row[0]) else ""
        u = val0.upper()
        if "CABALLERO" in u:
            current = "CAB"
        elif "DAMA" in u and "CAB" not in u:
            current = "DAMA"
        elif "KIDS" in u or "NIÑ" in u:
            current = "KIDS"
        elif val0.strip() in ("Color", "TOTAL", "nan", "") or not current:
            continue
        elif pd.notna(row[3]):
            try:
                result[current][val0.strip()] = int(float(row[3]))
            except (ValueError, TypeError):
                pass
    return result


def parse_consumo(path: Path) -> dict[str, float]:
    df = pd.read_excel(path, sheet_name="Compra de Tela", header=None)
    out: dict[str, float] = {}
    for _, row in df.iterrows():
        v0 = str(row[0]).strip() if pd.notna(row[0]) else ""
        if v0 in ("CABALLERO", "DAMA", "KIDS") and pd.notna(row[1]):
            out["CAB" if v0 == "CABALLERO" else v0] = float(row[1])
    return out


def parse_mix_full(path: Path) -> dict[tuple, dict]:
    """Mix min/max por (modelo, genero, color) desde Excel proyección."""
    df = pd.read_excel(path, sheet_name="Producción Color × Talla", header=None)
    modelo = "MAR" if "MAR" in path.name.upper() else "RIO"
    section = None
    tallas: list[str] = []
    data: dict[tuple, dict] = {}

    for _, row in df.iterrows():
        val0 = str(row[0]) if pd.notna(row[0]) else ""
        if "CABALLERO" in val0.upper() and "PRODUCCIÓN" in val0.upper():
            section = "CAB"
            continue
        if "DAMA" in val0.upper() and "PRODUCCIÓN" in val0.upper():
            section = "DAMA"
            continue
        if "KIDS" in val0.upper() and "PRODUCCIÓN" in val0.upper():
            section = "KIDS"
            continue
        if val0.strip() == "Color / Talla":
            tallas = [str(row[i]).replace(".0", "") for i in range(1, len(row))
                      if pd.notna(row[i]) and str(row[i]) != "Total"]
            continue
        if not section or ("Mín" not in val0 and "Máx" not in val0):
            continue
        color = val0.split("—")[0].strip()
        key = (modelo, section, color)
        bucket = data.setdefault(key, {
            "tallas": tallas[:], "min": {}, "max": {}, "total_min": 0, "total_max": 0,
        })
        kind = "min" if "Mín" in val0 else "max"
        total_col = len(tallas) + 1
        total = int(float(row[total_col])) if pd.notna(row[total_col]) else 0
        bucket[f"total_{kind}"] = total
        for i, t in enumerate(tallas):
            if pd.notna(row[i + 1]):
                bucket[kind][t] = int(float(row[i + 1]))
    return data


def lote_cortado_meta_kids() -> dict[str, float]:
    por = LOTE_CORTADO_TOTAL / len(LOTE_CORTADO_COLORES)
    out: dict[str, float] = {}
    for c in LOTE_CORTADO_COLORES:
        mc = LOTE_A_META_KIDS.get(c)
        if mc:
            out[mc] = out.get(mc, 0) + por
    return out


def load_all() -> tuple[dict, dict, dict]:
    mar, rio = parse_colores(MAR_XLSX), parse_colores(RIO_XLSX)
    metas = {"MAR": mar, "RIO": rio}
    consumos = {"MAR": parse_consumo(MAR_XLSX), "RIO": parse_consumo(RIO_XLSX)}
    mix = parse_mix_full(MAR_XLSX)
    mix.update(parse_mix_full(RIO_XLSX))
    return metas, consumos, mix


def build_plan_rows(metas: dict, consumos: dict) -> list[dict]:
    cortado = lote_cortado_meta_kids()
    cfg = {(m, g): (p, pct, cons) for m, g, p, pct, cons in LINEAS_CONFIG}
    rows: list[dict] = []
    for modelo, genero, prio, pct, cons_default in LINEAS_CONFIG:
        cons = consumos[modelo][genero]
        for color, meta in metas[modelo][genero].items():
            ya = cortado.get(color, 0) if modelo == "MAR" and genero == "KIDS" else 0
            rows.append({
                "prio": prio, "orden": f"{modelo} {genero}", "modelo": modelo,
                "genero": genero, "color": color, "meta": meta, "ya_cortado": ya,
                "pct": pct, "consumo": cons,
            })
    return rows


def style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def style_input(cell) -> None:
    cell.fill = INPUT_FILL
    cell.border = BORDER


def merge_block(ws, col: int, r1: int, r2: int) -> None:
    if r2 > r1:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
    cell = ws.cell(r1, col)
    cell.alignment = CENTER
    cell.border = BORDER


def auto_width(ws, mx: int = 20) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = min(
            max(len(str(c.value or "")) for c in col) + 2, mx)


def write_metas_sheet(ws, metas: dict) -> None:
    ws.append(["Modelo", "Género", "Color", "Meta máx Excel"])
    style_header(ws, 1, 4)
    row = 2
    for modelo in ("MAR", "RIO"):
        for gen in ("KIDS", "CAB", "DAMA"):
            colors = sorted(metas[modelo][gen].items(), key=lambda x: -x[1])
            if not colors:
                continue
            sec_start = row
            for color, mx in colors:
                ws.cell(row, 3, color)
                ws.cell(row, 4, mx)
                ws.cell(row, 3).border = BORDER
                ws.cell(row, 4).border = BORDER
                ws.cell(row, 4).alignment = Alignment(horizontal="right")
                row += 1
            sec_end = row - 1
            ws.cell(sec_start, 1, modelo)
            ws.cell(sec_start, 2, gen)
            merge_block(ws, 1, sec_start, sec_end)
            merge_block(ws, 2, sec_start, sec_end)
    auto_width(ws)


def write_mix_tallas_sheet(ws, metas: dict, mix: dict) -> None:
    ws.append(["Modelo", "Género", "Color", "Talla", "Und máx Excel", "Total color"])
    style_header(ws, 1, 6)
    row = 2
    for modelo in ("MAR", "RIO"):
        for gen in ("KIDS", "CAB", "DAMA"):
            colors = sorted(metas[modelo][gen].keys())
            if not colors:
                continue
            mod_start = row
            for color in colors:
                key = (modelo, gen, color)
                block = mix.get(key)
                if not block or not block["tallas"]:
                    continue
                col_start = row
                total_max = block["total_max"]
                for t in block["tallas"]:
                    ws.cell(row, 4, t)
                    ws.cell(row, 5, block["max"].get(t, 0))
                    for c in range(4, 6):
                        ws.cell(row, c).border = BORDER
                        ws.cell(row, c).alignment = CENTER
                    row += 1
                col_end = row - 1
                ws.cell(col_start, 3, color)
                ws.cell(col_start, 6, total_max)
                merge_block(ws, 3, col_start, col_end)
                merge_block(ws, 6, col_start, col_end)
                ws.cell(col_start, 6).font = Font(bold=True)
            mod_end = row - 1
            if mod_end >= mod_start:
                ws.cell(mod_start, 1, modelo)
                ws.cell(mod_start, 2, gen)
                merge_block(ws, 1, mod_start, mod_end)
                merge_block(ws, 2, mod_start, mod_end)
    auto_width(ws)


def _section_title(ws, row: int, title: str) -> int:
    ws.cell(row, 1, title)
    ws.cell(row, 1).font = Font(bold=True, size=11, color="1F4E79")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    return row + 2


def write_curva_ref_section(ws, row: int, modelo: str, genero: str, metas: dict, mix: dict) -> int:
    """Bloque MÍN/MÁX referencia Excel. Retorna siguiente fila libre."""
    row = _section_title(ws, row, f"▸ {modelo} {genero} — Curva referencia (Excel proyección)")
    colors = sorted(metas[modelo][genero].items(), key=lambda x: -x[1])
    if not colors:
        return row + 1
    sample = mix.get((modelo, genero, colors[0][0]))
    tallas = sample["tallas"] if sample else []
    if not tallas:
        return row + 1

    ws.cell(row, 1, "Color")
    for i, t in enumerate(tallas, 2):
        ws.cell(row, i, t)
    tot_col = len(tallas) + 2
    ws.cell(row, tot_col, "Total")
    style_header(ws, row, tot_col)
    row += 1
    min_rows, max_rows = [], []

    for color, _ in colors:
        block = mix.get((modelo, genero, color))
        if not block:
            continue
        ws.cell(row, 1, f"{color} — MÍNIMO")
        ws.cell(row, 1).font = MIN_FONT
        for i, t in enumerate(tallas, 2):
            ws.cell(row, i, block["min"].get(t, 0))
            ws.cell(row, i).font = MIN_FONT
            ws.cell(row, i).border = BORDER
            ws.cell(row, i).alignment = CENTER
        ws.cell(row, tot_col, f"=SUM(B{row}:{get_column_letter(len(tallas)+1)}{row})")
        min_rows.append(row)
        row += 1
        ws.cell(row, 1, f"{color} — MÁXIMO")
        ws.cell(row, 1).font = MAX_FONT
        for i, t in enumerate(tallas, 2):
            ws.cell(row, i, block["max"].get(t, 0))
            ws.cell(row, i).font = MAX_FONT
            ws.cell(row, i).border = BORDER
            ws.cell(row, i).alignment = CENTER
        ws.cell(row, tot_col, f"=SUM(B{row}:{get_column_letter(len(tallas)+1)}{row})")
        max_rows.append(row)
        row += 1

    if min_rows:
        ws.cell(row, 1, "TOTAL — MÍNIMO")
        ws.cell(row, 1).font = Font(bold=True)
        for i in range(2, tot_col + 1):
            col = get_column_letter(i)
            ws.cell(row, i, f"=SUM({','.join(f'{col}{r}' for r in min_rows)})")
            ws.cell(row, i).fill = TOTAL_MIN_FILL
            ws.cell(row, i).font = Font(bold=True)
            ws.cell(row, i).border = BORDER
        row += 1
        ws.cell(row, 1, "TOTAL — MÁXIMO")
        ws.cell(row, 1).font = Font(bold=True)
        for i in range(2, tot_col + 1):
            col = get_column_letter(i)
            ws.cell(row, i, f"=SUM({','.join(f'{col}{r}' for r in max_rows)})")
            ws.cell(row, i).fill = TOTAL_MAX_FILL
            ws.cell(row, i).font = Font(bold=True)
            ws.cell(row, i).border = BORDER
        row += 1
    return row + 1


def write_curva_corte_section(
    ws, row: int, modelo: str, genero: str, metas: dict, mix: dict,
    plan_row_map: dict[tuple, int], plan_start: int,
) -> int:
    """Curva de lo que SE CORTARÍA según PLAN COLORES (Und FINAL × mix)."""
    row = _section_title(ws, row, f"▸ {modelo} {genero} — Objetivo A CORTAR (desde PLAN COLORES)")
    colors = sorted(metas[modelo][genero].items(), key=lambda x: -x[1])
    if not colors:
        return row + 1
    sample = mix.get((modelo, genero, colors[0][0]))
    tallas = sample["tallas"] if sample else []
    if not tallas:
        return row + 1

    # Col A=color, B=Und FINAL (ref plan), C+=tallas, last=total check
    ws.cell(row, 1, "Color")
    ws.cell(row, 2, "Und FINAL")
    for i, t in enumerate(tallas, 3):
        ws.cell(row, i, t)
    tot_col = len(tallas) + 3
    ws.cell(row, tot_col, "Total")
    style_header(ws, row, tot_col)
    row += 1
    data_rows: list[int] = []

    for color, _ in colors:
        block = mix.get((modelo, genero, color))
        plan_r = plan_row_map.get((modelo, genero, color))
        if not block or not plan_r:
            continue
        ws.cell(row, 1, f"{color} — A CORTAR")
        ws.cell(row, 1).font = CORTE_FONT
        ws.cell(row, 2, f"='PLAN COLORES'!P{plan_r}")
        ws.cell(row, 2).fill = CORTE_FILL
        ws.cell(row, 2).border = BORDER
        for i, t in enumerate(tallas, 3):
            mix_val = block["max"].get(t, 0)
            total = block["total_max"] or 1
            col_b = f"$B{row}"
            ws.cell(row, i, f'=IF({col_b}=0,0,ROUND({col_b}*{mix_val}/{total},0))')
            ws.cell(row, i).border = BORDER
            ws.cell(row, i).alignment = CENTER
        first_t = get_column_letter(3)
        last_t = get_column_letter(len(tallas) + 2)
        ws.cell(row, tot_col, f"=SUM({first_t}{row}:{last_t}{row})")
        ws.cell(row, tot_col).font = Font(bold=True)
        data_rows.append(row)
        row += 1

    if data_rows:
        ws.cell(row, 1, "TOTAL — A CORTAR")
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 2, f"=SUM({','.join(f'B{r}' for r in data_rows)})")
        ws.cell(row, 2).fill = TOTAL_CORTE_FILL
        for i in range(3, tot_col + 1):
            col = get_column_letter(i)
            ws.cell(row, i, f"=SUM({','.join(f'{col}{r}' for r in data_rows)})")
            ws.cell(row, i).fill = TOTAL_CORTE_FILL
            ws.cell(row, i).font = Font(bold=True)
            ws.cell(row, i).border = BORDER
        row += 1
    return row + 1


def write_workbook(path: Path, metas: dict, consumos: dict, mix: dict) -> dict:
    wb = Workbook()
    cortado = lote_cortado_meta_kids()
    plan_rows = build_plan_rows(metas, consumos)
    refs: dict = {}

    # INSTRUCCIONES
    ws0 = wb.active
    ws0.title = "INSTRUCCIONES"
    for i, t in enumerate([
        "PLAN MAR / RIO — Guía",
        "",
        "• PARAMETROS: edita % objetivo y consumo por MODELO + GÉNERO (amarillo).",
        "• Inventario kg y und ya cortadas por modelo/género/color.",
        "• METAS: referencia agrupada como Excel proyección.",
        "• MIX TALLAS: mix por color (referencia Excel).",
        "• PLAN COLORES: und final por color (fórmulas).",
        "• CURVAS REFERENCIA: MÍN/MÁX del Excel proyección (todas las líneas).",
        "• CURVAS A CORTAR: lo que se cortaría hoy según PLAN × mix talla.",
    ], 1):
        ws0.cell(i, 1, t)
        if i == 1:
            ws0.cell(i, 1).font = Font(bold=True, size=14)

    # PARAMETROS
    ws_p = wb.create_sheet("PARAMETROS")
    ws_p["A1"] = "PARÁMETROS POR MODELO Y GÉNERO (editable)"
    ws_p["A1"].font = Font(bold=True, size=12)
    lin_start = 3
    ws_p.cell(lin_start, 1, "Modelo")
    ws_p.cell(lin_start, 2, "Género")
    ws_p.cell(lin_start, 3, "% Objetivo")
    ws_p.cell(lin_start, 4, "Consumo kg/und")
    ws_p.cell(lin_start, 5, "Prioridad (P)")
    style_header(ws_p, lin_start, 5)
    lr = lin_start + 1
    line_cfg_rows: dict[tuple, int] = {}
    for modelo, genero, prio, pct, cons in LINEAS_CONFIG:
        ws_p.cell(lr, 1, modelo)
        ws_p.cell(lr, 2, genero)
        ws_p.cell(lr, 3, pct)
        ws_p.cell(lr, 4, consumos[modelo][genero])
        ws_p.cell(lr, 5, prio)
        for c in (3, 4):
            style_input(ws_p.cell(lr, c))
        ws_p.cell(lr, 3).number_format = "0%"
        line_cfg_rows[(modelo, genero)] = lr
        lr += 1
    lin_end = lr - 1
    refs["lin_start"], refs["lin_end"] = lin_start, lin_end

    # Inventario
    inv_hdr = lr + 2
    ws_p.cell(inv_hdr, 1, "INVENTARIO TELA JABÓN (kg) — foto actual")
    ws_p.cell(inv_hdr, 1).font = Font(bold=True)
    inv_start = inv_hdr + 2
    ws_p.cell(inv_start, 1, "Color")
    ws_p.cell(inv_start, 2, "Inventario kg")
    style_header(ws_p, inv_start, 2)
    ir = inv_start + 1
    for color in sorted(TELA_KG, key=lambda c: -TELA_KG[c]):
        ws_p.cell(ir, 1, color)
        ws_p.cell(ir, 2, round(TELA_KG[color], 2))
        style_input(ws_p.cell(ir, 2))
        ir += 1
    inv_end = ir - 1
    refs["inv_start"], refs["inv_end"] = inv_start, inv_end

    # Und ya cortadas — todos modelos/géneros
    cort_hdr = ir + 2
    ws_p.cell(cort_hdr, 1, "UND YA CORTADAS (solo descuenta meta — NO resta inventario kg)")
    ws_p.cell(cort_hdr, 1).font = Font(bold=True)
    cort_start = cort_hdr + 2
    ws_p.cell(cort_start, 1, "Modelo")
    ws_p.cell(cort_start, 2, "Género")
    ws_p.cell(cort_start, 3, "Color")
    ws_p.cell(cort_start, 4, "Und cortadas")
    style_header(ws_p, cort_start, 4)
    cr = cort_start + 1
    for modelo in ("MAR", "RIO"):
        for gen in ("KIDS", "CAB", "DAMA"):
            for color in sorted(metas[modelo][gen].keys()):
                val = cortado.get(color, 0) if modelo == "MAR" and gen == "KIDS" else 0
                ws_p.cell(cr, 1, modelo)
                ws_p.cell(cr, 2, gen)
                ws_p.cell(cr, 3, color)
                ws_p.cell(cr, 4, round(val, 1) if val else 0)
                style_input(ws_p.cell(cr, 4))
                cr += 1
    cort_end = cr - 1
    refs["cort_start"], refs["cort_end"] = cort_start, cort_end
    auto_width(ws_p)

    # METAS
    ws_m = wb.create_sheet("METAS")
    write_metas_sheet(ws_m, metas)

    # MIX TALLAS
    ws_mix = wb.create_sheet("MIX TALLAS")
    write_mix_tallas_sheet(ws_mix, metas, mix)

    # PLAN COLORES (antes de curvas a cortar — las fórmulas lo referencian)
    ws_plan = wb.create_sheet("PLAN COLORES")
    headers = [
        "P", "Línea", "Modelo", "Género", "Color", "Meta máx", "Ya cortado", "Pendiente",
        "% Obj", "Und objetivo", "Consumo kg/und", "Inventario kg", "Kg usado antes",
        "Saldo kg", "Und max tela", "Und FINAL", "Kg usar", "Estado",
    ]
    ws_plan.append(headers)
    style_header(ws_plan, 1, len(headers))
    plan_start = 2

    def pct_formula(r: int) -> str:
        a, b, c = lin_start + 1, lin_end, lin_start + 1
        return (
            f'=IF(SUMPRODUCT((PARAMETROS!$A${a}:$A${b}=C{r})*(PARAMETROS!$B${a}:$B${b}=D{r}))=0,0.5,'
            f'SUMPRODUCT((PARAMETROS!$A${a}:$A${b}=C{r})*(PARAMETROS!$B${a}:$B${b}=D{r})*PARAMETROS!$C${c}:$C${b}))'
        )

    def cons_formula(r: int) -> str:
        a, b, d = lin_start + 1, lin_end, lin_start + 1
        return (
            f'=IF(SUMPRODUCT((PARAMETROS!$A${a}:$A${b}=C{r})*(PARAMETROS!$B${a}:$B${b}=D{r}))=0,0.3,'
            f'SUMPRODUCT((PARAMETROS!$A${a}:$A${b}=C{r})*(PARAMETROS!$B${a}:$B${b}=D{r})*PARAMETROS!$D${d}:$D${b}))'
        )

    def inv_formula(color_cell: str) -> str:
        return (
            f'=IFERROR(INDEX(PARAMETROS!$B${inv_start+1}:$B${inv_end},'
            f'MATCH({color_cell},PARAMETROS!$A${inv_start+1}:$A${inv_end},0)),'
            f'IFERROR(INDEX(PARAMETROS!$B${inv_start+1}:$B${inv_end},'
            f'MATCH("Morado",PARAMETROS!$A${inv_start+1}:$A${inv_end},0)),0))'
        )

    def cortado_formula(r: int) -> str:
        a, b = cort_start + 1, cort_end
        return (
            f'=SUMPRODUCT((PARAMETROS!$A${a}:$A${b}=C{r})*(PARAMETROS!$B${a}:$B${b}=D{r})*'
            f'(PARAMETROS!$C${a}:$C${b}=E{r})*PARAMETROS!$D${a}:$D${b})'
        )

    for i, pr in enumerate(plan_rows):
        r = plan_start + i
        ws_plan.cell(r, 1, pr["prio"])
        ws_plan.cell(r, 2, pr["orden"])
        ws_plan.cell(r, 3, pr["modelo"])
        ws_plan.cell(r, 4, pr["genero"])
        ws_plan.cell(r, 5, pr["color"])
        ws_plan.cell(r, 6, pr["meta"])
        ws_plan.cell(r, 7, cortado_formula(r))
        ws_plan.cell(r, 8, f"=MAX(0,F{r}-G{r})")
        ws_plan.cell(r, 9, pct_formula(r))
        ws_plan.cell(r, 9).number_format = "0%"
        ws_plan.cell(r, 10, f"=ROUND(H{r}*I{r},0)")
        ws_plan.cell(r, 11, cons_formula(r))
        ws_plan.cell(r, 12, inv_formula(f"E{r}"))
        ws_plan.cell(r, 13, 0 if r == plan_start else f'=SUMIFS($Q${plan_start}:$Q{r-1},$E${plan_start}:$E{r-1},E{r})')
        ws_plan.cell(r, 14, f"=MAX(0,L{r}-M{r})")
        ws_plan.cell(r, 15, f"=IF(N{r}<{MIN_KG},0,FLOOR(N{r}/K{r},1))")
        ws_plan.cell(r, 16, f"=MIN(J{r},O{r})")
        ws_plan.cell(r, 17, f"=P{r}*K{r}")
        ws_plan.cell(r, 18, f'=IF(P{r}=0,"SIN TELA",IF(P{r}<J{r},"PARCIAL","OK"))')
        if pr["genero"] == "KIDS":
            for c in range(1, len(headers) + 1):
                ws_plan.cell(r, c).fill = KIDS_FILL

    plan_end = plan_start + len(plan_rows) - 1
    refs["plan_start"], refs["plan_end"] = plan_start, plan_end
    plan_row_map = {
        (pr["modelo"], pr["genero"], pr["color"]): plan_start + i
        for i, pr in enumerate(plan_rows)
    }

    # CURVAS REFERENCIA — una sola pestaña, todas las líneas
    ws_ref = wb.create_sheet("CURVAS REFERENCIA")
    ws_ref["A1"] = "CURVAS DE REFERENCIA — MÍN / MÁX (Excel proyección)"
    ws_ref["A1"].font = Font(bold=True, size=13)
    rr = 3
    for modelo, genero, *_ in LINEAS_CONFIG:
        rr = write_curva_ref_section(ws_ref, rr, modelo, genero, metas, mix)
    auto_width(ws_ref)

    # CURVAS A CORTAR — objetivo real repartido por talla
    ws_corte = wb.create_sheet("CURVAS A CORTAR")
    ws_corte["A1"] = "CURVAS A CORTAR — Objetivo actual según PLAN COLORES × mix talla"
    ws_corte["A1"].font = Font(bold=True, size=13)
    ws_corte["A2"] = "Edita PARAMETROS o PLAN COLORES → estas curvas recalculan solas (columna B = Und FINAL)."
    cr = 4
    for modelo, genero, *_ in LINEAS_CONFIG:
        cr = write_curva_corte_section(
            ws_corte, cr, modelo, genero, metas, mix, plan_row_map, plan_start,
        )
    auto_width(ws_corte)

    # RESUMEN
    ws_r = wb.create_sheet("RESUMEN")
    ws_r["A1"] = "RESUMEN POR LÍNEA"
    ws_r["A1"].font = Font(bold=True, size=12)
    ws_r.cell(3, 1, "Línea")
    ws_r.cell(3, 2, "Und FINAL")
    ws_r.cell(3, 3, "Kg usar")
    style_header(ws_r, 3, 3)
    for i, etq in enumerate(PRIORIDAD_ORDEN, 4):
        ws_r.cell(i, 1, etq)
        ws_r.cell(i, 2, f"=SUMIF('PLAN COLORES'!$B${plan_start}:$B${plan_end},\"{etq}\",'PLAN COLORES'!$P${plan_start}:$P${plan_end})")
        ws_r.cell(i, 3, f"=SUMIF('PLAN COLORES'!$B${plan_start}:$B${plan_end},\"{etq}\",'PLAN COLORES'!$Q${plan_start}:$Q${plan_end})")
    tr = len(PRIORIDAD_ORDEN) + 5
    ws_r.cell(tr, 1, "TOTAL")
    ws_r.cell(tr, 2, f"=SUM(B4:B{tr-2})")
    ws_r.cell(tr, 3, f"=SUM(C4:C{tr-2})")
    auto_width(ws_r)

    wb.save(path)
    return refs


def eval_plan_python(metas: dict, consumos: dict) -> dict:
    cortado = lote_cortado_meta_kids()
    kg_usado: dict[str, float] = {}
    results = []
    for pr in build_plan_rows(metas, consumos):
        ya = cortado.get(pr["color"], 0) if pr["modelo"] == "MAR" and pr["genero"] == "KIDS" else 0
        pend = max(0, pr["meta"] - ya)
        und_obj = int(pend * pr["pct"])
        ck = canonical_color(pr["color"])
        saldo = max(0, tela_stock(pr["color"]) - kg_usado.get(ck, 0))
        und_final = min(und_obj, int(saldo / pr["consumo"]) if saldo >= MIN_KG else 0)
        kg = und_final * pr["consumo"]
        kg_usado[ck] = kg_usado.get(ck, 0) + kg
        results.append({**pr, "und_final": und_final, "kg": round(kg, 2)})
    by_line = {e: {"und": sum(r["und_final"] for r in results if r["orden"] == e),
                   "kg": round(sum(r["kg"] for r in results if r["orden"] == e), 2)}
               for e in PRIORIDAD_ORDEN}
    return {"filas": results, "por_linea": by_line,
            "total_und": sum(r["und_final"] for r in results),
            "total_kg": round(sum(r["kg"] for r in results), 2)}


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    metas, consumos, mix = load_all()
    path = OUT_DIR / "PLAN_PRIORIDADES_MAR_RIO_KIDS.xlsx"
    write_workbook(path, metas, consumos, mix)
    summary = eval_plan_python(metas, consumos)
    (OUT_DIR / "plan_prioridades.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(f"Excel: {path}")
    print("Hojas (8): INSTRUCCIONES | PARAMETROS | METAS | MIX TALLAS | PLAN COLORES | CURVAS REF | CURVAS CORTE | RESUMEN")
    for e in PRIORIDAD_ORDEN:
        d = summary["por_linea"][e]
        print(f"  {e}: {d['und']} und / {d['kg']} kg")


if __name__ == "__main__":
    main()
