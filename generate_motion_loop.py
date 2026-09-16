#!/usr/bin/env python3
"""Generate Motion Loop production suggestion file and analysis."""
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROCESO_PATH = "/home/ubuntu/.cursor/projects/workspace/uploads/MOTION_LOOP_EN_PROCESO_5753.xlsx"
INV_PATH = "/home/ubuntu/.cursor/projects/workspace/uploads/MOOTION_LOOP_INVENTARIO_ARREGLADO_3b55.xlsx"
OUT_PATH = "/workspace/MOTION_LOOP_SUGERENCIA_PRODUCCION.xlsx"
ANALYSIS_PATH = "/workspace/motion_loop_analysis.json"


def calc_priority(row):
    if row["FALTANTE"] == 0:
        return "COMPLETO"
    score = 0
    if row["INVENTARIO_TALLER"] == 0:
        score += 3
    elif row["INVENTARIO_TALLER"] < 10:
        score += 2
    if row["TALLA"] in ["S", "M"]:
        score += 2
    if row["FALTANTE"] >= 24:
        score += 3
    elif row["FALTANTE"] >= 15:
        score += 2
    elif row["FALTANTE"] >= 10:
        score += 1
    if row["GENERO"] == "DAMA" and row["TALLA"] in ["XS", "S", "M"]:
        score += 1
    if score >= 6:
        return "ALTA"
    if score >= 4:
        return "MEDIA"
    return "BAJA"


def main():
    proceso = pd.read_excel(PROCESO_PATH)
    inv = pd.read_excel(INV_PATH)

    proceso = proceso.rename(columns={"PROUDUCIDO": "PRODUCIDO"})
    proceso["GENERO"] = proceso["GENERO"].str.upper().str.strip()

    inv_by_sku = inv.groupby("SKU")["Cantidad en inventario"].sum().reset_index()
    inv_by_sku.columns = ["SKU", "INVENTARIO_TALLER"]

    df = proceso.merge(inv_by_sku, on="SKU", how="left")
    df["INVENTARIO_TALLER"] = df["INVENTARIO_TALLER"].fillna(0).astype(int)
    df["PRODUCIDO"] = df["PRODUCIDO"].fillna(0).astype(int)

    size_pct_cab = proceso[proceso["GENERO"] == "CAB"].groupby("TALLA")["CANTIDAD"].sum()
    size_pct_cab = (size_pct_cab / size_pct_cab.sum() * 100).round(1)
    size_pct_dama = proceso[proceso["GENERO"] == "DAMA"].groupby("TALLA")["CANTIDAD"].sum()
    size_pct_dama = (size_pct_dama / size_pct_dama.sum() * 100).round(1)

    def pct_talla(row):
        if row["GENERO"] == "CAB":
            return float(size_pct_cab.get(row["TALLA"], 0))
        return float(size_pct_dama.get(row["TALLA"], 0))

    df["PCT_TALLA_HIST"] = df.apply(pct_talla, axis=1)
    df["NECESIDAD_TOTAL"] = (df["CANTIDAD"] - df["INVENTARIO_TALLER"]).clip(lower=0).astype(int)
    df["SUGERENCIA_PRODUCCION"] = df["FALTANTE"].astype(int)
    df["PRIORIDAD"] = df.apply(calc_priority, axis=1)
    df["SUG_REP_T1"] = np.ceil(df["SUGERENCIA_PRODUCCION"] / 2).astype(int)
    df["SUG_REP_T2"] = df["SUGERENCIA_PRODUCCION"] - df["SUG_REP_T1"]

    prio_order = {"ALTA": 0, "MEDIA": 1, "BAJA": 2, "COMPLETO": 3}
    df["_sort"] = df["PRIORIDAD"].map(prio_order)
    df = df.sort_values(["_sort", "GENERO", "PRODUCTO", "COLOR", "TALLA"]).drop("_sort", axis=1)

    out_cols = [
        "SKU",
        "PRODUCTO",
        "GENERO",
        "COLOR",
        "TALLA",
        "CANTIDAD",
        "PRODUCIDO",
        "FALTANTE",
        "INVENTARIO_TALLER",
        "PCT_TALLA_HIST",
        "NECESIDAD_TOTAL",
        "SUGERENCIA_PRODUCCION",
        "PRIORIDAD",
        "SUG_REP_T1",
        "SUG_REP_T2",
    ]
    output = df[out_cols].copy()
    output.to_excel(OUT_PATH, index=False, sheet_name="Sugerencia Producción")

    wb = load_workbook(OUT_PATH)
    ws = wb.active
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    fills = {
        "ALTA": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
        "MEDIA": PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid"),
        "BAJA": PatternFill(start_color="6BCB77", end_color="6BCB77", fill_type="solid"),
        "COMPLETO": PatternFill(start_color="4D96FF", end_color="4D96FF", fill_type="solid"),
    }

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    prio_col = out_cols.index("PRIORIDAD") + 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        prio = row[prio_col - 1].value
        if prio in fills:
            row[prio_col - 1].fill = fills[prio]
            row[prio_col - 1].font = Font(bold=True)

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 13
    ws.column_dimensions["B"].width = 22
    wb.save(OUT_PATH)

    with pd.ExcelWriter(OUT_PATH, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df.groupby("GENERO").agg(
            PEDIDO=("CANTIDAD", "sum"),
            PRODUCIDO=("PRODUCIDO", "sum"),
            FALTANTE=("FALTANTE", "sum"),
            INVENTARIO=("INVENTARIO_TALLER", "sum"),
            SUGERENCIA=("SUGERENCIA_PRODUCCION", "sum"),
            NECESIDAD=("NECESIDAD_TOTAL", "sum"),
        ).reset_index().to_excel(writer, sheet_name="Resumen Género", index=False)

        df.groupby(["GENERO", "PRIORIDAD"]).agg(
            SKUs=("SKU", "count"),
            SUGERENCIA=("SUGERENCIA_PRODUCCION", "sum"),
        ).reset_index().to_excel(writer, sheet_name="Resumen Prioridad", index=False)

        talla_rows = []
        for genero, pct in [("CAB", size_pct_cab), ("DAMA", size_pct_dama)]:
            for talla, p in pct.items():
                talla_rows.append({"GENERO": genero, "TALLA": talla, "PCT_PLAN": p})
        pd.DataFrame(talla_rows).to_excel(writer, sheet_name="Curva Tallas", index=False)

    analysis = {
        "totals": {
            "faltante": int(df["FALTANTE"].sum()),
            "sugerencia": int(df["SUGERENCIA_PRODUCCION"].sum()),
            "inventario": int(df["INVENTARIO_TALLER"].sum()),
            "pedido": int(df["CANTIDAD"].sum()),
            "necesidad": int(df["NECESIDAD_TOTAL"].sum()),
        },
        "size_cab": {k: float(v) for k, v in size_pct_cab.items()},
        "size_dama": {k: float(v) for k, v in size_pct_dama.items()},
        "by_genero": {
            g: {
                "faltante": int(r.FALTANTE),
                "sugerencia": int(r.SUGERENCIA_PRODUCCION),
                "inventario": int(r.INVENTARIO_TALLER),
                "pedido": int(r.CANTIDAD),
            }
            for g, r in df.groupby("GENERO")
            .agg(
                FALTANTE=("FALTANTE", "sum"),
                SUGERENCIA_PRODUCCION=("SUGERENCIA_PRODUCCION", "sum"),
                INVENTARIO_TALLER=("INVENTARIO_TALLER", "sum"),
                CANTIDAD=("CANTIDAD", "sum"),
            )
            .iterrows()
        },
        "by_color_cab": df[df["GENERO"] == "CAB"].groupby("COLOR")["SUGERENCIA_PRODUCCION"].sum().to_dict(),
        "by_color_dama": df[df["GENERO"] == "DAMA"].groupby("COLOR")["SUGERENCIA_PRODUCCION"].sum().to_dict(),
        "by_prioridad": df.groupby("PRIORIDAD")["SUGERENCIA_PRODUCCION"].sum().to_dict(),
        "alta": output[output["PRIORIDAD"] == "ALTA"][
            ["SKU", "GENERO", "COLOR", "TALLA", "FALTANTE", "INVENTARIO_TALLER", "SUGERENCIA_PRODUCCION"]
        ].to_dict("records"),
        "rows": output.to_dict("records"),
    }
    with open(ANALYSIS_PATH, "w", encoding="utf-8") as f:
        json.dump(analysis, f, ensure_ascii=False, indent=2)

    print(f"Generated {OUT_PATH}")
    print(f"FALTANTE total: {df['FALTANTE'].sum()} | Inventario: {df['INVENTARIO_TALLER'].sum()}")


if __name__ == "__main__":
    main()
