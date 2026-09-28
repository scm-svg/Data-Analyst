#!/usr/bin/env python3
"""Propuesta Black Friday — inventario de baja rotación.

La lista sale del cuadro de rotación (SKU de cola, con stock, manufactura
y equipamiento). Jacket 1.0 y Short Sport entran con el inventario completo
de la línea. Zenit, Refresh y Motion Loop quedan fuera. El tablero y el
Excel hablan en unidades, rotación y cobertura.
"""

from __future__ import annotations

import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path("/workspace/propuesta-black-friday")
HTML_TEMPLATE = ROOT / "dashboard.template.html"
HTML_OUT = ROOT / "propuesta-black-friday-categoria-c.html"
XLSX_OUT = ROOT / "Propuesta_Black_Friday_Categoria_C.xlsx"
ABC_HTML = Path("/home/ubuntu/.cursor/projects/workspace/uploads/abc_ver_5c4b.html")
REF_HTML = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/black_friday_propuesta_categoria_c_0e74.html"
)
INV_XLSX = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/INVENTARIO_TOTAL_CUADRO_PARA_ABC_-_PROPUESTA_dc6f.xlsx"
)
VEN_XLSX = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/VENTAS_CUADRO_ACTUALIZADAS_ARREGLADO_4e9d.xlsx"
)

STORES = [
    "La Grieta",
    "Sambil Valencia",
    "Cerro Verde",
    "Sambil Chacao",
    "La Vela",
    "Grandplaz",
    "Tolon",
]
STORE_MAP = {
    "CERRO VERDE": "Cerro Verde",
    "CHACAO": "Sambil Chacao",
    "SAMBIL": "Sambil Valencia",
    "GRIETA": "La Grieta",
    "VELA": "La Vela",
    "GRANDPLAZ": "Grandplaz",
    "TOLON": "Tolon",
}
PERIODS = [
    "2025-10", "2025-11", "2025-12",
    "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06",
    "2026-07", "2026-08", "2026-09",
]
FLOOR = 1.10
MIN_DISC = 0.10
MIN_MODEL_STOCK = 5

# Líneas que el corte CC deja casi vacías y sí tienen bulto para ofrecer.
FORCE_FULL = {
    "JACKET CAB",
    "JACKET DAMA",
    "JACKET KIDS",
    "SHORT SPORT KIDS",
    "SHORT SPORT R1 DAMA",
    "SHORT SPORT R1 CAB",
}
DROP_RE = re.compile(r"ZENIT|REFRESH|MOTION LOOP", re.I)

TIERS = [
    ("<6", "Menos de 6 meses", 0.15),
    ("6-12", "6 a 12 meses", 0.20),
    ("12-18", "12 a 18 meses", 0.25),
    ("18-24", "18 a 24 meses", 0.30),
    ("24+", "Más de 24 meses", 0.35),
    ("dormido", "Sin salida en el año", 0.40),
]
TIER_PCT = {t[0]: t[2] for t in TIERS}


def load_json_script(path: Path, element_id: str):
    text = path.read_text(encoding="utf-8", errors="replace")
    marker = f'id="{element_id}"'
    start = text.find(">", text.find(marker)) + 1
    end = text.find("</script>", start)
    return json.loads(text[start:end])


def load_abc():
    data = load_json_script(ABC_HTML, "abc-embedded-data")
    skus = data["skus"]
    rows = []
    for r in data["salesRows"]:
        rows.append((skus[r[1]].upper(), float(r[4]), float(r[5]), float(r[6])))
    pos = pd.DataFrame(rows, columns=["sku", "qty", "rev", "cost"])
    pos = pos[pos.qty > 0]
    eco = pos.groupby("sku").agg(qty=("qty", "sum"), rev=("rev", "sum"), cost=("cost", "sum"))
    eco = eco[eco.qty > 0].copy()
    eco["uprice"] = eco.rev / eco.qty
    eco["ucost"] = eco.cost / eco.qty
    eco.loc[eco.uprice <= 0, "uprice"] = np.nan
    eco.loc[eco.ucost < 0, "ucost"] = np.nan
    sm = pd.DataFrame.from_dict(data["skuMaster"], orient="index").reset_index().rename(columns={"index": "sku"})
    sm["sku"] = sm.sku.str.upper()
    return eco, sm


def macro_of(cat):
    if cat is None or (isinstance(cat, float) and math.isnan(cat)):
        return None
    s = str(cat)
    if s.startswith("Manufactura"):
        return "Manufactura"
    if s.startswith("Equipamiento"):
        return "Equipamiento"
    if s.startswith("Materia"):
        return "Materia Prima"
    return None


def infer_category(modelo: str):
    name = str(modelo or "").upper()
    if re.search(r"STICKER|BOLSA|CINTA|KRAFT|NAVIDAD", name):
        return None, None
    if "PELOTA" in name:
        return "Equipamiento", "Equipamiento / Pelotas"
    if re.search(
        r"VITA |MOTION LOOP|BASIC LINE|MANGA LARGA|FALDA |VESTIDO|SHORT |PANT|BIKER|LEGGING|BIKINI|TOP |HOODIE|JACKET|POLO |ZENIT|REFRESH",
        name,
    ):
        return "Manufactura", "Manufactura / General"
    return None, None


def step5(p: float) -> float:
    if p <= 0:
        return 0.0
    return math.floor(p * 20 + 1e-9) / 20.0


def max_discount(price, cost) -> float:
    if price is None or not np.isfinite(price) or price <= 0:
        return 0.0
    if cost is None or not np.isfinite(cost) or cost < 0:
        return 0.0
    if cost == 0:
        return 0.70
    return max(0.0, 1.0 - FLOOR * float(cost) / float(price))


def sku_discount(price, cost, tier: float):
    cap = max_discount(price, cost)
    applied = step5(min(tier, cap))
    if applied + 1e-9 < MIN_DISC:
        return 0.0, False
    return applied, applied + 1e-9 < tier - 1e-9


def num(x, nd=2):
    if x is None:
        return None
    try:
        v = float(x)
    except (TypeError, ValueError):
        return None
    if not np.isfinite(v):
        return None
    return round(v, nd)


def clean(obj):
    if isinstance(obj, dict):
        return {k: clean(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [clean(v) for v in obj]
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating, float)):
        if not np.isfinite(obj):
            return None
        return float(obj)
    if isinstance(obj, (np.bool_,)):
        return bool(obj)
    if obj is pd.NA:
        return None
    return obj


def wavg(values, weights):
    v = np.asarray(values, dtype=float)
    w = np.asarray(weights, dtype=float)
    mask = np.isfinite(v) & np.isfinite(w) & (w > 0)
    if not mask.any():
        return None
    return float(np.average(v[mask], weights=w[mask]))


def rango_of(qty: float, cover: float) -> str:
    if qty <= 0:
        return "dormido"
    if cover < 6:
        return "<6"
    if cover < 12:
        return "6-12"
    if cover < 18:
        return "12-18"
    if cover < 24:
        return "18-24"
    return "24+"


def build_frames():
    eco, sm = load_abc()
    inv = pd.read_excel(INV_XLSX)
    ven = pd.read_excel(VEN_XLSX)
    for df, cols in (
        (inv, ["SKU", "MODELO", "GENERO", "COLOR", "TALLA", "Ubicación"]),
        (ven, ["SKU", "modelo", "GENERO", "COLOR", "TALLA", "tienda / ubicación"]),
    ):
        for c in cols:
            df[c] = df[c].astype(str).str.strip()
    inv["SKU"] = inv.SKU.str.upper()
    ven["SKU"] = ven.SKU.str.upper()
    inv["MODELO"] = inv.MODELO.replace({"None": np.nan, "nan": np.nan})
    ven["modelo"] = ven["modelo"].replace({"None": np.nan, "nan": np.nan})

    months = {
        "octubre": 10, "noviembre": 11, "diciembre": 12,
        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
        "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9,
    }
    ven["mes_n"] = ven["Mes"].str.lower().map(months)
    ven["period"] = ven["Año"].astype(str) + "-" + ven["mes_n"].astype(int).astype(str).str.zfill(2)
    ven = ven[ven.period.isin(PERIODS)].copy()
    ven["canal"] = np.where(ven["tienda / ubicación"].isin(["Pedidos", "CORPORATIVO"]), "mayor", "consumo")
    ven["tienda"] = ven["tienda / ubicación"].where(ven["tienda / ubicación"].isin(STORES))

    sm = sm.copy()
    sm["macro_real"] = sm.categoria.map(macro_of)
    attr_inv = inv.groupby("SKU").agg(
        modelo=("MODELO", "first"), genero=("GENERO", "first"), color=("COLOR", "first"), talla=("TALLA", "first")
    )
    attr_ven = ven.groupby("SKU").agg(
        modelo=("modelo", "first"), genero=("GENERO", "first"), color=("COLOR", "first"), talla=("TALLA", "first")
    )
    sm_i = sm.set_index("sku")
    all_skus = sorted(set(attr_inv.index) | set(attr_ven.index) | set(sm_i.index))
    base = pd.DataFrame(index=pd.Index(all_skus, name="sku"))
    base["modelo"] = attr_inv["modelo"].reindex(base.index).fillna(attr_ven["modelo"]).fillna(sm_i["modelo"])
    base["genero"] = attr_inv["genero"].reindex(base.index).fillna(attr_ven["genero"]).fillna(sm_i["genero"])
    base["color"] = attr_inv["color"].reindex(base.index).fillna(attr_ven["color"]).fillna(sm_i["color"])
    base["talla"] = attr_inv["talla"].reindex(base.index).fillna(attr_ven["talla"]).fillna(sm_i["talla"])
    base["categoria"] = sm_i["categoria"].reindex(base.index)
    base["macro"] = base["categoria"].map(macro_of)

    known = base[base["macro"].notna()]
    mode_cat = known.groupby("modelo")["categoria"].agg(lambda s: s.value_counts().index[0])
    mode_macro = known.groupby("modelo")["macro"].agg(lambda s: s.value_counts().index[0])
    inherit = base["macro"].isna() & base["modelo"].isin(mode_macro.index)
    base.loc[inherit, "categoria"] = base.loc[inherit, "modelo"].map(mode_cat)
    base.loc[inherit, "macro"] = base.loc[inherit, "modelo"].map(mode_macro)

    still = base["macro"].isna()
    inferred = [infer_category(m) for m in base.loc[still, "modelo"]]
    base.loc[still, "macro"] = [m for m, _ in inferred]
    base.loc[still, "categoria"] = [c for _, c in inferred]
    base["macro"] = base["macro"].fillna("Fuera de alcance")
    base["categoria"] = base["categoria"].fillna("Sin categoría")

    base["uprice"] = eco["uprice"].reindex(base.index)
    base["ucost"] = eco["ucost"].reindex(base.index)
    med_p = base.groupby("modelo")["uprice"].transform("median")
    med_c = base.groupby("modelo")["ucost"].transform("median")
    fill_p = base.uprice.isna() & med_p.notna()
    base.loc[fill_p, "uprice"] = med_p[fill_p]
    fill_c = base.ucost.isna() & med_c.notna()
    base.loc[fill_c, "ucost"] = med_c[fill_c]
    base["priced"] = base.uprice.notna() & base.ucost.notna() & (base.uprice > 0) & (base.ucost < base.uprice)

    cons = ven[ven.canal == "consumo"]
    base["qty_12"] = cons.groupby("SKU")["Cant. ordenada"].sum().reindex(base.index).fillna(0).clip(lower=0)
    inv["tienda"] = inv["Ubicación"].map(STORE_MAP)
    inv["pool"] = np.where(inv["Ubicación"].isin(["TALLER MANUFACTURADO", "TALLER EQUIPAMIENTO"]), "taller", "tienda")
    base["stock_tienda"] = inv[inv.pool == "tienda"].groupby("SKU")["Cantidad en inventario"].sum().reindex(base.index).fillna(0)
    base["stock_taller"] = inv[inv.pool == "taller"].groupby("SKU")["Cantidad en inventario"].sum().reindex(base.index).fillna(0)
    base["stock"] = base["stock_tienda"] + base["stock_taller"]

    store_stock = inv[inv.pool == "tienda"].pivot_table(
        index="SKU", columns="tienda", values="Cantidad en inventario", aggfunc="sum", fill_value=0
    )
    store_sales = cons[cons.tienda.notna()].pivot_table(
        index="SKU", columns="tienda", values="Cant. ordenada", aggfunc="sum", fill_value=0
    )
    for col in STORES:
        if col not in store_stock.columns:
            store_stock[col] = 0
        if col not in store_sales.columns:
            store_sales[col] = 0
    store_stock = store_stock[STORES].reindex(base.index).fillna(0)
    store_sales = store_sales[STORES].reindex(base.index).fillna(0).clip(lower=0)
    return base, store_stock, store_sales


def reference_skus():
    data = load_json_script(REF_HTML, "bf-embedded-data")
    skus = {str(r["sku"]).upper() for r in data["skus"]}
    return skus, data["meta"], data["summary"]


def plan_sku(store_stock: dict, taller: float, store_sales: dict, model_sales: dict):
    total = int(round(sum(store_stock.values()) + taller))
    if total <= 0:
        return []
    own = {s: store_sales.get(s, 0) for s in STORES if store_sales.get(s, 0) > 0}
    seeded = False
    if own:
        demand = sorted(own, key=lambda s: own[s], reverse=True)
    else:
        seeded = True
        ranked = sorted(STORES, key=lambda s: model_sales.get(s, 0), reverse=True)
        demand = [s for s in ranked if model_sales.get(s, 0) > 0][:3] or ranked[:3]
    if len(demand) > total:
        demand = demand[:total]
    targets = {s: 0 for s in STORES}
    remaining = total
    for s in demand:
        if remaining <= 0:
            break
        targets[s] += 1
        remaining -= 1
    for s in demand:
        if remaining <= 0:
            break
        ann = store_sales.get(s, 0)
        cap = 2 if seeded or ann <= 0 else max(2, int(round(ann / 6.0)))
        room = max(0, cap - targets[s])
        give = min(room, remaining)
        targets[s] += give
        remaining -= give
    need = {}
    surplus = {}
    for s in STORES:
        cur = store_stock.get(s, 0)
        tgt = targets[s]
        if cur + 1e-9 < tgt:
            need[s] = tgt - cur
        elif cur > tgt + 1e-9:
            surplus[s] = cur - tgt
    moves = []
    tall = float(taller)
    for s, n in sorted(need.items(), key=lambda kv: -kv[1]):
        take = min(n, tall)
        q = int(round(take))
        if q > 0:
            moves.append({"desde": "Taller", "hacia": s, "unidades": q})
            tall -= q
            need[s] -= q
    donors = sorted(surplus.items(), key=lambda kv: -kv[1])
    for s in list(need):
        n = need[s]
        if n < 0.5:
            continue
        new_donors = []
        for d, sup in donors:
            if n < 0.5 or sup < 0.5:
                new_donors.append((d, sup))
                continue
            take = min(n, sup)
            q = int(round(take))
            if q > 0:
                moves.append({"desde": d, "hacia": s, "unidades": q})
            n -= q
            new_donors.append((d, max(0, sup - q)))
        donors = new_donors
    return moves


def assemble(base, store_stock, store_sales, ref_skus):
    base = base.copy()
    base["modelo_u"] = base["modelo"].fillna("").astype(str).str.strip()
    drop_mask = base["modelo_u"].str.contains(DROP_RE)
    in_ref = base.index.to_series().isin(ref_skus) & ~drop_mask & base["stock"].gt(0)
    in_ref &= base["macro"].isin(["Manufactura", "Equipamiento"])
    forced = base["modelo_u"].isin(FORCE_FULL) & ~drop_mask & base["stock"].gt(0)
    forced &= base["macro"].isin(["Manufactura", "Equipamiento"])
    offer = base[in_ref | forced].copy()
    offer["origen"] = np.where(offer["modelo_u"].isin(FORCE_FULL), "Línea completa", "CC")
    offer["matriz"] = np.where(offer["origen"] == "Línea completa", "Línea", "CC")

    # Polvo: modelos CC con menos de 5 unidades. No arman evento.
    stock_modelo = offer.groupby("modelo_u")["stock"].sum()
    polvo_names = [
        m for m, st in stock_modelo.items() if st < MIN_MODEL_STOCK and m not in FORCE_FULL
    ]
    polvo = offer[offer.modelo_u.isin(polvo_names)].copy()
    offer = offer[~offer.modelo_u.isin(polvo_names)].copy()

    rows_m = []
    rows_s = []
    for modelo, part in offer.groupby("modelo_u"):
        qty = float(part.qty_12.sum())
        stock = float(part.stock.sum())
        rate = qty / 12.0
        cover = (stock / rate) if rate > 0 else None
        rango = rango_of(qty, cover if cover is not None else 999)
        tier = TIER_PCT[rango]
        di, topi, promo = [], [], []
        for price, cost, priced in zip(part.uprice, part.ucost, part.priced):
            if not priced:
                di.append(0.0)
                topi.append(False)
                promo.append(None)
                continue
            applied, capped = sku_discount(float(price), float(cost), tier)
            di.append(applied)
            topi.append(capped)
            promo.append(float(price) * (1 - applied) if applied >= MIN_DISC else None)
        part = part.copy()
        part["descuento"] = di
        part["tope"] = topi
        part["promo"] = promo
        part["rango"] = rango
        part["tier"] = tier
        priced_disc = part[part.descuento >= MIN_DISC]
        if len(priced_disc):
            d_eff = 1 - (
                (priced_disc.stock * priced_disc.uprice * (1 - priced_disc.descuento)).sum()
                / (priced_disc.stock * priced_disc.uprice).sum()
            )
            pvp = wavg(priced_disc.uprice, priced_disc.stock)
            pbf = wavg(priced_disc.uprice * (1 - priced_disc.descuento), priced_disc.stock)
        else:
            d_eff = None
            pvp = wavg(part.uprice, part.stock)
            pbf = None
        sin_precio = int((~part.priced).sum())
        rows_m.append(
            {
                "modelo": modelo,
                "macro": part["macro"].value_counts().index[0],
                "categoria": part["categoria"].value_counts().index[0],
                "origen": "Línea completa" if modelo in FORCE_FULL else "CC",
                "matriz": "Línea" if modelo in FORCE_FULL else "CC",
                "skus": int(part.index.nunique()),
                "rotacion": num(rate, 2),
                "qty12": num(qty, 0),
                "stock": num(stock, 0),
                "tienda": num(part.stock_tienda.sum(), 0),
                "taller": num(part.stock_taller.sum(), 0),
                "cover": num(cover, 1) if cover is not None else None,
                "rango": rango,
                "descuento": num(d_eff, 4),
                "precio": num(pvp, 2),
                "promo": num(pbf, 2),
                "tope": bool(part.tope.any()),
                "sinPrecio": sin_precio == len(part),
                "skusSinPrecio": sin_precio,
            }
        )
        for sku, r in part.iterrows():
            sku_rate = float(r.qty_12) / 12.0
            sku_cover = (float(r.stock) / sku_rate) if sku_rate > 0 else None
            rows_s.append(
                {
                    "sku": sku,
                    "modelo": modelo,
                    "gen": "" if str(r.genero) in ("nan", "None") else str(r.genero or ""),
                    "color": "" if str(r.color) in ("nan", "None") else str(r.color or ""),
                    "talla": "" if str(r.talla) in ("nan", "None") else str(r.talla or ""),
                    "matriz": r.matriz,
                    "stock": num(r.stock, 0),
                    "tienda": num(r.stock_tienda, 0),
                    "taller": num(r.stock_taller, 0),
                    "rotacion": num(sku_rate, 2),
                    "cover": num(sku_cover, 1) if sku_cover is not None else None,
                    "qty12": num(r.qty_12, 0),
                    "descuento": num(r.descuento, 4) if r.descuento >= MIN_DISC else None,
                    "precio": num(r.uprice, 2),
                    "promo": num(r.promo, 2) if r.promo is not None else None,
                    "tope": bool(r.tope),
                    "priced": bool(r.priced),
                }
            )
    models = pd.DataFrame(rows_m)
    skus = pd.DataFrame(rows_s)
    return offer, models, skus, polvo


def build_supply(skus, offer, store_stock, store_sales):
    disc = skus[skus.descuento.notna() & (skus.descuento >= MIN_DISC)]
    model_sales = {}
    for modelo, part in offer.groupby("modelo_u"):
        model_sales[modelo] = {s: float(store_sales.reindex(part.index).fillna(0)[s].sum()) for s in STORES}
    moves = []
    for row in disc.itertuples(index=False):
        sku = row.sku
        ss = {s: float(store_stock.at[sku, s]) if sku in store_stock.index else 0.0 for s in STORES}
        sv = {s: float(store_sales.at[sku, s]) if sku in store_sales.index else 0.0 for s in STORES}
        for m in plan_sku(ss, float(offer.at[sku, "stock_taller"]) if sku in offer.index else 0.0, sv, model_sales.get(row.modelo, {})):
            moves.append({"sku": sku, "modelo": row.modelo, "desde": m["desde"], "hacia": m["hacia"], "unidades": m["unidades"]})
    moves_df = pd.DataFrame(moves)
    tiendas = []
    offer_idx = offer.index
    for store in STORES:
        stock_hoy = float(store_stock.reindex(offer_idx).fillna(0)[store].sum())
        venta = float(store_sales.reindex(offer_idx).fillna(0)[store].sum())
        if moves_df.empty:
            entrar = salir = desde = 0
        else:
            entrar = int(moves_df.loc[moves_df.hacia == store, "unidades"].sum())
            salir = int(moves_df.loc[moves_df.desde == store, "unidades"].sum())
            desde = int(moves_df.loc[(moves_df.hacia == store) & (moves_df.desde == "Taller"), "unidades"].sum())
        quiebres = 0
        for modelo, part in offer.groupby("modelo_u"):
            sold = float(store_sales.reindex(part.index).fillna(0)[store].sum())
            have = float(store_stock.reindex(part.index).fillna(0)[store].sum())
            if sold > 0 and have <= 0:
                quiebres += 1
        tiendas.append(
            {
                "tienda": store,
                "stock": num(stock_hoy, 0),
                "venta": num(venta, 0),
                "entrar": entrar,
                "desdeTaller": desde,
                "salir": salir,
                "quiebres": quiebres,
            }
        )
    if moves_df.empty:
        pull = []
        top_moves = []
        bajar = mover = 0
    else:
        bajar = int(moves_df.loc[moves_df.desde == "Taller", "unidades"].sum())
        mover = int(moves_df.loc[moves_df.desde != "Taller", "unidades"].sum())
        pull_s = moves_df[moves_df.desde == "Taller"].groupby("modelo")["unidades"].sum().sort_values(ascending=False).head(8)
        pull = [{"modelo": i, "unidades": int(v)} for i, v in pull_s.items()]
        top = moves_df.sort_values("unidades", ascending=False).head(40)
        top_moves = top.to_dict(orient="records")
    return moves_df, tiendas, pull, top_moves, bajar, mover


def exclusion_rows(base):
    """Zenit, Refresh y Motion Loop, con el stock real, para que directiva vea la salida."""
    part = base[base.modelo.fillna("").astype(str).str.contains(DROP_RE) & base.stock.gt(0)].copy()
    rows = []
    for modelo, g in part.groupby(part.modelo.fillna("").astype(str).str.strip()):
        qty = float(g.qty_12.sum())
        rate = qty / 12.0
        if "ZENIT" in modelo.upper():
            motivo = "Zenit"
        elif "REFRESH" in modelo.upper():
            motivo = "Refresh"
        else:
            motivo = "Motion Loop"
        rows.append(
            {
                "motivo": motivo,
                "modelo": modelo,
                "macro": g["macro"].value_counts().index[0],
                "stock": num(g.stock.sum(), 0),
                "tienda": num(g.stock_tienda.sum(), 0),
                "taller": num(g.stock_taller.sum(), 0),
                "qty12": num(qty, 0),
                "rotacion": num(rate, 2),
                "cover": num(g.stock.sum() / rate, 1) if rate > 0 else None,
            }
        )
    rows.sort(key=lambda r: ({"Zenit": 0, "Refresh": 1, "Motion Loop": 2}[r["motivo"]], -r["stock"]))
    return rows


def polvo_rows(polvo):
    rows = []
    if polvo.empty:
        return rows
    for modelo, g in polvo.groupby(polvo.modelo.fillna("").astype(str).str.strip()):
        rows.append(
            {
                "motivo": "Menos de 5 unidades",
                "modelo": modelo,
                "stock": num(g.stock.sum(), 0),
                "qty12": num(g.qty_12.sum(), 0),
            }
        )
    rows.sort(key=lambda r: r["modelo"])
    return rows


def resumen(models, skus, bajar, mover):
    priced = skus[skus.descuento.notna()]
    if len(priced) and (priced.stock * priced.precio).sum():
        desc = 1 - (priced.stock * priced.promo).sum() / (priced.stock * priced.precio).sum()
    else:
        desc = None
    cover_ok = models.cover.dropna()
    u24 = int(models.loc[models.rango.isin(["24+", "dormido"]), "stock"].sum())
    return {
        "modelos": int(models.modelo.nunique()) if len(models) else 0,
        "skus": int(len(skus)),
        "unidades": num(models.stock.sum(), 0) if len(models) else 0,
        "tienda": num(models.tienda.sum(), 0) if len(models) else 0,
        "taller": num(models.taller.sum(), 0) if len(models) else 0,
        "desc_efectivo": num(desc, 4),
        "bajar": int(bajar),
        "mover": int(mover),
        "rotacion_mediana": num(models.rotacion.median(), 2) if len(models) else None,
        "cobertura_mediana": num(cover_ok.median(), 1) if len(cover_ok) else None,
        "unidades_largas": u24,
    }


def write_excel(models, skus, moves, tiendas, excl, polvo, res, texto):
    wb = Workbook()
    navy = PatternFill("solid", fgColor="0F6E56")
    cream = PatternFill("solid", fgColor="F7F4EE")
    sand = PatternFill("solid", fgColor="FFF6EA")
    white = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
    title = Font(name="Calibri", bold=True, size=18, color="1C1915")
    body = Font(name="Calibri", size=11, color="1C1915")
    thin = Border(
        left=Side(style="thin", color="E4DDD0"),
        right=Side(style="thin", color="E4DDD0"),
        top=Side(style="thin", color="E4DDD0"),
        bottom=Side(style="thin", color="E4DDD0"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    def head(ws, headers, fill=navy):
        for i, h in enumerate(headers, 1):
            cell = ws.cell(1, i, h)
            cell.fill = fill
            cell.font = white
            cell.alignment = wrap
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
        ws.oddHeader.left.text = "Black Friday · inventario de baja rotación"
        ws.oddFooter.right.text = "Página &P de &N"

    def widths(ws, cols):
        for i, w in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def paint(ws, row, cols, zebra=False):
        for c in range(1, cols + 1):
            cell = ws.cell(row, c)
            cell.font = body
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if zebra and row % 2 == 0:
                cell.fill = cream

    ws = wb.active
    ws.title = "01_Decision"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Black Friday · inventario de baja rotación"
    ws["A1"].font = title
    ws.merge_cells("A1:F1")
    ws["A2"] = "Octubre 2025 – septiembre 2026 · manufactura y equipamiento · para ofrecer en tienda"
    ws["A2"].font = Font(name="Calibri", size=12, color="5E584E")
    ws.merge_cells("A2:F2")
    ws["A4"] = texto
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A4"].font = body
    ws.merge_cells("A4:F4")
    ws.row_dimensions[4].height = 78
    labels = [
        ("Modelos", res["modelos"]),
        ("Unidades", res["unidades"]),
        ("En tienda", res["tienda"]),
        ("En taller", res["taller"]),
        ("Cobertura mediana (meses)", res["cobertura_mediana"]),
        ("Rotación mediana (u/mes)", res["rotacion_mediana"]),
        ("Descuento efectivo", res["desc_efectivo"]),
        ("Bajar de taller", res["bajar"]),
        ("Mover entre tiendas", res["mover"]),
        ("Unidades con cobertura larga", res["unidades_largas"]),
    ]
    ws["A6"] = "Indicador"
    ws["B6"] = "Valor"
    ws["A6"].fill = navy
    ws["B6"].fill = navy
    ws["A6"].font = white
    ws["B6"].font = white
    for i, (lab, val) in enumerate(labels, 7):
        ws.cell(i, 1, lab).font = body
        ws.cell(i, 2, val).font = Font(name="Calibri", bold=True, size=12)
        ws.cell(i, 1).border = thin
        ws.cell(i, 2).border = thin
        if lab == "Descuento efectivo":
            ws.cell(i, 2).number_format = "0.0%"
        elif isinstance(val, (int, float)):
            ws.cell(i, 2).number_format = "#,##0.00" if isinstance(val, float) and val < 100 else "#,##0"
    ws["A18"] = "Escalera de descuento"
    ws["A18"].font = Font(name="Calibri", bold=True, size=14)
    for i, h in enumerate(["Tramo de cobertura", "Descuento", "Modelos", "Unidades"], 1):
        cell = ws.cell(19, i, h)
        cell.fill = navy
        cell.font = white
    for i, (rid, label, pct) in enumerate(TIERS):
        part = models[models.rango == rid]
        ws.cell(20 + i, 1, label).font = body
        ws.cell(20 + i, 2, pct).font = body
        ws.cell(20 + i, 2).number_format = "0%"
        ws.cell(20 + i, 3, int(len(part))).font = body
        ws.cell(20 + i, 4, int(part.stock.sum()) if len(part) else 0).font = body
        ws.cell(20 + i, 4).number_format = "#,##0"
        for c in range(1, 5):
            ws.cell(20 + i, c).border = thin
            if i % 2:
                ws.cell(20 + i, c).fill = cream
    ws["A28"] = "Líneas que entran con el stock completo"
    ws["A28"].font = Font(name="Calibri", bold=True, size=14)
    for i, h in enumerate(["Modelo", "Unidades", "Tienda", "Taller", "Rotación/mes", "Cobertura", "Descuento"], 1):
        cell = ws.cell(29, i, h)
        cell.fill = navy
        cell.font = white
    full = models[models.origen == "Línea completa"].sort_values("stock", ascending=False)
    for i, r in enumerate(full.itertuples(index=False), 30):
        vals = [r.modelo, r.stock, r.tienda, r.taller, r.rotacion, r.cover if r.cover is not None else "Sin salida", r.descuento]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.font = body
            cell.border = thin
            cell.fill = sand
        ws.cell(i, 2).number_format = "#,##0"
        ws.cell(i, 3).number_format = "#,##0"
        ws.cell(i, 4).number_format = "#,##0"
        ws.cell(i, 5).number_format = "0.00"
        if isinstance(r.cover, float):
            ws.cell(i, 6).number_format = "0.0"
        if r.descuento is not None:
            ws.cell(i, 7).number_format = "0%"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 28
    for col in "CDEFG":
        ws.column_dimensions[col].width = 18
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.oddHeader.left.text = "Black Friday · decisión"
    ws.print_title_rows = "1:2"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.zoomScale = 120

    headers = [
        "Modelo", "Línea", "Categoría", "Origen", "SKUs", "Rotación/mes", "Venta 12 meses",
        "Stock total", "Stock tiendas", "Stock taller", "Cobertura (meses)", "Descuento",
        "Precio", "Precio Black Friday", "Piso de precio",
    ]

    def option_sheet(name, frame):
        ws = wb.create_sheet(name)
        head(ws, headers)
        for i, r in enumerate(frame.itertuples(index=False), 2):
            vals = [
                r.modelo, r.macro, r.categoria, r.origen, r.skus, r.rotacion, r.qty12,
                r.stock, r.tienda, r.taller, r.cover if r.cover is not None else "Sin salida",
                r.descuento, r.precio, r.promo, "Sí" if r.tope else "",
            ]
            for c, v in enumerate(vals, 1):
                ws.cell(i, c, None if v is None else v)
            paint(ws, i, len(headers), zebra=True)
            ws.cell(i, 6).number_format = "0.00"
            ws.cell(i, 7).number_format = "#,##0"
            for col in (8, 9, 10):
                ws.cell(i, col).number_format = "#,##0"
            if isinstance(r.cover, float):
                ws.cell(i, 11).number_format = "0.0"
            if r.descuento is not None:
                ws.cell(i, 12).number_format = "0.0%"
            for col in (13, 14):
                if ws.cell(i, col).value is not None:
                    ws.cell(i, col).number_format = '"$"#,##0.00'
        widths(ws, [36, 16, 32, 18, 10, 14, 16, 14, 14, 14, 18, 12, 12, 20, 16])
        ws.auto_filter.ref = f"A1:O{max(1, len(frame)+1)}"
        ws.auto_filter.add_sort_condition("A2:A2")
        ws.print_title_rows = "1:1"
        ws.page_setup.fitToHeight = 0
        ws.oddFooter.left.text = name.replace("_", " ")

    by_rot = models.sort_values(["rotacion", "stock"], ascending=[True, False])
    by_stock = models.sort_values(["stock", "rotacion"], ascending=[False, True])
    option_sheet("02_Por_rotacion", by_rot)
    option_sheet("03_Por_inventario", by_stock)

    ws = wb.create_sheet("04_Detalle_SKU")
    h = ["Modelo", "SKU", "Género", "Color", "Talla", "Origen", "Rotación/mes", "Venta 12 meses",
         "Stock", "Tienda", "Taller", "Cobertura (meses)", "Descuento", "Precio", "Precio Black Friday"]
    head(ws, h)
    sk_sorted = skus.sort_values(["modelo", "stock"], ascending=[True, False])
    origen = models.set_index("modelo")["origen"]
    for i, r in enumerate(sk_sorted.itertuples(index=False), 2):
        vals = [
            r.modelo, r.sku, r.gen, r.color, r.talla, origen.get(r.modelo, "CC"),
            r.rotacion, r.qty12, r.stock, r.tienda, r.taller,
            r.cover if r.cover is not None else "Sin salida",
            r.descuento, r.precio, r.promo,
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, None if v is None else v)
        paint(ws, i, len(h), True)
        ws.cell(i, 7).number_format = "0.00"
        for col in (8, 9, 10, 11):
            ws.cell(i, col).number_format = "#,##0"
        if isinstance(r.cover, float):
            ws.cell(i, 12).number_format = "0.0"
        if r.descuento is not None:
            ws.cell(i, 13).number_format = "0.0%"
        for col in (14, 15):
            if ws.cell(i, col).value is not None:
                ws.cell(i, col).number_format = '"$"#,##0.00'
    widths(ws, [32, 16, 14, 18, 12, 18, 14, 16, 12, 12, 12, 18, 12, 12, 20])
    ws.auto_filter.ref = f"A1:O{max(1, len(sk_sorted)+1)}"
    ws.print_title_rows = "1:1"

    ws = wb.create_sheet("05_Movimientos")
    h = ["Modelo", "SKU", "Desde", "Hacia", "Unidades"]
    head(ws, h)
    if not moves.empty:
        ordered = moves.sort_values(["unidades", "modelo"], ascending=[False, True])
        for i, r in enumerate(ordered.itertuples(index=False), 2):
            for c, v in enumerate([r.modelo, r.sku, r.desde, r.hacia, int(r.unidades)], 1):
                ws.cell(i, c, v)
            paint(ws, i, 5, True)
            ws.cell(i, 5).number_format = "#,##0"
        ws.auto_filter.ref = f"A1:E{len(ordered)+1}"
    widths(ws, [36, 16, 22, 22, 14])
    ws.print_title_rows = "1:1"

    ws = wb.create_sheet("06_No_entra")
    h = ["Motivo", "Modelo", "Línea", "Unidades", "Tienda", "Taller", "Venta 12 meses", "Rotación/mes", "Cobertura (meses)"]
    head(ws, h)
    extra = []
    for r in excl:
        extra.append([r["motivo"], r["modelo"], r["macro"], r["stock"], r["tienda"], r["taller"], r["qty12"], r["rotacion"], r["cover"] if r["cover"] is not None else "Sin salida"])
    for r in polvo:
        extra.append([r["motivo"], r["modelo"], "", r["stock"], None, None, r["qty12"], None, None])
    sinp = models[models.sinPrecio]
    for r in sinp.itertuples(index=False):
        extra.append(["Sin precio en el ABC", r.modelo, r.macro, r.stock, r.tienda, r.taller, r.qty12, r.rotacion, r.cover if r.cover is not None else "Sin salida"])
    for i, vals in enumerate(extra, 2):
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, None if v is None else v)
        paint(ws, i, len(h), True)
        for col in (4, 5, 6, 7):
            if isinstance(ws.cell(i, col).value, (int, float)):
                ws.cell(i, col).number_format = "#,##0"
        if isinstance(ws.cell(i, 8).value, float):
            ws.cell(i, 8).number_format = "0.00"
        if isinstance(ws.cell(i, 9).value, float):
            ws.cell(i, 9).number_format = "0.0"
    widths(ws, [24, 36, 16, 14, 12, 12, 16, 14, 18])
    if extra:
        ws.auto_filter.ref = f"A1:I{len(extra)+1}"
    ws.print_title_rows = "1:1"
    note_row = len(extra) + 3
    ws.cell(note_row, 1, "Zenit, Refresh y Motion Loop salen de la oferta por decisión de esta revisión. El polvo (menos de 5 unidades) es cola CC del cuadro de rotación: queda visible aquí y no arma el evento.").alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws.row_dimensions[note_row].height = 36

    ws = wb.create_sheet("07_Abasto_tiendas")
    h = ["Tienda", "Stock hoy", "Venta 12 meses", "Deben entrar", "Desde taller", "Salen hacia otra tienda", "Modelos en quiebre"]
    head(ws, h)
    for i, t in enumerate(tiendas, 2):
        vals = [t["tienda"], t["stock"], t["venta"], t["entrar"], t["desdeTaller"], t["salir"], t["quiebres"]]
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, v)
        paint(ws, i, len(h), True)
        for col in range(2, 8):
            ws.cell(i, col).number_format = "#,##0"
    widths(ws, [22, 14, 16, 16, 16, 24, 22])
    ws.cell(11, 1, "Cada tienda que vendió el SKU en el año queda con al menos una unidad. El taller cubre primero. El tope por tienda es cerca de dos meses de su propia venta. Lo que sobra se queda en taller para la semana del evento.").alignment = Alignment(wrap_text=True)
    ws.merge_cells("A11:G11")
    ws.row_dimensions[11].height = 32
    ws.print_title_rows = "1:1"

    # Gráfico de unidades por tramo, al lado de la escalera.
    chart = BarChart()
    chart.type = "col"
    chart.title = "Unidades por tramo de descuento"
    chart.y_axis.title = "Unidades"
    chart.x_axis.title = None
    data = Reference(wb["01_Decision"], min_col=4, min_row=19, max_row=19 + len(TIERS))
    cats = Reference(wb["01_Decision"], min_col=1, min_row=20, max_row=19 + len(TIERS))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.style = 10
    chart.y_axis.majorGridlines = None
    chart.width = 18
    chart.height = 8
    wb["01_Decision"].add_chart(chart, "A38")

    # Validación suave: el origen solo admite dos valores, para quien filtre el libro.
    dv = DataValidation(type="list", formula1='"CC,Línea completa"', allow_blank=False)
    dv.error = "Origen"
    dv.errorTitle = "Origen"
    dv.add("D2:D500")
    wb["02_Por_rotacion"].add_data_validation(dv)

    wb.properties.title = "Propuesta Black Friday · baja rotación"
    wb.properties.creator = "Cuadro"
    wb.save(XLSX_OUT)


def main():
    base, store_stock, store_sales = build_frames()
    ref_skus, ref_meta, ref_summary = reference_skus()
    offer, models, skus, polvo = assemble(base, store_stock, store_sales, ref_skus)
    moves, tiendas, pull, top_moves, bajar, mover = build_supply(skus, offer, store_stock, store_sales)
    excl = exclusion_rows(base)
    polvo_list = polvo_rows(polvo)
    res = resumen(models, skus, bajar, mover)

    wanted = ["JACKET CAB", "JACKET DAMA", "JACKET KIDS", "SHORT SPORT KIDS", "SHORT SPORT R1 DAMA", "SHORT SPORT R1 CAB"]
    have = set(models.modelo)
    missing = [m for m in wanted if m not in have]
    if missing:
        raise SystemExit(f"Faltan líneas pedidas: {missing}")
    banned = models.modelo.str.contains(DROP_RE)
    if banned.any():
        raise SystemExit(f"Se colaron excluidos: {models.loc[banned, 'modelo'].tolist()}")
    if models.stock.sum() <= 0:
        raise SystemExit("La oferta quedó sin unidades")

    def miles(n):
        return f"{int(round(float(n))):,}".replace(",", ".")

    focus = models[models.modelo.isin(wanted)].set_index("modelo")
    bits = []
    for name, label in (
        ("JACKET CAB", "Jacket Caballero"),
        ("JACKET DAMA", "Jacket Dama"),
        ("JACKET KIDS", "Jacket Kids"),
        ("SHORT SPORT KIDS", "Short Sport Kids"),
        ("SHORT SPORT R1 DAMA", "Short Sport R1 Dama"),
        ("SHORT SPORT R1 CAB", "Short Sport R1 Caballero"),
    ):
        r = focus.loc[name]
        cover = "sin salida en el año" if r.cover is None else f"{r.cover:.0f} meses de cobertura"
        disc = "sin precio" if r.descuento is None else f"{r.descuento * 100:.0f}%"
        bits.append(f"{label}: {miles(r.stock)} unidades, {cover}, descuento {disc}")

    texto = (
        f"Recomendamos leer la oferta por rotación: lo más lento primero. "
        f"Son {miles(res['modelos'])} modelos y {miles(res['unidades'])} unidades "
        f"({miles(res['tienda'])} en tienda y {miles(res['taller'])} en taller). "
        f"El descuento va del 15% al 40% según los meses de cobertura, y el precio promo se sostiene en costo × 1,10 o por encima. "
        f"La otra lectura ordena esta misma oferta por inventario, para ver primero lo que más unidades ocupa. "
        f"Jacket 1.0 y Short Sport entran con el stock completo de la línea. "
        + ". ".join(bits)
        + "."
    )

    payload = {
        "meta": {
            "periodo": "Octubre 2025 – septiembre 2026",
            "generado": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "fuente": "Cuadro de rotación categoría C, inventario y ventas de la propuesta",
            "refSkus": ref_summary.get("skus_c_total"),
            "refUnidades": ref_summary.get("unidades_stock"),
        },
        "recomendacion": {"id": "rotacion", "texto": texto},
        "opciones": {
            "rotacion": {
                "nombre": "Opción A · Rotación",
                "frase": "Lo más lento primero. La misma oferta, ordenada por unidades al mes.",
                "sort": "rotacion_asc",
            },
            "inventario": {
                "nombre": "Opción B · Inventario",
                "frase": "Lo que más stock tiene primero. La misma oferta, ordenada por unidades en tienda y taller.",
                "sort": "stock_desc",
            },
        },
        "resumen": res,
        "tiers": [{"id": i, "label": lab, "pct": pct} for i, lab, pct in TIERS],
        "modelos": models.to_dict(orient="records"),
        "lineas": skus.to_dict(orient="records"),
        "tiendas": tiendas,
        "pull": pull,
        "movimientos": top_moves,
        "excluidos": excl,
        "polvo": {"modelos": len(polvo_list), "unidades": num(sum(r["stock"] for r in polvo_list), 0), "lista": polvo_list[:12]},
        "foco": models[models.origen == "Línea completa"].sort_values("stock", ascending=False).to_dict(orient="records"),
    }
    html = HTML_TEMPLATE.read_text(encoding="utf-8")
    if html.count("__DATA__") != 1:
        raise SystemExit("La plantilla debe tener un solo __DATA__")
    html = html.replace("__DATA__", json.dumps(clean(payload), ensure_ascii=False, separators=(",", ":")))
    HTML_OUT.write_text(html, encoding="utf-8")
    write_excel(models, skus, moves, tiendas, excl, polvo_list, res, texto)

    print(texto)
    print("--- foco ---")
    print(focus[["stock", "tienda", "taller", "rotacion", "cover", "descuento", "rango", "skus"]].to_string())
    print("--- rotación (15 más lentos) ---")
    print(models.sort_values("rotacion")[["modelo", "rotacion", "stock", "cover", "descuento"]].head(15).to_string(index=False))
    print("--- inventario (12 más grandes) ---")
    print(models.sort_values("stock", ascending=False)[["modelo", "rotacion", "stock", "cover", "descuento", "origen"]].head(12).to_string(index=False))
    print("excluidos", [(r["motivo"], r["modelo"], r["stock"]) for r in excl])
    print("polvo", len(polvo_list), "modelos", payload["polvo"]["unidades"], "u")
    print("sin precio", models.loc[models.sinPrecio, "modelo"].tolist())
    print("modelos", res["modelos"], "skus", res["skus"], "uds", res["unidades"], "bajar", bajar, "mover", mover)
    print("html", HTML_OUT, "xlsx", XLSX_OUT)


if __name__ == "__main__":
    main()
